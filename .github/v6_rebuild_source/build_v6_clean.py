# coding: utf-8
"""Clean-rebuild V6 using the 9.30 original Word and audited V5 source research."""
from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from zoneinfo import ZoneInfo
import re,shutil,tarfile
ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/'.github/v6_rebuild_source/10201024_V6_重建发布稿.md'
OUT=ROOT/'1019/v6/修改稿'
OLD=ROOT/'1019/v5/修改稿'
stamp='10201024'
raw=SRC.read_text(encoding='utf-8')
assert '作为评委' not in raw and '站在评委' not in raw and '我更想追问' not in raw
groups=[]
for line in raw.splitlines():
    if line.startswith('## P'):
        m=re.match(r'## (P\d+(?:[—-]P\d+)?)｜(.+)',line)
        assert m,line
        groups.append((m.group(1),m.group(2),[]))
    elif line.startswith(('A：','B：')):
        groups[-1][2].append((line[0],line[2:]))
han=lambda s:len(re.findall('[\u4e00-\u9fff]',str(s)))
a=sum(han(t) for _,_,lines in groups for r,t in lines if r=='A')
b=sum(han(t) for _,_,lines in groups for r,t in lines if r=='B')
assert a==1057 and b==444,(a,b)
assert len(groups)==15,len(groups)
page_map={}
for code,title,lines in groups:
    m=re.fullmatch(r'P(\d+)(?:[—-]P(\d+))?',code)
    assert m,code
    for page in range(int(m.group(1)),int(m.group(2) or m.group(1))+1):page_map[page]=(code,title,lines)
assert sorted(page_map)==list(range(1,26))
# Delete the previous V6 directory only after document inputs and ratios have passed their checks.
if OUT.parent.exists():shutil.rmtree(OUT.parent)
OUT.mkdir(parents=True)
def font(r,name,size,bold=None):
    r.font.name=name;r.font.size=Pt(size)
    if bold is not None:r.bold=bold
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
def para(d,text,style=None):return d.add_paragraph(text,style=style)
def head(d,text):return para(d,text,'Heading 2')
def table(d,heads,rows):
    t=d.add_table(rows=1,cols=len(heads));t.style='Table Grid'
    for c,s in zip(t.rows[0].cells,heads):c.text=str(s)
    for line in rows:
        cells=t.add_row().cells
        for i,s in enumerate(line):cells[i].text=str(s)
    return t
def make_speech(stage):
    d=Document(ROOT/'智烤管家演讲讲稿9.30.docx')
    for node in list(d._element.body):
        if node.tag!=qn('w:sectPr'):d._element.body.remove(node)
    if stage:
        d.sections[0].top_margin=Cm(1.65)
        d.sections[0].bottom_margin=Cm(1.65)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('智烤管家——基于健康评估的烤房用投管一体化解决方案'),'黑体',14,True)
    for code,title,lines in groups:
        if stage:
            p=d.add_paragraph()
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(0)
            r=p.add_run('【'+code+'】'+title)
            font(r,'黑体',9,True);r.font.color.rgb=RGBColor(85,104,134)
            if code=='P2':
                para(d,'（静默播放30秒开场视频；其他演示不得额外叠加播放。）')
        for r,t in lines:
            p=d.add_paragraph()
            p.paragraph_format.line_spacing=1.07 if stage else 1.12
            p.paragraph_format.space_after=Pt(2 if stage else 4)
            font(p.add_run(r+'：'),'黑体',14,True)
            font(p.add_run(t),'仿宋_GB2312',14)
    return d
formal=OUT/f'{stamp}_04_V6重建_8分钟双人正式演讲净稿.docx'
stage=OUT/f'{stamp}_05_V6重建_25页PPT双人舞台稿.docx'
make_speech(False).save(formal);make_speech(True).save(stage)
# Award study: preserve all V5 prior source-by-source analysis, then add actual QC learnings.
award=Document(OLD/'10192342_01_V5_十二篇获奖现场解说词学习与创新借鉴.docx')
head(award,'V6重建：中质协15分钟QC发布稿专项学习')
for h,t in [
('借鉴双人互动逻辑，不模仿身份','QC发布稿B是有一线维护经验的共同发布人，能自然答技术细节、回应方法难度、成本、长效问题。V6没有演员式角色介绍，也删除一切“作为评委”等不自然自称。'),
('先证问题，再讲技术','QC从单次维护40分钟调查到瓶颈工序分析，再呈现措施验证和标准化巩固。V6将这种“问题—方法—结果—长效”结构缩进8分钟，而不是罗列软件菜单。'),
('技术句要有业务解释','模型1讲AHP部件权重、80/10/10健康分、寿命估算、风险TOP3及清单；模型2讲有效烤能与理论需求、AHP加熵权组合、TOPSIS及地图沙盘，解释为什么能指导修建决策。'),
('绝不混用项目成绩','QC课题维护40分钟至17.9分钟、经济效益63.6万元等只属于独立QC成果，不属于智烤管家。智烤仍采用2442建档、48调剂、18报修、531问卷96.8%等官方数据。'),
('A/B台词不是“角色扮演”','B承接一线问题、压实技术目标、归纳服务价值，A主讲方法与事实。篇幅为1057:444，约70.42%:29.58%。')
]:
    head(award,h);para(award,t)
table(award,['来源','V6的实际改法'],[
('去年12篇优秀解说词','借鉴三问开题、反差排比、节奏转换，不套用别人的项目指标'),
('QC 0907发布词','借鉴问题归因、证据检验与组织固化，不把15分钟直接压缩成8分钟'),
('9.30原稿','保留匠心工坊、资源无统筹、投入无量化、管护无闭环'),
('V1—V5','吸收现场感、技术深度、双模耦合、服务实效与四队制度')])
award.save(OUT/f'{stamp}_01_12篇获奖解说词及QC发布方法学习.docx')
# Comprehensive deep study retains original image evidence from V5.
research=Document(OLD/'10192342_02_V5_项目五万字级全文件研究与双模创新增订.docx')
head(research,'V6重建增订：证据导向的技术与管理体系研究')
for text in [
'本报告继承V5完整项目全文件研究与图片证据。继承的摘录和图件属于原资料，并非本轮全部重新创作的新内容；本轮新增的是双人发布论证结构、技术边界核对与项目组织逻辑。',
'4大部位、18个关键清查项用于现场摸底，P8图示的12个计算参数用于健康评估。两种粒度需要字段映射表，不应在演讲中混称为18个算法参数。',
'健康模型以基础档案、部件数据、检查评价及管护维修为输入，以AHP确定部件权重，80%部件状态、10%管护、10%人工评价算整栋健康分，估算剩余寿命，输出风险部位前三和维修清单。',
'PPT第8页使用状态和剩余寿命合成公式，第9页出现Weibull拟合表述。正式上场前须由实施团队确认现行版本，V6主稿不虚构Weibull模型的训练过程或预测精度。',
'村级发展导引以“数、需、评、推”为技术路线：汇总资产与生产村情、折算健康状态后的有效烤能Q与理论需求D、用AHP＋熵权组合赋权和TOPSIS候选排序、通过地图沙盘模拟新建栋数和位置。',
'模型1输出的健康状态改变村级有效烤能，模型2把有效烤能纳入供需缺口，从而影响新建方案。模型协同不等于自动批准建设，仍须现场踏勘和既定审批。',
'三智和四队组成数据闭环与线下执行链。业务申请、派单、处置、验收、评价、回写依次产生数据，算法得以更新；四队、制度和基金使运维拥有长期责任主体。',
'官方指标边界：2442座全口径建档；健康分低于60的340余座是预警对象；48次调剂、18次报修，16条平台工单平均响应0.38小时；531份问卷满意度96.8%；1363栋统筹管护，年度在用1672栋约81%；近190万元是基金总额。'
]:para(research,text)
research.save(OUT/f'{stamp}_02_项目五万字级深度研究及V6创新增订.docx')
official=Document(OLD/'10192342_03_V5_官方叙述一万字级研究校订版.docx')
head(official,'V6重建：QC式专业语言与严谨口径')
for text in [
'台词不采用“我作为评委”“站在评委视角”“评委更想知道”等表演式设问；B只在项目发布者内部自然承接业务问题。开场常规问候可以称领导和专家，正文无需对评审身份作任何假定。',
'官方成果表达优先采用明确范围、分母和时点的数据。16条平台工单的0.38小时平均响应不能误称18次报修全部平均维修时长。',
'基金近190万元是总额；1363栋与1672栋的比值约81%；340多座低分烤房属于风险研判对象，均不等于已经维修完成。',
'技术内容可通俗解释，但必须保留AHP、权重结构、TOPSIS、熵权、有效烤能和沙盘推演，使信息技术与烟叶业务领域听众都能准确理解。',
'中质协QC课题是独立案例，维护设备改进、检测工具、节省工时及经济效益数值均不得计入智烤管家。',
'本项目创新的准确说法是存量健康评估与村级有效烤能测算形成联动，支撑维修和新建的比较与优先建议，决策权仍在制度和人员。'
]:para(official,text)
official.save(OUT/f'{stamp}_03_官方叙述一万字级规范研究与QC校订.docx')
# PPT 25-page coverage document.
d=Document();d.add_heading('V6重建：25页PPT与AB口播和时间核验',0)
para(d,f'口述A={a}字、B={b}字，共{a+b}字，A占{a/(a+b)*100:.2f}%。按220汉字/分、30秒视频和20秒换页，预计总时长{(a+b)*60/220+50:.1f}秒；为预算而非真实计时。')
for p in range(1,26):
    code,title,lines=page_map[p];head(d,f'第{p}页｜{title}')
    first=int(re.search(r'P(\d+)',code).group(1))
    if p==2:para(d,'静默播放视频30秒。原素材时长应先裁剪，现场不加读台词。')
    elif p==first:
        for r,t in lines:para(d,r+'：'+t)
    else:para(d,'本页承接同组前一页的口述，以界面或技术图形支撑，不另加重复台词。')
    if p in (8,9):para(d,'P8寿命合成公式与P9 Weibull文字说明，赛前须由模型研发团队确认一致性。')
    if p in (14,16,17,19):para(d,'本页附加视频不得直接叠加播放，若播放必须等量缩减台词。')
d.save(OUT/f'{stamp}_06_25页PPT逐页口播及8分钟计时表.docx')
# Comparison and defence QA
d=Document();d.add_heading('V6重建：9.30原稿与V1—V5优势融合、QC对照',0)
table(d,['原版优点或问题','V6处理'],[
('旧V6自称评委等生硬台词','全部删除；B以发布同伴身份自然设问和承接事实'),
('9.30匠心工坊和三缺','保留烟农收成的情感锚点、三个管理问题'),
('V1—V3情绪与技术','保留排比，完整讲两模输入、计算、输出，不泛化'),
('V4—V5业绩与治理','保留调剂、报修、满意度、四队和基金的真实统计口径'),
('QC发布词论证方式','问题→方法→实证→固化，保留专业性而不混用成果'),
('PPT与技术一致性','备答确认18清查项与12计算项映射、P8/P9寿命估算方式')])
head(d,'需要赛前核实的技术细节')
for s in [
'1. 现场清查18项与模型计算12参数如何映射；请研发人员携完整字段字典。',
'2. PPT第8页寿命合成路径和第9页Weibull文字是否对应同一模型版本；不要将未证实的拟合精度写进口播。',
'3. 18次报修与16条平台工单的统计范围差异；响应时间与完成时间不得混淆。',
'4. 秩序口径：基金总额不是期末余额；预警对象不是已修复对象；新建候选排名不等于已经获批。',
'5. 30秒视频、25页换页和现场停顿需要联合计时；本稿书面预算约7分39秒，仍须实际彩排。'
]:para(d,s)
d.save(OUT/f'{stamp}_07_原稿历版QC融合修改对照与技术备答.docx')
# Name/manifest for transparent handoff
(OUT/f'{stamp}_README_核验说明.md').write_text(
f'# V6重新建立（已清理旧V6）\n- A={a}汉字，B={b}汉字，A占比={a/(a+b)*100:.2f}%\n- 220字/分＋30秒视频＋20秒翻页≈{(a+b)*60/220+50:.1f}秒；非现场实测。\n- 项目与官方研究报告是在V5已完成的五万字/一万字源证研究卷基础上的增订，不是本次独立重新写出五万、一万字原创。\n- 双模型技术与官方指标保留，删掉全部作为评委/评委视角表演式台词。\n- 原9.30稿和最新版PPT没有改动。\n',encoding='utf-8')
(OUT/f'{stamp}_V6完整演讲正文.md').write_text(raw,encoding='utf-8')
docs=sorted(OUT.glob('*.docx'))
assert len(docs)==7,[x.name for x in docs]
with tarfile.open(OUT/f'{stamp}_智烤管家V6重建_七份Word.tar.gz','w:gz') as tf:
    for f in docs:tf.add(f,arcname=f.name)
    tf.add(OUT/f'{stamp}_V6完整演讲正文.md',arcname='V6完整演讲正文.md')
print('V6 built',stamp,a,b,len(docs),[(p.name,p.stat().st_size) for p in docs])
