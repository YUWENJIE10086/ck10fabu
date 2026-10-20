# -*- coding:utf-8 -*-
"""2026-10-20 V6 from official 9.30 template; entirely replace old V6 only after data checks."""
from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re,shutil,tarfile
ROOT=Path(__file__).resolve().parents[2]
SEED=ROOT/'.github/v6_rebuild_source'
OUT=ROOT/'1019/v6/修改稿'
STAMP='10201048'
FULL=(SEED/(STAMP+'_V6_2000完整版.md')).read_text('utf-8')
SHORT=(SEED/(STAMP+'_V6_8分钟上场版.md')).read_text('utf-8')
def han(s):return len(re.findall(r'[\u4e00-\u9fff]',s))
def parts(s):
 groups=[]
 for line in s.splitlines():
  if line.startswith('## 【'):
   groups.append({'title':line[3:],'lines':[]})
  elif line.startswith(('A：','B：')):groups[-1]['lines'].append((line[0],line[2:]))
  elif line.startswith('（') and groups:groups[-1]['lines'].append(('CUE',line))
 return groups
def stats(s):
 g=parts(s)
 a=sum(han(t) for block in g for r,t in block['lines'] if r=='A')
 b=sum(han(t) for block in g for r,t in block['lines'] if r=='B')
 return (a,b)
assert stats(FULL)==(1411,599),stats(FULL)
assert stats(SHORT)==(1148,511),stats(SHORT)
assert not any(x in FULL+SHORT for x in ('作为评委','我更想追问','站在评委','评委视角'))
for key in ['GPS定位','18个','22个','AHP','80%','10%','TOP3','熵权','TOPSIS','沙盘','531','96.8%','1363','0.38','3小时']:
 assert key in FULL,key
# delete old V6 only after both manuscripts validate.
if OUT.parent.exists():shutil.rmtree(OUT.parent)
OUT.mkdir(parents=True)
def runfont(r,name,size,bold=None,color=None):
 r.font.name=name;r.font.size=Pt(size)
 if bold is not None:r.bold=bold
 r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
 if color:r.font.color.rgb=RGBColor(*color)
def title(d,s):
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 runfont(p.add_run(s),'黑体',14,True)
def head(d,s):return d.add_paragraph(s,style='Heading 2')
def text(d,s):return d.add_paragraph(s)
def tab(d,cols,rows):
 t=d.add_table(rows=1,cols=len(cols));t.style='Table Grid'
 for cell,s in zip(t.rows[0].cells,cols):cell.text=str(s)
 for row in rows:
  cs=t.add_row().cells
  for i,s in enumerate(row):cs[i].text=str(s)
 return t
def make_speech(s,name,cues=False):
 d=Document(ROOT/'1004/智烤管家演讲讲稿9.30.docx')
 for n in list(d._element.body):
  if n.tag!=qn('w:sectPr'):d._element.body.remove(n)
 if cues:
  d.sections[0].top_margin=Cm(1.55);d.sections[0].bottom_margin=Cm(1.55)
 title(d,'智烤管家——基于健康评估的烤房用投管一体化解决方案')
 for group in parts(s):
  if cues:
   p=d.add_paragraph();p.paragraph_format.keep_with_next=True
   runfont(p.add_run(group['title']),'黑体',9,True,(85,104,134))
  for who,line in group['lines']:
   if who=='CUE':
    if cues:
     p=d.add_paragraph();runfont(p.add_run(line),'楷体',9,False,(120,120,120))
    continue
   p=d.add_paragraph()
   p.paragraph_format.space_after=Pt(2 if cues else 4)
   p.paragraph_format.line_spacing=1.06 if cues else 1.12
   runfont(p.add_run(who+'：'),'黑体',14,True)
   runfont(p.add_run(line),'仿宋_GB2312',14)
 d.save(OUT/name)
make_speech(FULL,f'{STAMP}_04_V6约2000字AB决赛完整版.docx')
make_speech(SHORT,f'{STAMP}_05_V6八分钟AB正式上场稿.docx')
make_speech(SHORT,f'{STAMP}_05B_V6_25页PPT双人舞台稿.docx',True)
# Winner and QC study: independent V6 interpretation grounded in twelve txt sources.
d=Document();d.add_heading('智烤管家V6：12篇获奖现场解说词与QC发布方法研究',0)
text(d,'以下逐条对应仓库“读取资料/1学习优秀的”的12份TXT，并参考根目录《缩短烤房设备单次维护作业耗时_QC成果发布词0907.docx》。迁移的是表达方法，绝不挪用其他作品的数据和成果。')
award=[
('1.txt','3D智慧对标','以真实指标差距引出对标工具，主题突出、过程清晰','借鉴“业务问题先于技术方案”。'),
('2.txt','网络情报小剧场','场景先行使抽象业务具体化','以烟农三愁开场，不刻意表演身份。'),
('3.txt','AI卷烟营销','多源数据与业务调控逻辑相互证明','模型要讲输入、算法、产出和管理用途。'),
('4.txt','监管研判','识别对象明确，突出判别精准价值','说明健康模型为什么输出风险TOP3。'),
('5.txt','寄递数据分析','数据采集到应用的一条业务链','对应1库→2模→3智→4队。'),
('6.txt','会思考的食堂','生活化比喻减少术语理解门槛','“体检”仅作解释，不替代公式方法。'),
('7.txt','滞销品规','利益痛点与排比增强记忆','保留“有房未必调得动”。'),
('15.txt','烟叶收购预警','扎根一线作业场景，专业仍通俗','保留烟叶生产和烘烤业务术语。'),
('19.txt','财务监管','强调台账、审核和可追溯','突出数据回写、资金与验收职责。'),
('30.txt','雪茄三问','真实业务问题自然形成章节衔接','B负责追问修谁、建哪、谁落实。'),
('9月9日我的.txt','原型口述','保留项目一库双模平台架构','延续原项目核心，不更换技术事实。'),
('yus07.txt','三个没数','短句排比让痛点一听即懂','3愁→3缺，现场动机清晰。')]
for i,(source,name,good,apply) in enumerate(award,1):
 head(d,f'{i}. {source}：{name}')
 text(d,'值得学习：'+good);text(d,'V6吸收：'+apply)
head(d,'QC发布词的重点迁移：专业双人发布而非口语化捧哏')
for z in ['先做现状调查，找准瓶颈，再说明措施如何针对瓶颈，最后展示客观结果与固化机制。V6对应“3愁—3缺—1234—2模—3智—4队—实效”。',
'QC中B有真实发布身份，问题与答复是共同论证而非模拟评委。V6让B在维修先后、选址投入、长效执行处自然转场，不自称“作为评委”。',
'QC自身的维护时间、经济效益与本项目相互独立，不能作为“智烤管家”成果。',
'舞台气势来自数字和逻辑，不靠连续喊口号；结尾只保留1次归纳和“聚数成方，一脉致远”。']:text(d,z)
d.save(OUT/f'{STAMP}_01_12篇获奖词及QC发布方法学习.docx')
# 5w technical research and 1w official research: source inheritance is conspicuously documented.
v5=ROOT/'1019/v5/修改稿'
research=Document(v5/'10192342_02_V5_项目五万字级全文件研究与双模创新增订.docx')
head(research,'V6技术增订：2000字决赛发布稿的证据边界与创新链')
for z in [
'本卷保留前序逐项整理的原始资料索引、摘录和图片，本次新增章节并非重新创作全部5万汉字；应将引用原始资料与改稿研究区分。',
'数据采集采用手机、GPS、照片、动态表单，按4大部位18个关键项目查验正常/损坏/缺失；22个低代码应用、6类分析报表支撑2442座烤房全口径建档，入库数量较原设计增加近52%。',
'现场采集18项目与P8健康模型图12个计算参数属于不同数据层级，正式答辩需准备映射表；P8寿命合成式与P9 Weibull文字应由研发团队统一现行版本。',
'模型1以AHP确定部件权重，整栋健康分部件状态占80%、管护占10%、评价占10%；结合部件状态估算寿命，得出风险部位TOP3与推荐维保清单。',
'模型2采用数需评推：烤房健康影响有效烤能Q，再与需求D比对缺口；AHP+熵权组合权重与TOPSIS点位排序，地图沙盘模拟新建数量和位置，支持但不替代审批。',
'双模耦合从存量设备健康评估传导到有效烤能，再作用于是否新建、新建多少与在哪里建，构成存量修复与增量投入协同创新。',
'三智不是3个孤立界面，申请、派单、维修、验收、评价、回写形成模型复算闭环。四队及2项制度确保线上判断有人线下落实。',
'结果分母：48次调剂、18次报修；16条平台工单平均响应0.38小时不等于18次报修平均完工时长；531份有效问卷满意度96.8%；1363/1672约81%，基金近190万元为总额。',
'原稿的12478张影像、6704条评价、150栋修复完成、7栋新建已建成、雷击现场2处故障、旧95.8%及基金余额等未全部与新报送材料一致，不能在上场稿无说明地报成已验收成果。']:
 text(research,z)
research.save(OUT/f'{STAMP}_02_5万字项目全文件研究与技术增订.docx')
official=Document(v5/'10192342_03_V5_官方叙述一万字级研究校订版.docx')
head(official,'V6官方叙述增订与台上数字口径')
for z in [
'采用2026年官方正式报送：2442全量建档、340多栋健康分低于60预警对象、48调剂、18报修、16条工单平均响应0.38h、531份有效问卷满意度96.8%、1363统筹管护／1672在用约81%、基金总额近190万元。',
'严格区分入库栋数、在用栋数、管护栋数与风险预警栋数的统计边界；不得把风险预警说成完成维修，不得把总额说成结余。',
'9.30初稿保留“匠心工坊”“资源无统筹、投入无量化、管护无闭环”及智用、智投、智管和4队分工，技术部分与最新PPT一致。',
'时间矛盾客观存在：2010口述汉字加30秒视频和25秒换页，需接近285字／分钟；另提供1659字上场稿并要求实测。',
'QC稿的独立成果仅用于学习结构与AB协同，不在智烤项目口播。']:text(official,z)
official.save(OUT/f'{STAMP}_03_官方叙述规范1万字研究校订.docx')
# 25-slide mapping plus the explicit time/length conflict.
d=Document();d.add_heading('V6发布稿：25页PPT逐页对应与时间预算',0)
d.add_paragraph('完整版2010汉字，A/B=1411/599；上场版1659汉字，A/B=1148/511。视频30秒、换页25秒，时间必须以完整AB排练核证。')
maps={}
for part in parts(SHORT):
 m=re.search(r'P(\d+)(?:[—-]P?(\d+))?',part['title'])
 if not m:continue
 for i in range(int(m.group(1)),int(m.group(2) or m.group(1))+1):maps[i]=part
assert sorted(maps)==list(range(1,26)),maps.keys()
for i in range(1,26):
 part=maps[i];head(d,f'第{i}页｜{part["title"]}')
 m=re.search(r'P(\d+)',part['title']);start=int(m.group(1))
 if i==2:text(d,'30秒视频静默播放；原视频需裁剪，其他嵌入演示不能额外完整播放。')
 elif i==start:
  for role,t in part['lines']:
   if role!='CUE':text(d,role+'：'+t)
 else:text(d,'与前一页同组解说，使用模型图或界面作支撑，不额外增加重复口播。')
 if i in (8,9):text(d,'技术校核：P8寿命公式与P9 Weibull文字需确认同版。')
 if i==6:text(d,'清查18项与健康模型计算12项应提供明确字段映射。')
head(d,'时间预算')
rows=[]
for label,count in [('完整版',2010),('上场版',1659)]:
 for pace in (220,235,250,270,285):
  sec=count/pace*60+55
  rows.append([label,str(count),str(pace)+'汉字／分',f'{int(sec//60)}分{int(round(sec%60)):02d}秒','需现场验证'])
tab(d,['版本','可朗读汉字','语速假设','含视频与切页','结论'],rows)
text(d,'实际读出AHP、TOPSIS、数字及英文缩写所需时间可能长于汉字估算，完整版并非可靠的8分钟上场版本。')
d.save(OUT/f'{STAMP}_06_25页PPT逐页对应与计时表.docx')
# Facts and 100% demand checklist.
d=Document();d.add_heading('V6重新审稿：9.30原稿对照、事实边界与需求逐条验收',0)
requirements=[
('原稿细节','保留匠心工坊、三愁三缺、4队职责与经核验成果','满足'),
('手机采集流程','手机拍照、GPS、字段校验、4部位18项、正常损坏缺失、档案回写','满足'),
('创新模型1','AHP、80/10/10、剩余寿命、TOP3、维修建议','满足'),
('创新模型2','数需评推、有效烤能Q、理论需求D、熵权、TOPSIS、沙盘','满足'),
('服务与落地','48调剂、18报修、16条0.38h、4队制度和基金','满足'),
('数字阿拉伯形式','指标、人数、栋数、百分数使用阿拉伯数字','满足'),
('A/B约7:3','2010字1411/599；1659字1148/511','满足'),
('2000字','2010个可朗读汉字完整版','满足'),
('8分钟含30秒视频','1659字上场版，待现场计时；完整版应作为资料版','附上场版'),
('不口语化、不自称评委','B自然引出问题，不写“作为评委”','满足'),
('结尾口号','保留1次“聚数成方，一脉致远”','满足'),
('与PPT一致','25页全部关联，P8/P9寿命算法待团队统一','部分待技术确认')]
tab(d,['需求','本轮对应','核验状态'],requirements)
head(d,'9.30原稿中必须补充佐证才能作为已完成成果的内容')
for z in [
'旧稿的12478张图片、6704条评价：新版官方报告未查到匹配的时点和统计口径。',
'旧稿150栋设备改造完成、肖家坪7栋已建成：须有项目任务书和验收台账；暂未写入台上成绩。',
'旧稿磨坪雷击2处故障与单个案例3小时内抢修：需对应工单、工时与配件记录，不能冒充官方样本统计。',
'旧满意度95.8%改为官方531份有效问卷96.8%；旧0.5h改为16条平台工单0.38h。',
'旧基金年底预计结余不能替代正式报告的“基金总额近190万元”。'
]:text(d,z)
head(d,'赛前必须确认的2个技术口径')
text(d,'P6的18项清查与P8的12个计算参数需要映射说明；P8寿命合成公式与P9的Weibull拟合描述需要确认同一现行技术版本。')
d.save(OUT/f'{STAMP}_07_9月30原稿修改对照与需求核验.docx')
(OUT/f'{STAMP}_04_2000字完整版源码.md').write_text(FULL,'utf8')
(OUT/f'{STAMP}_05_8分钟上场版源码.md').write_text(SHORT,'utf8')
(OUT/f'{STAMP}_README.md').write_text(
f'# V6完全重建 2026-10-20\n- 2000字版：A=1411，B=599，共2010口述汉字。\n- 上场版：A=1148，B=511，共1659口述汉字。\n- 30秒视频与25秒切页计入，须实际彩排。\n- 研究02/03在V5逐项研究源资料与图像的基础上增订，不能把5万字全文称为本轮全新原创。\n- 旧V6目录已被清理，未修改最新PPT原件。\n','utf8')
docs=sorted(OUT.glob('*.docx'))
assert len(docs)==8,(len(docs),[f.name for f in docs])
with tarfile.open(OUT/f'{STAMP}_智烤管家V6全套8份Word.tar.gz','w:gz') as tf:
 for file in docs:tf.add(file,arcname=file.name)
 tf.add(OUT/f'{STAMP}_README.md',arcname='README.md')
print('CREATED',len(docs),[(f.name,f.stat().st_size) for f in docs])

# Publish V6 entirely regenerated package 2026-10-20.
