from pathlib import Path
from docx import Document
import subprocess
OUT=Path(".tmp_sources_09300909"); OUT.mkdir(exist_ok=True)
def docx_text(p):
    d=Document(p); out=[]
    for x in d.paragraphs:
        if x.text.strip(): out.append(x.text.strip())
    for t in d.tables:
        for row in t.rows:
            vals=[c.text.replace("\n"," | ").strip() for c in row.cells]
            if any(vals): out.append("\t".join(vals))
    return "\n".join(out)
def old_doc(p):
    cp=subprocess.run(["antiword",str(p)],stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    for enc in ("utf-8","gb18030"):
        try:return cp.stdout.decode(enc)
        except:pass
    return cp.stdout.decode("utf-8","ignore")
files=[
("official.txt","读取资料/3之前项目官方叙述过的/附件2：宜昌市局+智烤管家-基于健康评估的烤房用投管一体化解决方案+数创成果报告.docx","docx"),
("old_draft.txt","读取资料/4一版初稿/讲稿初稿.doc","doc")
]
for name,p,k in files:
    txt=docx_text(p) if k=="docx" else old_doc(p)
    (OUT/name).write_text(txt,encoding="utf-8")
    print(name,len(txt))
