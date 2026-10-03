from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT=Path('.')
DEST=ROOT/'决赛发布材料/1002PPT决赛完善版/1003'
SRC=DEST/'智烤管家_10031957_原图1比1全矢量可编辑重构_15页合并版.pptx'
OUT=DEST/'智烤管家_10032026_全矢量可编辑_亮蓝科技黄标母版精修版.pptx'
prs=Presentation(SRC)

FILL_MAP={
'03162B':'062B49','031427':'062742','061F36':'0A3456','082944':'0D3A5D','0A2E4A':'124B70',
'0A4B70':'12628E','0B4C70':'146C99','0B4B6D':'13648E','0B4668':'125E86','0D4B6C':'156386',
'0B4768':'126087','0A4A70':'13668F','0A2A46':'0E3C60','0A3557':'104A70','0B3557':'114A70',
'0A304E':'0D456A','0D3A59':'114F75','0B405F':'145778','B8E6FF':'D8F7FF','78CFFF':'8DE8FF',
'46D9FF':'63ECFF','80E8FF':'A8F4FF','CFF7FF':'E7FCFF','DDF7FF':'F0FCFF','F8FCFF':'FFFFFF',
'FFD65A':'FFC928','EAF7FF':'F2FBFF','EAFBFF':'F5FDFF','F3F8FC':'F7FCFF'
}
LINE_MAP={
'46D9FF':'63ECFF','0B6A9C':'1682B8','0B6F9D':'1A8FC5','3ABDF3':'50D8FF','59DEFF':'79EEFF',
'0A6F9E':'188FC5','8BEAFF':'B0F5FF','32BBF2':'4DD5FF','407D99':'5E9AB5','0C638F':'157FAE',
'4B7890':'6498AE','0A5E8C':'137AA9','0C6694':'1685B5','54859E':'6E9FAF','1C7FAB':'2B9CC5','236D9C':'3489B5',
'FFD65A':'FFC928'
}
CYAN='63ECFF'; YELLOW='FFC928'; FONT_DEFAULT='F7FCFF'

def rgb(h): return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))

def set_alpha(shape, transparency):
    solid=shape.fill._xPr.find(qn('a:solidFill'))
    if solid is not None and len(solid):
        clr=solid[0]
        for a in list(clr.findall(qn('a:alpha'))): clr.remove(a)
        a=OxmlElement('a:alpha'); a.set('val',str(int((100-transparency)*1000))); clr.append(a)

def set_text_color(sh, color):
    if not getattr(sh,'has_text_frame',False): return
    for p in sh.text_frame.paragraphs:
        for r in p.runs: r.font.color.rgb=rgb(color)

yellow_exact={'01','02','03','04','分不到','修不及','投不准','60分','6.8年','缺口','先调剂','再修复','才新建','该修哪栋','该补哪里','该建多少','少跑腿','少等待','少协调','28%','档案不断档，责任不断链','2442座','1363次','96.8%','≈190万','48个','一库清底','双模决策','三智落地','四队长效','让每一座烤房，'}
yellow_contains=['先调剂、再维修、最后才新建','仍有结构性缺口','把“看不清、修不及、投不准”','健康评分','基金保障']
cyan_exact={'资源调度','抢修服务','投入决策','用房管理断链','维修保障断链','投入决策断链','现场采集','标准判定','一房一档','状态标准化','AHP组合评分','Weibull估寿','算有效烤能','判真实缺口','做建设推演','健康评估','供需推演','双模耦合输出','智用','智投','智管','存量体检','缺口识别','选址推演','效果复盘','建档','巡检','维修','保障','专班推进','维修服务','运营管护'}
slide_tags={2:'01 痛点',3:'02 诊断',4:'03 方案',5:'04 一库',6:'05 模型1',7:'06 模型2',8:'07 双模',9:'08 三智',10:'09 智用',11:'10 智投',12:'11 智管',13:'12 长效',14:'13 成效'}

for si,slide in enumerate(prs.slides,1):
    first_content_el=None
    for sh in slide.shapes:
        try:
            if sh.fill.type is not None and sh.fill.fore_color.type is not None:
                old=str(sh.fill.fore_color.rgb)
                if old in FILL_MAP: sh.fill.fore_color.rgb=rgb(FILL_MAP[old])
        except: pass
        try:
            if sh.line.color.type is not None:
                old=str(sh.line.color.rgb)
                if old in LINE_MAP:
                    sh.line.color.rgb=rgb(LINE_MAP[old])
                    if sh.line.width and sh.line.width < Pt(2.2): sh.line.width=int(sh.line.width*1.10)
        except: pass
        n=sh.name
        if n.startswith('背景_深蓝底_'):
            sh.fill.fore_color.rgb=rgb('062B49')
        elif '背景_山体层_1' in n:
            sh.fill.fore_color.rgb=rgb('124D72'); set_alpha(sh,28)
        elif '背景_山体层_2' in n:
            sh.fill.fore_color.rgb=rgb('0F4165'); set_alpha(sh,18)
        elif '背景_山体层_3' in n:
            sh.fill.fore_color.rgb=rgb('0A3456'); set_alpha(sh,8)
        elif '背景_云雾_' in n:
            sh.fill.fore_color.rgb=rgb('DDF9FF'); set_alpha(sh,80)
        elif '背景_左上光束_' in n:
            sh.fill.fore_color.rgb=rgb('8DEBFF'); set_alpha(sh,73)
        elif '背景_星点_' in n:
            sh.fill.fore_color.rgb=rgb('E8FDFF')
        elif '背景_数据光柱_' in n:
            try: sh.line.color.rgb=rgb('5BE3FF'); sh.line.width=Pt(1.15)
            except: pass
        elif '背景_光路底_' in n:
            try: sh.line.color.rgb=rgb('1577A9')
            except: pass
        elif '背景_光路亮线_' in n:
            try: sh.line.color.rgb=rgb('55E3FF')
            except: pass
        elif '背景_地平线光环' in n:
            try: sh.line.color.rgb=rgb('69E9FF')
            except: pass
        if first_content_el is None and not n.startswith('背景_'): first_content_el=sh._element

        if getattr(sh,'has_text_frame',False) and sh.text.strip():
            text=sh.text.strip()
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try: old=str(r.font.color.rgb)
                    except: old=''
                    if old in FILL_MAP: r.font.color.rgb=rgb(FILL_MAP[old])
                    elif old in ('','None'): r.font.color.rgb=rgb(FONT_DEFAULT)
            if sh.top < Inches(1.45) and sh.width > Inches(5.0) and len(text) > 8:
                set_text_color(sh,'F7FCFF')
            if text in yellow_exact or any(k in text for k in yellow_contains):
                set_text_color(sh,YELLOW)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs: r.font.bold=True
            elif text in cyan_exact:
                set_text_color(sh,CYAN)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs: r.font.bold=True
            if si==14 and any(k in text for k in ['2442','1363','96.8','190','48个']):
                set_text_color(sh,YELLOW)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs: r.font.bold=True

    if si in slide_tags:
        tag=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(0.48),Inches(0.48),Inches(1.02),Inches(0.32))
        tag.name=f'母版黄标_S{si:02d}_可编辑'
        tag.fill.solid(); tag.fill.fore_color.rgb=rgb(YELLOW); tag.line.fill.background()
        tf=tag.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=slide_tags[si]; r.font.name='Arial'; r.font.size=Pt(8.2); r.font.bold=True; r.font.color.rgb=rgb('082944')
        if first_content_el is not None:
            tree=slide.shapes._spTree; el=tag._element; tree.remove(el); tree.insert(tree.index(first_content_el),el)

prs.save(OUT)
print(OUT)
