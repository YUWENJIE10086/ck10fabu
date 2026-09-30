from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
import re

OUT=Path("决赛发布材料/09300004发布稿")
STAMP="09300004"; BF="仿宋_GB2312"; HF="黑体"

KEYS=[
"三大痛点","资源调度缺少数字支撑","基础设施投入缺少量化依据","管护运维缺少闭环工具",
"1234","一库清底、双模决策、三智协同、四队长效","2442","12个关键部件","80%","10%","AHP",
"威布尔分布","维修可以续命，但不能回春","340多栋","有效烤能","50%","30%","20%","熵权法","TOPSIS",
"增长2%、5%、10%","29栋","7栋","智用、智投、智管","不到3小时","0.38小时","96.8%","81%","190万元",
"先调剂、再修复、仍有结构性缺口才新建","好房充分用，旧房精准修，缺口科学建，建后持续管"
]

def rf(run,name=BF,size=15,bold=None):
    run.font.name=name; run.font.size=Pt(size)
    if bold is not None: run.font.bold=bold
    rpr=run._element.get_or_add_rPr(); rfnt=rpr.rFonts
    if rfnt is None: rfnt=OxmlElement("w:rFonts"); rpr.insert(0,rfnt)
    for k in ("ascii","hAnsi","eastAsia","cs"): rfnt.set(qn("w:"+k),name)

def bold_para(p):
    text=p.text
    hits=[x for x in KEYS if x in text]
    if not hits:return
    prefix=""
    if text.startswith("A：") or text.startswith("B："):
        prefix=text[:2]; body=text[2:]
    else: body=text
    for r in p.runs:r._element.getparent().remove(r._element)
    if prefix:
        r=p.add_run(prefix); rf(r,HF,15,True)
    pat="("+"|".join(map(re.escape,sorted([x for x in hits if x in body],key=len,reverse=True)))+")" if any(x in body for x in hits) else None
    if pat:
        for x in re.split(pat,body):
            if not x:continue
            r=p.add_run(x); rf(r,BF,15,x in KEYS)
    else:
        r=p.add_run(body); rf(r,BF,15)

p1=OUT/f"01_智烤管家_湖北省创客大赛决赛双人发布稿_正式参赛版_{STAMP}.docx"
d=Document(p1)
for p in d.paragraphs:bold_para(p)
d.save(p1)

WRITE=[
("开场身份+三问","先介绍A；B定位项目组成员，再说没房、故障、修还是建","优秀稿几乎都让评委30秒内知道人物和场景。"),
("三愁+三痛点+一个根因","先人物痛，再管理痛，再上升为“用投管割裂”","把“为什么做”从功能需求上升到行业管理问题。"),
("1234框架","一库、双模、三智、四队","形成可复述记忆点，且全部来自真实项目构成。"),
("健康模型","先问哪栋先修，再讲输入—健康分—寿命—输出—维修动作","算法不是为了炫技，而是改变维修排序。"),
("发展模型","先问修完还缺怎么办，再讲有效烤能—指数—候选点—沙盘","把新建决策从经验审批变成可量化、可推演。"),
("双模耦合","明确第一模型输出成为第二模型输入","这是区别普通“两个评分模型”的核心。"),
("35秒视频","不重复口播，只展示结果怎样改变业务动作","视频负责证明系统会跑，口播负责解释为什么这样跑。"),
("三智+四队","模型做决策，平台跑业务，队伍和基金保长效","补齐管理创新，不让项目只像算法比赛。"),
("结果","场景案例+官方季末硬数据","数字按“数据—效率—用户—机制”四层验证。"),
("结尾","三大创新→口诀→三组“从…到…”","先理性总结，再情绪收束，给评委明确记忆点。")
]

def shade(c,fill):
    pr=c._tc.get_or_add_tcPr(); e=OxmlElement("w:shd"); e.set(qn("w:fill"),fill); pr.append(e)
def cell(c,t,b=False,sz=9):
    c.text=""; p=c.paragraphs[0]; p.paragraph_format.line_spacing=1.05
    r=p.add_run(str(t)); rf(r,"Noto Sans CJK SC",sz,b); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

p2=OUT/f"02_优秀发布稿逐句拆解与本版写法说明_{STAMP}.docx"
d=Document(p2)
h=d.add_paragraph(style="Heading 1"); r=h.add_run("九、逐段写法复核：为什么每一部分这样写"); rf(r,HF,18,True)
tb=d.add_table(rows=1,cols=3); tb.style="Table Grid"; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,x in enumerate(["段落","新版写法","写作目的"]):cell(tb.rows[0].cells[i],x,True,9);shade(tb.rows[0].cells[i],"D9E2F3")
for row in WRITE:
    cs=tb.add_row().cells
    for i,x in enumerate(row):cell(cs[i],x,False,9)
d.save(p2)
print("patched")
