from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from PIL import Image, ImageEnhance, ImageFilter
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import io, re, math, numpy as np, cv2

BASE=Path('决赛发布材料/1002PPT决赛完善版/1003')
A=BASE/'智烤管家_10031539_1003原图1比1字效原样+文字可编辑下层+图标分层+背景无光路版.pptx'
S3=BASE/'智烤管家_142106_1比1蓝白山海光路抠图还原版_副本3.pptx'
S2=BASE/'智烤管家_142106_1比1蓝白山海光路抠图还原版_副本2.pptx'
STAMP='10031616'
OUT=Path(f'智烤管家_{STAMP}_A版原图视觉保持+文字渐变可编辑+小图标原生组件+亮背景无底部光路版.pptx')
TMP=Path('/tmp')/f'a_refine2_{STAMP}'; TMP.mkdir(exist_ok=True)
prs=Presentation(A); src3=Presentation(S3); src2=Presentation(S2)
sources=list(src3.slides)+list(src2.slides); assert len(sources)==15

A_NS='http://schemas.openxmlformats.org/drawingml/2006/main'
def rm(parent,tag):
    for c in list(parent):
        if c.tag==qn('a:'+tag): parent.remove(c)
def set_grad(run, cols, glow=False, outline=False):
    rPr=run._r.get_or_add_rPr()
    for t in ['solidFill','gradFill','noFill','pattFill','blipFill']: rm(rPr,t)
    if len(cols)==1:
        rPr.append(parse_xml(f'<a:solidFill xmlns:a="{A_NS}"><a:srgbClr val="{cols[0]}"/></a:solidFill>'))
    else:
        stops=''.join(f'<a:gs pos="{p}"><a:srgbClr val="{c}"/></a:gs>' for p,c in zip([0,50000,100000],cols))
        rPr.append(parse_xml(f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>{stops}</a:gsLst><a:lin ang="5400000" scaled="1"/></a:gradFill>'))
    for c in list(rPr):
        if c.tag in [qn('a:glow'),qn('a:ln')]: rPr.remove(c)
    if outline: rPr.append(parse_xml(f'<a:ln xmlns:a="{A_NS}" w="7600"><a:solidFill><a:srgbClr val="E8FCFF"><a:alpha val="50000"/></a:srgbClr></a:solidFill></a:ln>'))
    if glow: rPr.append(parse_xml(f'<a:glow xmlns:a="{A_NS}" rad="18000"><a:srgbClr val="36D8FF"><a:alpha val="42000"/></a:srgbClr></a:glow>'))
def sample(pic):
    im=Image.open(io.BytesIO(pic.image.blob)).convert('RGBA'); a=np.asarray(im); al=a[:,:,3].astype(float); v=a[:,:,:3].max(2)
    m=(al>60)
    if m.sum()<10: return ['F3FAFF']
    thr=np.percentile(v[m],62)
    core=m & (v>=max(thr,135))
    if core.sum()<8: core=m
    rgb=a[:,:,:3][core].astype(float); ww=al[core]
    avg=(rgb*ww[:,None]).sum(0)/max(ww.sum(),1)
    cyan=(avg[2]-avg[0] > 18 and avg[1]-avg[0] > 6)
    if cyan: return ['FAFEFF','C9F4FF','49D8FF']
    return ['F9FCFF']

def remove(sh): el=sh._element; el.getparent().remove(el)
def back(slide,sh): tree=slide.shapes._spTree; el=sh._element; tree.remove(el); tree.insert(2,el)

CYAN='3AD9FF'; ICE='EAFBFF'; NAVY='0A345B'; WHITE='F8FCFF'; BLUE='1C8FD0'
def C(h): return RGBColor(int(h[:2],16),int(h[2:4],16),int(h[4:],16))
def shp(slide,t,x,y,w,h,fill=None,line=CYAN,lw=1.3,name=''):
    s=slide.shapes.add_shape(t,Inches(x),Inches(y),Inches(w),Inches(h)); s.name=name
    if fill: s.fill.solid(); s.fill.fore_color.rgb=C(fill)
    else: s.fill.background()
    if line: s.line.color.rgb=C(line); s.line.width=Pt(lw)
    else: s.line.fill.background()
    return s
def lin(slide,x1,y1,x2,y2,col=ICE,lw=1.5,name=''):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); s.name=name; s.line.color.rgb=C(col); s.line.width=Pt(lw); return s
def tx(slide,t,x,y,w,h,fs=12,name=''):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); s.name=name; tf=s.text_frame; tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=t; r.font.name='Arial'; r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=C(ICE); return s
def ring(slide,x,y,w,h,p):
    shp(slide,SHAPE.OVAL,x-.05,y-.05,w+.1,h+.1,None,'1A8BC6',.8,p+'_外环'); shp(slide,SHAPE.OVAL,x,y,w,h,None,CYAN,1.1,p+'_内环')
def server(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p)
    for i in range(3):
        yy=y+h*(.18+.22*i); shp(slide,SHAPE.ROUNDED_RECTANGLE,x+w*.26,yy,w*.48,h*.15,'0E6790',ICE,.7,p+f'_服务器{i}'); shp(slide,SHAPE.OVAL,x+w*.31,yy+h*.045,w*.045,w*.045,CYAN,None,0,p+f'_灯{i}')
def bars(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p)
    for i,hh in enumerate([.34,.56,.82]): shp(slide,SHAPE.RECTANGLE,x+w*(.23+.19*i),y+h*(.78-hh*.55),w*.12,h*hh*.55,ICE,None,0,p+f'_柱{i}')
    lin(slide,x+w*.20,y+h*.80,x+w*.78,y+h*.80,CYAN,1,p+'_轴')
def house(slide,x,y,w,h,p,pinflag=False):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.RECTANGLE,x+w*.28,y+h*.43,w*.44,h*.32,'0D6791',ICE,.8,p+'_房体'); shp(slide,SHAPE.ISOSCELES_TRIANGLE,x+w*.21,y+h*.18,w*.58,h*.36,ICE,CYAN,.8,p+'_屋顶'); shp(slide,SHAPE.RECTANGLE,x+w*.47,y+h*.58,w*.10,h*.18,NAVY,ICE,.5,p+'_门')
    if pinflag: pin(slide,x+w*.58,y+h*.45,w*.28,h*.35,p+'_定位',False)
def gear(slide,x,y,w,h,p): ring(slide,x,y,w,h,p); shp(slide,SHAPE.GEAR_6,x+w*.25,y+h*.23,w*.50,h*.50,'0D6791',ICE,.8,p+'_齿轮')
def tools(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.GEAR_6,x+w*.52,y+h*.50,w*.24,h*.24,'0D6791',ICE,.7,p+'_齿轮'); lin(slide,x+w*.30,y+h*.70,x+w*.62,y+h*.32,ICE,3,p+'_扳手柄'); shp(slide,SHAPE.OVAL,x+w*.22,y+h*.65,w*.15,w*.15,None,ICE,2,p+'_扳手尾')
def people(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p)
    for i,dx in enumerate([.33,.53,.70]): shp(slide,SHAPE.OVAL,x+w*(dx-.06),y+h*(.24+(i%2)*.04),w*.12,w*.12,ICE,CYAN,.5,p+f'_头{i}'); shp(slide,SHAPE.ROUNDED_RECTANGLE,x+w*(dx-.10),y+h*.45,w*.20,h*.23,'0D6791',ICE,.6,p+f'_身{i}')
def pin(slide,x,y,w,h,p,ringflag=True):
    if ringflag:ring(slide,x,y,w,h,p)
    s=shp(slide,SHAPE.TEAR,x+w*.31,y+h*.18,w*.38,h*.48,ICE,CYAN,.7,p+'_针'); s.rotation=180; shp(slide,SHAPE.OVAL,x+w*.44,y+h*.31,w*.12,w*.12,NAVY,None,0,p+'_孔')
def doc(slide,x,y,w,h,p,mode='plain'):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.ROUNDED_RECTANGLE,x+w*.29,y+h*.18,w*.40,h*.50,'0D6791',ICE,.8,p+'_纸')
    for i in range(3): lin(slide,x+w*.35,y+h*(.30+.10*i),x+w*.56,y+h*(.30+.10*i),ICE,.8,p+f'_线{i}')
    if mode=='search': shp(slide,SHAPE.OVAL,x+w*.56,y+h*.48,w*.19,w*.19,None,ICE,1.7,p+'_镜'); lin(slide,x+w*.69,y+h*.63,x+w*.80,y+h*.75,ICE,1.7,p+'_镜柄')
    if mode=='check': lin(slide,x+w*.58,y+h*.59,x+w*.64,y+h*.65,ICE,1.5,p+'_勾1'); lin(slide,x+w*.64,y+h*.65,x+w*.75,y+h*.52,ICE,1.5,p+'_勾2')
    if mode=='home': house(slide,x+w*.50,y+h*.46,w*.28,h*.28,p+'_小房')
def folder(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.FOLDED_CORNER,x+w*.23,y+h*.32,w*.54,h*.32,'0D6791',ICE,.8,p+'_夹'); lin(slide,x+w*.58,y+h*.57,x+w*.63,y+h*.62,ICE,1.4,p+'_勾1'); lin(slide,x+w*.63,y+h*.62,x+w*.73,y+h*.50,ICE,1.4,p+'_勾2')
def search(slide,x,y,w,h,p,chart=False):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.OVAL,x+w*.28,y+h*.22,w*.33,w*.33,None,ICE,1.8,p+'_镜'); lin(slide,x+w*.54,y+h*.50,x+w*.70,y+h*.68,ICE,2,p+'_柄')
    if chart:
        for i,hh in enumerate([.15,.25,.36]): shp(slide,SHAPE.RECTANGLE,x+w*(.34+.07*i),y+h*(.45-hh),w*.04,h*hh,ICE,None,0,p+f'_柱{i}')
def crane(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p); lin(slide,x+w*.45,y+h*.20,x+w*.45,y+h*.72,ICE,1.8,p+'_塔'); lin(slide,x+w*.25,y+h*.28,x+w*.72,y+h*.28,ICE,1.8,p+'_臂'); lin(slide,x+w*.62,y+h*.28,x+w*.62,y+h*.53,ICE,1,p+'_线'); shp(slide,SHAPE.RECTANGLE,x+w*.58,y+h*.53,w*.09,h*.10,ICE,None,0,p+'_吊')
def clock(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.OVAL,x+w*.28,y+h*.22,w*.42,h*.42,None,ICE,1.7,p+'_表'); lin(slide,x+w*.49,y+h*.28,x+w*.49,y+h*.45,ICE,1.6,p+'_时'); lin(slide,x+w*.49,y+h*.45,x+w*.62,y+h*.52,ICE,1.6,p+'_分')
def alert(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.ISOSCELES_TRIANGLE,x+w*.24,y+h*.18,w*.52,h*.48,None,ICE,1.5,p+'_三角'); tx(slide,'!',x+w*.40,y+h*.32,w*.20,h*.20,11,p+'_叹')
def shield(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p); shp(slide,SHAPE.PENTAGON,x+w*.28,y+h*.20,w*.44,h*.44,'0D6791',ICE,.8,p+'_盾'); tx(slide,'¥',x+w*.36,y+h*.28,w*.28,h*.22,11,p+'_元')
def wave(slide,x,y,w,h,p):
    ring(slide,x,y,w,h,p); pts=[(.26,.52),(.36,.52),(.42,.34),(.48,.65),(.55,.43),(.62,.52),(.75,.52)]
    for a,b in zip(pts,pts[1:]): lin(slide,x+w*a[0],y+h*a[1],x+w*b[0],y+h*b[1],ICE,1.6,p+'_波')
def trend(slide,x,y,w,h,p):
    bars(slide,x,y,w,h,p); lin(slide,x+w*.28,y+h*.62,x+w*.68,y+h*.27,ICE,1.6,p+'_趋势')
def eye(slide,x,y,w,h,p): ring(slide,x,y,w,h,p); shp(slide,SHAPE.OVAL,x+w*.26,y+h*.36,w*.48,h*.25,None,ICE,1.4,p+'_眼'); shp(slide,SHAPE.OVAL,x+w*.45,y+h*.41,w*.10,w*.10,ICE,None,0,p+'_瞳')
def send(slide,x,y,w,h,p): ring(slide,x,y,w,h,p); t=shp(slide,SHAPE.ISOSCELES_TRIANGLE,x+w*.30,y+h*.27,w*.42,h*.42,None,ICE,1.5,p+'_纸飞机'); t.rotation=45
def monitor(slide,x,y,w,h,p): ring(slide,x,y,w,h,p); shp(slide,SHAPE.RECTANGLE,x+w*.27,y+h*.25,w*.46,h*.33,'0D6791',ICE,.8,p+'_屏'); lin(slide,x+w*.34,y+h*.49,x+w*.45,y+h*.41,ICE,1,p+'_线1'); lin(slide,x+w*.45,y+h*.41,x+w*.55,y+h*.44,ICE,1,p+'_线2'); lin(slide,x+w*.55,y+h*.44,x+w*.66,y+h*.34,ICE,1,p+'_线3')
def pie(slide,x,y,w,h,p): ring(slide,x,y,w,h,p); shp(slide,SHAPE.PIE,x+w*.29,y+h*.23,w*.42,h*.42,ICE,CYAN,.8,p+'_饼')
FUN={'server':server,'bars':bars,'house':house,'housepin':lambda s,x,y,w,h,p:house(s,x,y,w,h,p,True),'gear':gear,'tools':tools,'people':people,'pin':pin,'doc':lambda *a:doc(*a,mode='plain'),'doc_search':lambda *a:doc(*a,mode='search'),'doc_check':lambda *a:doc(*a,mode='check'),'doc_home':lambda *a:doc(*a,mode='home'),'folder':folder,'search':search,'search_chart':lambda s,x,y,w,h,p:search(s,x,y,w,h,p,True),'crane':crane,'clock':clock,'alert':alert,'shield':shield,'wave':wave,'trend':trend,'eye':eye,'send':send,'monitor':monitor,'pie':pie}
IC={2:[(8.42,2.10,.54,.52,'server'),(8.40,3.62,.58,.54,'tools'),(8.41,5.03,.58,.54,'bars')],3:[(5.72,2.39,.82,.78,'house'),(4.77,4.47,.84,.78,'tools'),(8.25,4.48,.76,.78,'bars')],4:[(1.72,4.82,.62,.60,'house'),(5.00,4.22,.66,.62,'bars'),(8.08,3.92,.70,.66,'gear'),(11.08,3.34,.74,.66,'people')],5:[(2.02,2.15,.70,.68,'doc'),(6.25,2.14,.70,.68,'doc_search'),(10.30,2.12,.80,.70,'folder')],6:[(2.10,1.65,.72,.68,'doc_search'),(6.43,1.62,.70,.66,'gear'),(10.95,1.60,.78,.68,'trend')],7:[(2.66,1.78,.70,.66,'gear'),(6.87,1.78,.72,.66,'search_chart'),(11.08,1.72,.80,.70,'crane'),(11.22,4.46,.56,.56,'tools')],8:[(1.34,3.28,.54,.54,'gear'),(2.96,3.30,.54,.54,'trend'),(8.74,3.30,.54,.54,'wave'),(10.24,3.30,.54,.54,'bars'),(11.86,3.28,.54,.54,'house')],9:[(1.43,4.48,.46,.46,'eye'),(1.43,5.00,.46,.46,'send'),(5.83,4.48,.46,.46,'search'),(5.83,5.00,.46,.46,'pie'),(10.33,4.48,.46,.46,'monitor'),(10.33,5.00,.46,.46,'tools')],10:[(2.10,1.84,.66,.64,'housepin'),(8.55,1.84,.68,.64,'tools'),(11.72,1.84,.66,.64,'doc_check')],11:[(.84,1.99,.58,.58,'wave'),(4.89,1.99,.58,.58,'bars'),(8.04,1.99,.58,.58,'pin'),(11.14,1.99,.58,.58,'trend')],12:[(3.43,2.13,.74,.70,'doc_home'),(7.70,2.13,.74,.70,'doc_search'),(8.02,4.82,.74,.68,'tools'),(4.40,5.54,.60,.60,'shield')],13:[(1.96,3.45,.70,.68,'people'),(4.84,3.45,.70,.68,'tools'),(7.86,3.45,.70,.68,'monitor'),(10.86,3.45,.70,.68,'shield')]}

for sno,(slide,srcslide) in enumerate(zip(prs.slides,sources),1):
    text_boxes=[sh for sh in slide.shapes if sh.name.startswith('可编辑文字下层_')]
    style_pics={sh.name.replace('原字效切片_','').split('_删除后')[0]:sh for sh in slide.shapes if sh.name.startswith('原字效切片_')}
    full=Image.open(io.BytesIO(srcslide.shapes[0].image.blob)).convert('RGB'); arr=np.array(full); Hpx,Wpx=arr.shape[:2]; sx=Wpx/(prs.slide_width/914400); sy=Hpx/(prs.slide_height/914400)
    mask=np.zeros((Hpx,Wpx),np.uint8)
    for tb in text_boxes:
        x=int(tb.left/914400*sx); y=int(tb.top/914400*sy); w=int(tb.width/914400*sx); h=int(tb.height/914400*sy); pad=max(5,int(0.015*Wpx)); cv2.rectangle(mask,(max(0,x-pad),max(0,y-pad//2)),(min(Wpx-1,x+w+pad),min(Hpx-1,y+h+pad//2)),255,-1)
    for x,y,w,h,kind in IC.get(sno,[]):
        x0=int(x*sx); y0=int(y*sy); x1=int((x+w)*sx); y1=int((y+h)*sy); pad=5; cv2.rectangle(mask,(max(0,x0-pad),max(0,y0-pad)),(min(Wpx-1,x1+pad),min(Hpx-1,y1+pad)),255,-1)
    y0=int(5.75*sy); lower=arr[y0:]; hsv=cv2.cvtColor(lower,cv2.COLOR_RGB2HSV); bright=((hsv[:,:,2]>145)&(hsv[:,:,1]>45))|((hsv[:,:,2]>195)&(hsv[:,:,1]<70)); road=np.zeros_like(mask); road[y0:][bright]=255; road=cv2.dilate(road,np.ones((5,13),np.uint8),iterations=1); mask=np.maximum(mask,road)
    textmask=np.zeros((Hpx,Wpx),np.uint8)
    for tb in text_boxes:
        x=int(tb.left/914400*sx); y=int(tb.top/914400*sy); w=int(tb.width/914400*sx); h=int(tb.height/914400*sy); pad=max(5,int(0.010*Wpx)); cv2.rectangle(textmask,(max(0,x-pad),max(0,y-pad//2)),(min(Wpx-1,x+w+pad),min(Hpx-1,y+h+pad//2)),255,-1)
    nontext=cv2.subtract(mask,textmask)
    bgr=cv2.cvtColor(arr,cv2.COLOR_RGB2BGR); cleaned=cv2.inpaint(bgr,nontext,6,cv2.INPAINT_TELEA); cleaned=cv2.cvtColor(cleaned,cv2.COLOR_BGR2RGB)
    blur=cv2.GaussianBlur(cleaned,(0,0),sigmaX=22,sigmaY=12); feather=cv2.GaussianBlur(textmask,(0,0),sigmaX=8,sigmaY=5).astype(np.float32)/255.0; feather=feather[...,None]
    cleaned=(cleaned*(1-feather)+blur*feather).clip(0,255).astype(np.uint8)
    im=Image.fromarray(cleaned); im=ImageEnhance.Brightness(im).enhance(1.08); im=ImageEnhance.Color(im).enhance(1.04); im=ImageEnhance.Contrast(im).enhance(1.04)
    bgp=TMP/f'slide_{sno:02d}_bg_bright_no_text_icon_road.jpg'; im.save(bgp,quality=94)
    for sh in list(slide.shapes):
        if sh.shape_type==13: remove(sh)
    bg=slide.shapes.add_picture(str(bgp),0,0,width=prs.slide_width,height=prs.slide_height); bg.name=f'背景_S{sno:02d}_原图视觉保留_稍提亮_无文字_无小图标_无底部光路'; back(slide,bg)
    for tb in text_boxes:
        key=tb.name.replace('可编辑文字下层_','').split('_双击编辑')[0]; pic=style_pics.get(key); cols=sample(pic) if pic else ['F4FBFF']
        try: tb._element.xpath('.//p:cNvPr')[0].attrib.pop('hidden',None)
        except: pass
        txt=''.join(r.text for p in tb.text_frame.paragraphs for r in p.runs); hi=tb.height/914400
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                if re.fullmatch(r'0[1-4]',txt.strip()): r.font.name='Times New Roman'; r.font.bold=True; r.font.italic=True
                elif txt.strip().startswith('Z H I'): r.font.name='Arial Narrow'; r.font.bold=False
                elif hi>=.55 or len(txt.strip())<=8: r.font.name='华文中宋'; r.font.bold=True
                else: r.font.name='微软雅黑'
                set_grad(r,cols,glow=hi>.65,outline=hi>.85)
        tb.name=tb.name.replace('可编辑文字下层_','可编辑文字_').replace('_双击编辑_上方保留原字效切片','_直接双击修改_渐变颜色按原图采样')
    for i,(x,y,w,h,k) in enumerate(IC.get(sno,[]),1): FUN[k](slide,x,y,w,h,f'可编辑小图标_S{sno:02d}_{i:02d}_{k}')

prs.save(OUT)
print(OUT)
