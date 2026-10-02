de11 智投
s=prs.slides[10]; remove_all_shapes(s); add_hud_header(s,'智投增效：每一笔投入，先过“双模”','管理人员从“凭经验申报”转向“有清单、有缺口、有沙盘”',section=3,tabs=[('智投',True),('智用',False),('智管',False)])
add_panel(s,.48,1.25,3.75,3.95,glow=True,fill1='073F5A',fill2='01182E',alpha1=65000,alpha2=43000)
for i,(t,d) in enumerate([('看健康','模型①筛风险烤房｜生成巡检/修复优先序'),('看缺口','模型②核算有效烤能｜判断村级真实供需缺口'),('看沙盘','候选点排序 + 情景推演｜比较建前/建后改善幅度')]):
    add_number_badge(s,i+1,t,.72,1.63+i*1.0)
    add_text(s,d,2.07,1.58+i*1.0,1.9,.58,9,SOFT,False)
add_text(s,'管理动作',.74,4.55,.9,.24,9,YELLOW,True); add_text(s,'巡检清单 · 修复清单 · 新建清单 · 资金优先序',1.72,4.48,2.2,.42,9,WHITE,True)
# evidence right
add_panel(s,4.55,1.25,4.95,3.95,glow=True,fill1='053B58',fill2='01162B',alpha1=62000,alpha2=40000)
image_frame(s,AS/'invest_strip.jpg',4.8,1.55,4.45,2.0,'INVEST')
add_grad_text(s,'29',5.02,3.82,.9,.48,25,True,PP_ALIGN.CENTER,YELLOW,ORANGE); add_text(s,'栋建设方案论证',5.85,3.86,1.55,.32,10,SOFT,True)
add_grad_text(s,'7',7.45,3.82,.7,.48,25,True,PP_ALIGN.CENTER,YELLOW,ORANGE); add_text(s,'栋落在肖家坪村',8.08,3.86,1.1,.32,10,SOFT,True)
add_text(s,'先有缺口证据，再谈建设；先看投入效果，再定位置。',5.0,4.52,4.1,.34,11,WHITE,True,PP_ALIGN.CENTER)
add_fade_transition(s)

# Slide12 智用
s=prs.slides[11]; remove_all_shapes(s); add_hud_header(s,'智用解忧：找得到 · 约得到 · 修得快','烟农 + 技术员：共享调剂与抢修服务一条线跑通',section=3,tabs=[('智投',False),('智用',True),('智管',False)])
add_panel(s,.45,1.25,3.65,3.95,glow=True,fill1='073E59',fill2='01172E',alpha1=65000,alpha2=43000)
flows=[('1','查可用','附近烤房 / 空闲时段 / 状态评分'),('2','预约调剂','线上申请 / 就近匹配'),('3','一键报修','匹配人员 / 配件 / 风险部位'),('4','双向评价','服务评价回写档案')]
for i,(n,t,d) in enumerate(flows):
    y=1.52+i*.82; add_number_badge(s,n,t,.68,y); add_text(s,d,2.02,y-.02,1.82,.42,9,SOFT,False)
add_panel(s,4.45,1.25,5.05,3.95,glow=True,fill1='053B58',fill2='01162B',alpha1=62000,alpha2=40000)
image_frame(s,AS/'phone1.jpg',5.15,1.55,1.55,2.95,'调剂')
image_frame(s,AS/'phone2.jpg',7.2,1.55,1.55,2.95,'报修')
add_text(s,'磨坪实战',4.75,4.72,1.0,.28,10,YELLOW,True); add_text(s,'雷击故障 → 线上报修 → 协同抢修 → 3小时内恢复',5.75,4.68,3.25,.34,10,WHITE,True)
add_fade_transition(s)

# Slide13 智管 + four teams
s=prs.slides[12]; remove_all_shapes(s); add_hud_header(s,'智管提质：线上闭环 + 四队长效','技术负责“算得准”，机制负责“跑得久”',section=3,tabs=[('智投',False),('智用',False),('智管',True)])
# process beam
add_text(s,'线上闭环',.6,1.35,1.0,.3,11,YELLOW,True)
proc=['申请','派单','处置','验收','评价','回写']
for i,t in enumerate(proc):
    x=1.55+i*1.22
    c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(1.28),Inches(.72),Inches(.72)); c.fill.solid();c.fill.fore_color.rgb=rgb('07435A');set_alpha_on_solid(c,76000);set_gradient_line(c,CYAN2,'1376FF',1.2);add_glow(c,CYAN,5,48)
    add_text(s,t,x,1.3,.72,.68,9,WHITE,True,PP_ALIGN.CENTER)
    if i<5:add_arrow(s,x+.75,1.52,x+1.05,1.78,CYAN)
# teams panels
add_text(s,'四支专属管家队伍',.6,2.4,1.65,.32,11,YELLOW,True)
teams=[('烟叶调制','烘烤指导'),('紧急抢修','故障处置'),('日常管护','巡检养护'),('综合管理','监督统筹')]
for i,(t,d) in enumerate(teams):
    x=.6+i*2.3
    add_panel(s,x,2.85,2.05,1.5,glow=True,fill1='07415B',fill2='01172E',alpha1=65000,alpha2=43000)
    add_grad_text(s,t,x+.15,3.12,1.75,.34,14,True,PP_ALIGN.CENTER,WHITE,CYAN2)
    add_text(s,d,x+.2,3.62,1.65,.25,9,SOFT,False,PP_ALIGN.CENTER)
    add_text(s,['技术员','抢修队','管护员','管理员'][i],x+.55,4.02,.95,.22,8,GREEN,True,PP_ALIGN.CENTER)
add_panel(s,1.15,4.65,7.7,.5,glow=True,fill1='0B4B53',fill2='051D32',alpha1=70000,alpha2=52000,line1=GREEN,line2=CYAN2)
add_text(s,'1363栋统筹管护   ·   在用烤房覆盖81%   ·   管护基金累计近190万元',1.28,4.72,7.45,.34,12,WHITE,True,PP_ALIGN.CENTER)
add_fade_transition(s)

# Slide14 outcomes - story + metrics
s=prs.slides[13]; remove_all_shapes(s); add_hud_header(s,'应用成效：三类人的工作方式变了','成效不是“多了几个数字”，而是经验管理 → 数据决策，被动补救 → 主动预判')
cols=[('烟农 / 技术员','用得顺 · 修得快',[('48','次调剂'),('18','次抢修均<3h'),('96.8%','综合满意度')],['找房靠问 → 在线匹配','坏了再等 → 工单协同'],CYAN2),('烟站 / 管理人员','看得清 · 投得准',[('2442','座健康评价'),('340+','预警/危险烤房'),('29','栋方案论证')],['凭经验排查 → 风险排序','凭感觉申报 → 数据推演'],YELLOW),('合作社','管得住 · 跑得久',[('1363','栋统筹管护'),('81%','在用烤房覆盖'),('≈190万','管护基金')],['责任悬空 → 四队履职','一次投入 → 基金滚动'],GREEN)]
for i,(role,val,metrics,changes,col) in enumerate(cols):
    x=.35+i*3.22
    add_panel(s,x,1.2,2.98,3.98,glow=True,fill1='073E59',fill2='01172E',alpha1=65000,alpha2=43000,line1=col,line2=CYAN2)
    add_grad_text(s,role,x+.18,1.45,1.68,.32,14,True,PP_ALIGN.LEFT,WHITE,CYAN2); add_text(s,val,x+1.82,1.47,.95,.28,9,col,True,PP_ALIGN.RIGHT)
    for j,(num,lab) in enumerate(metrics):
        y=2.02+j*.7
        add_grad_text(s,num,x+.18,y,.95,.4,19,True,PP_ALIGN.LEFT,col,WHITE)
        add_text(s,lab,x+1.18,y+.04,1.55,.28,9,SOFT,True)
    add_panel(s,x+.18,4.25,2.62,.65,glow=False,fill1='0B4054',fill2='031B30',alpha1=70000,alpha2=52000)
    add_text(s,changes[0]+'\n'+changes[1],x+.28,4.3,2.42,.5,9,WHITE,True,PP_ALIGN.CENTER)
add_fade_transition(s)

# Slide15 closing
s=prs.slides[14]; remove_all_shapes(s)
# central aura
for rad,alp in [(7.8,11000),(5.7,15000),(3.6,22000)]:
    e=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(5-rad/2),Inches(2.55-rad*.22),Inches(rad),Inches(rad*.44)); e.fill.background();set_gradient_line(e,CYAN2,'116AFF',1.0);add_glow(e,CYAN,8,25)
add_grad_text(s,'智烤管家',2.15,.95,5.7,.7,30,True,PP_ALIGN.CENTER,WHITE,CYAN2)
add_text(s,'ZHI KAO MANAGER',3.55,1.62,2.9,.24,10,YELLOW,True,PP_ALIGN.CENTER)
add_grad_text(s,'好房充分用  ·  旧房精准修  ·  缺口科学建  ·  建后持续管',1.15,2.45,7.7,.45,18,True,PP_ALIGN.CENTER,WHITE,CYAN2)
add_panel(s,2.1,3.25,5.8,.52,glow=True,fill1='0B4659',fill2='061D32',alpha1=68000,alpha2=50000)
add_text(s,'一库清底  ·  双模决策  ·  三智落地  ·  四队长效',2.2,3.33,5.6,.34,12,YELLOW,True,PP_ALIGN.CENTER)
add_text(s,'从一座烤房开始，把烟区基础设施变成可体检、可推演、可服务、可持续的“数字管家”',1.25,4.12,7.5,.4,11,SOFT,False,PP_ALIGN.CENTER)
add_grad_text(s,'谢谢大家',3.45,4.8,3.1,.42,18,True,PP_ALIGN.CENTER,WHITE,CYAN2)
add_fade_transition(s)

prs.save(OUT)
print(OUT)


# ===== PREMIUM POLISH STAGE =====
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFilter
from pathlib import Path
import math, os, re, datetime

SRC = Path('/tmp/智烤管家_PREMIUM_BASE.pptx')
OUT_DIR = Path('决赛发布材料/1002PPT决赛完善版')
ASSET = OUT_DIR / '_assets'

stamp='10021140'
OUT = OUT_DIR / '智烤管家_10021140_高级科技动态一等奖精修版.stage1.pptx'
GIF = Path('/tmp/_动态扫光边框_10021140.gif')

# ---------- animated sweep frame ----------
def make_sweep_gif(path, W=1280, H