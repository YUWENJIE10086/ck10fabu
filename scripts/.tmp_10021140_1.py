from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from copy import deepcopy
from pathlib import Path
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.enum.dml import MSO_LINE_DASH_STYLE

SRC=Path('决赛发布材料/1002PPT决赛完善版/智烤管家_10021021_演讲稿主线母版科技完善版.pptx')
AS=Path('决赛发布材料/1002PPT决赛完善版/_assets')
OUT=Path('/tmp/智烤管家_PREMIUM_BASE.pptx')
prs=Presentation(SRC)
SW,SH=prs.slide_width,prs.slide_height

# palette
NAVY='041E3B'; NAVY2='062A50'; CYAN='00D9FF'; CYAN2='52EEFF'; WHITE='F7FDFF'; SOFT='BCEEFF'; YELLOW='FFC928'; ORANGE='FF9E2C'; RED='FF5A5A'; GREEN='38E6A5'; MID='0B5B8D'; BLUE='129BFF'; DEEP='001528'; MUTED='83BCD3'
FONT='Microsoft YaHei'
FONT_B='Microsoft YaHei'


def rgb(h): return RGBColor.from_string(h)

def remove_all_shapes(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)

def set_alpha_on_solid(shape, alpha=42000):
    solid=shape.fill._xPr.solidFill
    if solid is not None and solid.srgbClr is not None:
        for c in list(solid.srgbClr):
            if c.tag==qn('a:alpha'): solid.srgbClr.remove(c)
        a=parse_xml(f'<a:alpha {nsdecls("a")} val="{alpha}"/>'); solid.srgbClr.append(a)

def add_glow(shape, color=CYAN, radius_pt=6, alpha=55):
    spPr=shape._element.spPr
    eff=spPr.find(qn('a:effectLst'))
    if eff is None:
        eff=parse_xml(f'<a:effectLst {nsdecls("a")}/>' )
        spPr.append(eff)
    glow=eff.find(qn('a:glow'))
    if glow is not None: eff.remove(glow)
    rad=int(radius_pt*12700)
    glow=parse_xml(f'<a:glow {nsdecls("a")} rad="{rad}"><a:srgbClr val="{color}"><a:alpha val="{int(alpha*1000)}"/></a:srgbClr></a:glow>')
    eff.append(glow)

def set_gradient_line(shape, c1=CYAN2, c2='1370FF', width_pt=1.5, alpha1=100000, alpha2=55000):
    ln=shape._element.spPr.find(qn('a:ln'))
    if ln is None:
        ln=parse_xml(f'<a:ln {nsdecls("a")} w="{int(width_pt*12700)}"/>')
        shape._element.spPr.append(ln)
    else: ln.set('w',str(int(width_pt*12700)))
    for tag in [qn('a:solidFill'),qn('a:gradFill'),qn('a:noFill')]:
        n=ln.find(tag)
        if n is not None: ln.remove(n)
    grad=parse_xml(f'''<a:gradFill {nsdecls('a')} rotWithShape="1"><a:gsLst>
    <a:gs pos="0"><a:srgbClr val="{c1}"><a:alpha val="{alpha1}"/></a:srgbClr></a:gs>
    <a:gs pos="100000"><a:srgbClr val="{c2}"><a:alpha val="{alpha2}"/></a:srgbClr></a:gs>
    </a:gsLst><a:lin ang="0" scaled="1"/></a:gradFill>''')
    ln.insert(0,grad)

def set_gradient_fill(shape,c1='0A4366',c2='031C35',alpha1=70000,alpha2=42000,angle=2700000):
    spPr=shape._element.spPr
    for tag in [qn('a:solidFill'),qn('a:gradFill'),qn('a:noFill')]:
        n=spPr.find(tag)
        if n is not None: spPr.remove(n)
    grad=parse_xml(f'''<a:gradFill {nsdecls('a')} rotWithShape="1"><a:gsLst>
    <a:gs pos="0"><a:srgbClr val="{c1}"><a:alpha val="{alpha1}"/></a:srgbClr></a:gs>
    <a:gs pos="100000"><a:srgbClr val="{c2}"><a:alpha val="{alpha2}"/></a:srgbClr></a:gs>
    </a:gsLst><a:lin ang="{angle}" scaled="1"/></a:gradFill>''')
    # after xfrm before ln ideally
    idx=1
    spPr.insert(idx,grad)

def set_text_gradient(run,c1=WHITE,c2=CYAN2):
    rPr=run._r.get_or_add_rPr()
    for tag in [qn('a:solidFill'),qn('a:gradFill')]:
        n=rPr.find(tag)
        if n is not None: rPr.remove(n)
    grad=parse_xml(f'''<a:gradFill {nsdecls('a')} rotWithShape="1"><a:gsLst>
      <a:gs pos="0"><a:srgbClr val="{c1}"/></a:gs>
      <a:gs pos="100000"><a:srgbClr val="{c2}"/></a:gs>
      </a:gsLst><a:lin ang="0" scaled="1"/></a:gradFill>''')
    rPr.append(grad)

def add_fade_transition(slide):
    sld=slide._element
    old=sld.find(qn('p:transition'))
    if old is not None: sld.remove(old)
    trans=parse_xml(f'<p:transition {nsdecls("p")} spd="fast" advClick="1"><p:fade/></p:transition>')
    # insert before timing/extLst, after clrMapOvr if present
    idx=len(sld)
    for i,ch in enumerate(sld):
        if ch.tag in [qn('p:timing'),qn('p:extLst')]: idx=i; break
    sld.insert(idx,trans)

def add_text(slide,text,x,y,w,h,size=18,color=WHITE,bold=False,align=PP_ALIGN.LEFT,font=FONT,anchor=MSO_ANCHOR.MIDDLE,margin=0.03):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(margin)
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return tb

def add_grad_text(slide,text,x,y,w,h,size=30,bold=True,align=PP_ALIGN.LEFT,c1=WHITE,c2=CYAN2):
    tb=add_text(slide,'',x,y,w,h,size=size,color=WHITE,bold=bold,align=align)
    p=tb.text_frame.paragraphs[0]; r=p.add_run(); r.text=text; r.font.name=FONT_B; r.font.size=Pt(size); r.font.bold=bold
    set_text_gradient(r,c1,c2)
    return tb

def add_panel(slide,x,y,w,h, radius='round', glow=True, line1=CYAN2,line2='126EFF', fill1='0B4160',fill2='02182F', alpha1=64000,alpha2=46000):
    shape_type=MSO_SHAPE.ROUNDED_RECTANGLE if radius=='round' else MSO_SHAPE.RECTANGLE
    sh=slide.shapes.add_shape(shape_type,Inches(x),Inches(y),Inches(w),Inches(h))
    set_gradient_fill(sh,fill1,fill2,alpha1,alpha2)
    set_gradient_line(sh,line1,line2,1.4)
    if glow: add_glow(sh,CYAN,5,48)
    return sh

def add_hud_header(slide,title,subtitle='',section=None,tabs=None):
    # upper beam
    line=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(.67), Inches(8.7), Inches(.012))
    line.fill.solid(); line.fill.fore_color.rgb=rgb(CYAN); set_alpha_on_solid(line,45000); line.line.fill.background()
    add_glow(line,CYAN,5,55)
    add_text(slide,'///',.62,.12,.5,.2,10,CYAN,bold=True)
    add_text(slide,'///',8.88,.12,.5,.2,10,CYAN,bold=True,align=PP_ALIGN.RIGHT)
    if section:
        add_grad_text(slide,str(section),2.12,.07,.45,.42,26,True,PP_ALIGN.CENTER,YELLOW,ORANGE)
        tx=2.62; ww=4.9
    else: tx=1.7; ww=6.6
    add_grad_text(slide,title,tx,.04,ww,.5,23,True,PP_ALIGN.CENTER,WHITE,CYAN2)
    if subtitle: add_text(slide,subtitle,2.1,.5,5.8,.22,9,SOFT,False,PP_ALIGN.CENTER)
    if tabs:
        x=7.05
        for label,active in tabs:
            p=add_panel(slide,x,.06,.62,.24,glow=active,fill1=('087796' if active else '08283F'),fill2=('0A4366' if active else '031C2B'),alpha1=70000,alpha2=