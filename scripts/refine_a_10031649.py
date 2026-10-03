from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from PIL import Image, ImageEnhance, ImageDraw
from copy import deepcopy
from collections import defaultdict
import io,re

BASE=Path('决赛发布材料/1002PPT决赛完善版/1003')
SRC=BASE/'智烤管家_10031539_1003原图1比1字效原样+文字可编辑下层+图标分层+背景无光路版.pptx'
ICONS=BASE/'智烤管家_10031616_A版原图视觉保持+文字渐变可编辑+小图标原生组件+亮背景无底部光路版.pptx'
OUT=BASE/'智烤管家_10031649_A版精修终稿_文字渐变可编辑_小图标原生矢量_统一亮背景无光路.pptx'
TMP=Path('/tmp/zhikao_10031649'); TMP.mkdir(parents=True,exist_ok=True)
EMU=914400

prs=Presentation(SRC); iprs=Presentation(ICONS)
# Extract clean background from A, brighten slightly; no bottom road.
bgshape=next(sh for sh in prs.slides[0].shapes if sh.shape_type==13 and sh.name.startswith('背景_'))
bg=Image.open(io.BytesIO(bgshape.image.blob)).convert('RGB')
bg=ImageEnhance.Brightness(bg).enhance(1.11); bg=ImageEnhance.Contrast(bg).enhance(1.02)
ov=Image.new('RGBA',bg.size,(0,0,0,0)); d=ImageDraw.Draw(ov); W,H=bg.size
for rr in range(460,60,-20):
    a=max(0,int(10*(rr/460)**2)); d.ellipse((W//2-rr*2,int(H*.66)-rr//4,W//2+rr*2,int(H*.66)+rr//4),fill=(80,190,255,a))
bg=Image.alpha_composite(bg.convert('RGBA'),ov).convert('RGB'); bgp=TMP/'bg.jpg'; bg.save(bgp,quality=96)

def rm(sh):
    e=sh._element; e.getparent().remove(e)
def back(slide,sh):
    t=slide.shapes._spTree; e=sh._element; t.remove(e); t.insert(2,e)
def setea(run,name):
    run.font.name=name; rp=run._r.get_or_add_rPr(); ea=rp.find(qn('a:ea'))
    if ea is None: ea=OxmlElement('a:ea'); rp.append(ea)
    ea.set('typeface',name)
def clrfill(rp):
    for tag in ('a:solidFill','a:gradFill','a:noFill','a:pattFill'):
        for e in list(rp.findall(qn(tag))): rp.remove(e)
def solid(run,col):
    rp=run._r.get_or_add_rPr(); clrfill(rp); sf=OxmlElement('a:solidFill'); c=OxmlElement('a:srgbClr'); c.set('val',col); sf.append(c); rp.insert(0,sf)
def grad(run,top='FFFFFF',mid='D9FBFF',bottom='2FD4FF',outline='0B365D'):
    rp=run._r.get_or_add_rPr(); clrfill(rp); gf=OxmlElement('a:gradFill'); gf.set('rotWithShape','1'); gsl=OxmlElement('a:gsLst')
    for pos,col in ((0,top),(52000,mid),(100000,bottom)):
        gs=OxmlElement('a:gs'); gs.set('pos',str(pos)); c=OxmlElement('a:srgbClr'); c.set('val',col); gs.append(c); gsl.append(gs)
    gf.append(gsl); lin=OxmlElement('a:lin'); lin.set('ang','5400000'); lin.set('scaled','1'); gf.append(lin); rp.insert(0,gf)
    for e in list(rp.findall(qn('a:ln'))): rp.remove(e)
    ln=OxmlElement('a:ln'); ln.set('w','9000'); sf=OxmlElement('a:solidFill'); c=OxmlElement('a:srgbClr'); c.set('val',outline); sf.append(c); ln.append(sf); rp.append(ln)
def fit(txt,w,h,base,mn,mx):
    n=max(1,len(txt.replace(' ',''))); return max(mn,min(mx,base,w*72/n*.92,h*72*.66))
def style(sh,sno):
    txt=sh.text.strip();
    if not txt:return
    tf=sh.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; r=p.runs[0] if p.runs else p.add_run(); r.text=txt
    x,y,w,h=sh.left/EMU,sh.top/EMU,sh.width/EMU,sh.height/EMU; cp=txt.replace(' ','')
    seq=txt in {'01','02','03','04'}; numeric=bool(re.fullmatch(r'[0-9０-９,.:≈%万座次个年+/-]+',cp)); title=(y<1.28 and h>=.46 and len(cp)<=30) or (sno==1 and '智烤' in txt); subtitle=(y<1.55 and len(cp)>=18 and h<=.72); keyword=(len(cp)<=10 and h>=.46 and not subtitle) or any(k in txt for k in ('一库清底','双模决策','三智落地','四队长效'))
    p.alignment=PP_ALIGN.CENTER if (seq or y<1.55) else PP_ALIGN.LEFT
    if seq: setea(r,'Times New Roman'); r.font.bold=True; r.font.italic=True; r.font.size=Pt(fit(txt,w,h,42,30,44)); grad(r)
    elif title: setea(r,'华文中宋'); r.font.bold=False; r.font.size=Pt(fit(txt,w,h,56 if (sno==1 and '智烤' in txt) else 36,42 if sno==1 else 24,58 if sno==1 else 37)); grad(r)
    elif subtitle: setea(r,'微软雅黑'); r.font.bold=False; r.font.size=Pt(13.0); solid(r,'EAF7FF')
    elif keyword: setea(r,'华文中宋'); r.font.bold=False; r.font.size=Pt(fit(txt,w,h,29,17,31)); grad(r)
    elif numeric and h>=.4: setea(r,'Times New Roman'); r.font.bold=True; r.font.size=Pt(fit(txt,w,h,28,15,32)); grad(r)
    else: setea(r,'微软雅黑'); r.font.bold=False; r.font.size=Pt(fit(txt,w,h,15,9.5,15.5)); solid(r,'F3F8FC')
def ibbox(parts): return min(s.left for s in parts),min(s.top for s in parts),max(s.left+s.width for s in parts),max(s.top+s.height for s in parts)
def mask(slide,bb,name):
    l,t,r,b=bb; cx=(l+r)/2; cy=(t+b)/2; size=max(r-l,b-t)*1.26; x=cx-size/2; y=cy-size/2
    sh=slide.shapes.add_shape(MSO_SHAPE.OVAL,x,y,size,size); sh.name=name+'_底板_遮住原图标_图标本身为原生矢量'; sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(5,35,67); sh.line.color.rgb=RGBColor(39,199,241); sh.line.width=Pt(.9)
    sf=sh.fill._xPr.find(qn('a:solidFill'))
    if sf is not None:
        a=OxmlElement('a:alpha'); a.set('val','82000'); sf[0].append(a)

for si,slide in enumerate(prs.slides,1):
    for sh in list(slide.shapes):
        if sh.shape_type==13 and sh.name.startswith('背景_'): rm(sh)
    b=slide.shapes.add_picture(str(bgp),0,0,width=prs.slide_width,height=prs.slide_height); b.name=f'背景_S{si:02d}_统一亮蓝山海_无底部光路'; back(slide,b)
    for sh in list(slide.shapes):
        if sh.shape_type==13 and (sh.name.startswith('底部光路_') or sh.name.startswith('原字效切片_')): rm(sh)
    for sh in slide.shapes:
        if sh.shape_type==17 and sh.name.startswith('可编辑文字下层_'):
            c=sh._element.xpath('.//p:cNvPr')[0]; c.attrib.pop('hidden',None); sh.name=sh.name.replace('可编辑文字下层_','可编辑文字_').replace('_双击编辑_上方保留原字效切片','_直接双击修改'); style(sh,si)
    groups=defaultdict(list)
    for sh in iprs.slides[si-1].shapes:
        if sh.name.startswith('可编辑小图标_') and not any(k in sh.name for k in ('_外环','_内环')):
            groups['_'.join(sh.name.split('_')[:4])].append(sh)
    for pfx,parts in groups.items(): mask(slide,ibbox(parts),pfx.replace('可编辑小图标_','小图标原生可编辑_'))
    for pfx,parts in groups.items():
        for ssh in parts:
            el=deepcopy(ssh._element); c=el.xpath('.//p:cNvPr')[0]; c.set('name',c.get('name').replace('可编辑小图标_','小图标原生可编辑_')); slide.shapes._spTree.insert_element_before(el,'p:extLst')
prs.save(OUT)
print(OUT)
