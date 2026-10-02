')
set_text(find_shape(s,'成效不是'), '成效的本质：资源“看得见”，风险“提前判”，投入“算得准”，管护“跑得久”')

# Slide15 closing: innovation + replication rather than pure slogan
s=prs.slides[14]
set_text(find_shape(s,'从一座烤房开始'), '一库把家底变成数据资产，双模把经验变成可解释决策，三智四队把算法变成长期可运行的服务机制')
# add small innovation chips near bottom
add_chip(s,'创新① 动态健康体检',.18,.823,.19,True)
add_chip(s,'创新② 双模耦合决策',.405,.823,.19,True)
add_chip(s,'创新③ 技术+机制闭环',.63,.823,.19,True)
# shift thank you down a bit if exists
sh=find_shape(s,'谢谢大家')
if sh: sh.top=emu_y(.91)

# ---------- visual system pass ----------
for idx,slide in enumerate(prs.slides,1):
    if idx>1:
        add_neon_title_bar(slide)
    # title gradient on prominent text boxes near top
    for sh in slide.shapes:
        if not hasattr(sh,'text_frame'): continue
        txt=sh.text.strip() if hasattr(sh,'text') else ''
        if not txt: continue
        y=sh.top/SH
        # main slide title or big title words
        if (y < .10 and len(txt) <= 32) or txt in {'分不到','修不及','投不准','智烤管家'}:
            set_text_gradient(sh)
        # yellow / model numbering stays yellow
        if re.fullmatch(r'[1234]|0[123]', txt):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb=RGBColor(255,204,48)
    # convert key large rounded rectangles to layered glass/glow, excluding tiny chips
    for sh in slide.shapes:
        if sh.shape_type==1 and sh.width/SW>.18 and sh.height/SH>.09:
            try:
                set_shape_glass(sh, fill='062B4A', fill_alpha=58, line_colors=('62F0FF','087AD8'), glow='15CFFF', glow_alpha=33, glow_pt=5.5, line_pt=1.2)
            except Exception:
                pass
    # add corner brackets around large screenshots
    for sh in slide.shapes:
        if sh.shape_type==13 and sh.width/SW>.18 and sh.height/SH>.14:
            add_corner_brackets(slide,sh.left/SW-.006,sh.top/SH-.008,sh.width/SW+.012,sh.height/SH+.016)

# add animated border to key slides; move behind other content
for idx in [4,6,7,8,10,14]:
    slide=prs.slides[idx-1]
    pic=slide.shapes.add_picture(str(GIF),0,0,width=SW,height=SH)
    # move near back (after nvGrpSpPr and grpSpPr)
    spTree=slide.shapes._spTree
    el=pic._element
    spTree.remove(el)
    spTree.insert(2,el)

# ensure fonts globally consistent
for slide in prs.slides:
    for sh in slide.shapes:
        if hasattr(sh,'text_frame'):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.font.name:
                        r.font.name='Microsoft YaHei'

prs.save(str(OUT))
print(OUT)
print(stamp)


# ===== TEXT/COVER FINALIZATION =====
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pathlib import Path
import datetime, re
src=Path('决赛发布