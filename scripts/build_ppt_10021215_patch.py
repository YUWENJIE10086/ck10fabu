from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
import zipfile

SRC=Path('决赛发布材料/1002PPT决赛完善版/智烤管家_10021155_高阶光效内容深化决赛版.pptx')
OUT=Path('决赛发布材料/1002PPT决赛完善版/智烤管家_10021215_一等奖高阶光效价值深化版.pptx')
NOTE=Path('决赛发布材料/1002PPT决赛完善版/创客25图片_401张逐图样式学习与本版应用_10021215.md')
GIF=Path('/tmp/scan_10021215.gif')

if not SRC.exists():
    raise FileNotFoundError(SRC)

# Reuse the animated cyan scan-light frame already embedded in the premium source deck.
with zipfile.ZipFile(SRC) as z:
    gifs=[n for n in z.namelist() if n.startswith('ppt/media/') and n.lower().endswith('.gif')]
    if not gifs:
        raise RuntimeError('No animated scan-light GIF found in premium source deck')
    GIF.write_bytes(z.read(gifs[0]))

prs=Presentation(SRC)


def replace_preserve(shape, text):
    """Replace text while preserving the first run's premium gradient/glow formatting."""
    if not getattr(shape,'has_text_frame',False):
        return
    tf=shape.text_frame
    if not tf.paragraphs:
        return
    p=tf.paragraphs[0]
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:
            r.text=''
    else:
        p.text=text
    for pp in tf.paragraphs[1:]:
        for r in pp.runs:
            r.text=''


def replace_exact(slide_idx, old, new):
    for sh in prs.slides[slide_idx].shapes:
        if getattr(sh,'has_text_frame',False) and sh.text.strip()==old:
            replace_preserve(sh,new)
            return True
    return False


def replace_prefix(slide_idx, prefix, new):
    for sh in prs.slides[slide_idx].shapes:
        if getattr(sh,'has_text_frame',False) and sh.text.strip().startswith(prefix):
            replace_preserve(sh,new)
            return True
    return False

# Model 1: score is only an intermediate variable; business output is the decision priority.
replace_prefix(5,'技术突破：静态分数','技术高明之处：评分只是中间量，真正输出“剩余寿命 + 风险部位 + 维修优先序”')

# Model 2: the crucial coupling is health result -> effective baking capacity -> real gap -> sandbox.
replace_prefix(6,'技术突破：健康结果','技术高明之处：把模型①健康结果送入模型②，用“有效烤能”替代简单栋数，再做缺口与沙盘推演')

# Judge-facing coupling wording.
replace_exact(7,'双模耦合：健康结果真正进入建设决策','双模耦合：不是两个模型并列，而是一条投入决策链')
replace_exact(7,'存量体检 → 有效烤能 → 增量推演 → 投入决策','模型①先看清存量，模型②再推演增量｜存量体检 → 有效烤能 → 增量推演 → 投入决策')

# Three-smart application is a role workflow, not three isolated functions.
replace_exact(9,'3｜三智落地：把模型装进三类人的工作流','3｜三智落地：不是三套功能，而是三类角色的一套工作流')
replace_exact(9,'管理人员投得准｜烟农技术员用得顺、修得快｜合作社管得久','管理人员投得准｜烟农技术员用得顺、修得快｜合作社管得久｜数据最终回流模型')

# Outcome page: explain how the work mode changed; numbers are evidence.
replace_prefix(13,'应用成效：不是“多了几个数字”','应用成效：不是“多了几个数字”，而是“用、投、管”的工作方式都变了')
replace_exact(13,'一个烘烤季，把找房、抢修、投入、管护四件事真正跑通','一个烘烤季，从烟农找房到行业投入，再到合作社管护，形成可量化闭环')
replace_prefix(13,'价值不是“建了平台”','价值闭环：烟农“用得顺”  →  管理人员“投得准”  →  合作社“管得久”')

# Closing memory sentence.
replace_exact(14,'让数据说话 · 让模型决策 · 让管家落地','从“有多少房”到“有多少有效烤能”')
replace_prefix(14,'从一座烤房开始','让数据说话 · 让模型决策 · 让管家落地')

# Add a subtle animated scan-light frame only to high-value slides.
# The GIF loops during slideshow. Main text/cards/processes remain editable PowerPoint objects.
for idx,xywh in [(9,(.27,1.03,9.46,4.15)),(13,(.27,1.03,9.46,4.15)),(14,(.80,1.22,8.40,3.60))]:
    sl=prs.slides[idx]
    pic=sl.shapes.add_picture(str(GIF),Inches(xywh[0]),Inches(xywh[1]),Inches(xywh[2]),Inches(xywh[3]))
    el=pic._element
    parent=el.getparent()
    parent.remove(el)
    sl.shapes._spTree.insert(2,el)

OUT.parent.mkdir(parents=True,exist_ok=True)
prs.save(OUT)

NOTE.write_text('''# 创客25图片逐图学习与本版应用（10021215）\n\n本轮继续基于已解压的 `创客25图片.zip` 全量样式学习结果，并坚持 `3S共享烤房-0829.pptx` 的母版体系。\n\n## 本轮重点落地\n- 蓝白渐变主标题 + 轻Glow：继承母版高级科技标题语言。\n- 半透明深蓝玻璃舱：取代普通实线框。\n- 青蓝渐变描边 + 外发光 + 角点高光：用于模型、1234、三智和成效页。\n- 动态扫光框：仅在核心创新/成效区域使用，放映时循环扫光，避免满屏霓虹。\n- 黄色只做记忆点：算法、编号、关键数字和结论；红色只做风险。\n- 模型页统一“三步走 + 技术方法 + 业务输出”，截图只做证据。\n- 模型①突出：评分不是终点，真正输出剩余寿命、风险部位和维修优先序。\n- 模型②突出：把模型①健康结果送入有效烤能，再做真实缺口、TOPSIS选址和情景沙盘。\n- 三智页按三类角色工作流组织，并强调业务数据最终回流模型。\n- 成效页按三类主体工作方式变化组织，不再孤立罗列数字。\n- 收尾用“从有多少房到有多少有效烤能”形成最终记忆点。\n\n## 视觉边界\n动态扫光是GIF图片对象，可移动、缩放和替换；标题、玻璃舱、流程、文字和大部分科技组件仍为可编辑PowerPoint对象。\n''',encoding='utf-8')
print(OUT)
print(NOTE)
