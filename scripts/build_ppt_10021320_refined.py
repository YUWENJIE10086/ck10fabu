from __future__ import annotations

from pathlib import Path
import colorsys
from copy import deepcopy

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches
from pptx.oxml.ns import qn

SRC = Path('决赛发布材料/1002PPT决赛完善版/智烤管家_10021215_一等奖高阶光效价值深化版.pptx')
OUT = Path('决赛发布材料/1002PPT决赛完善版/智烤管家_10021320_母版原生克制高级美学重构版.pptx')
NOTE = Path('决赛发布材料/1002PPT决赛完善版/去荧光审美重构说明_10021320.md')

if not SRC.exists():
    raise FileNotFoundError(SRC)

prs = Presentation(SRC)
SW, SH = prs.slide_width, prs.slide_height

# ---------- color / xml helpers ----------
def _hex_to_rgb(h: str):
    h = h.strip().replace('#', '')
    if len(h) != 6:
        return None
    try:
        return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))
    except ValueError:
        return None


def _rgb_to_hex(rgb):
    return ''.join(f'{max(0, min(255, int(round(c)))):02X}' for c in rgb)


def soften_neon_hex(h: str) -> str:
    """Mute only obviously fluorescent cyan/blue/purple; preserve gold/red/white."""
    rgb = _hex_to_rgb(h)
    if not rgb:
        return h
    r, g, b = [x / 255.0 for x in rgb]
    hh, ll, ss = colorsys.rgb_to_hls(r, g, b)
    deg = hh * 360.0

    # Bright cyan / electric blue -> restrained steel-sky blue.
    if 165 <= deg <= 225 and ss >= 0.72 and ll >= 0.42:
        ss2 = min(ss, 0.48)
        ll2 = min(max(ll, 0.58), 0.72)
        rr, gg, bb = colorsys.hls_to_rgb(hh, ll2, ss2)
        return _rgb_to_hex((rr * 255, gg * 255, bb * 255))

    # Neon violet/magenta -> low-saturation blue-violet, avoiding nightclub feel.
    if 270 <= deg <= 330 and ss >= 0.68 and ll >= 0.42:
        hh2 = 0.64  # calm blue-violet rather than pink/magenta
        ss2 = 0.22
        ll2 = min(max(ll, 0.58), 0.68)
        rr, gg, bb = colorsys.hls_to_rgb(hh2, ll2, ss2)
        return _rgb_to_hex((rr * 255, gg * 255, bb * 255))

    return h.upper()


def remove_node(node):
    p = node.getparent()
    if p is not None:
        p.remove(node)


def strip_effects(el):
    # Glow/soft-edge/reflection are the main source of the fluorescent look.
    for tag in ('glow', 'softEdge', 'reflection'):
        for n in list(el.xpath(f'.//a:{tag}')):
            remove_node(n)

    # Cyan outer shadows often read as another glow layer; keep only neutral shadows.
    for sh in list(el.xpath('.//a:outerShdw')):
        srgb = sh.xpath('.//a:srgbClr')
        remove = False
        for c in srgb:
            old = c.get('val', '')
            new = soften_neon_hex(old)
            if new != old.upper():
                remove = True
        if remove:
            remove_node(sh)


def soften_colors(el):
    for c in el.xpath('.//a:srgbClr'):
        old = c.get('val', '')
        if old:
            c.set('val', soften_neon_hex(old))

    # Thin down bright, decorative strokes. 12700 EMU = 1 pt.
    for ln in el.xpath('.//a:ln'):
        w = ln.get('w')
        if not w:
            continue
        try:
            width = int(w)
        except ValueError:
            continue
        changed = False
        for c in ln.xpath('.//a:srgbClr'):
            old = c.get('val', '')
            new = soften_neon_hex(old)
            if old and new != old.upper():
                changed = True
        if changed and width > 15875:  # >1.25 pt
            ln.set('w', '15875')


def shape_is_large_outline_frame(sh) -> bool:
    if sh.shape_type not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
        return False
    if getattr(sh, 'has_text_frame', False) and sh.text.strip():
        return False
    if sh.width < SW * 0.70 or sh.height < SH * 0.46:
        return False
    # Only touch line-only / near-empty decorative containers, never filled background art.
    el = sh._element
    has_solid = bool(el.xpath('.//a:solidFill'))
    has_grad = bool(el.xpath('.//a:gradFill'))
    has_pic = bool(el.xpath('.//a:blip'))
    # If a fill exists under spPr, do not delete; the shape may be structural background.
    sppr = el.xpath('./p:spPr')
    if sppr:
        fills = sppr[0].xpath('./a:solidFill|./a:gradFill|./a:blipFill|./a:pattFill')
        if fills:
            return False
    return not has_pic


def delete_shape(sh):
    el = sh._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


# ---------- 1) remove animated scan-light / giant decorative frames ----------
removed_gifs = 0
removed_frames = 0
for si, slide in enumerate(prs.slides):
    for sh in list(slide.shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                if sh.image.ext.lower() == 'gif':
                    delete_shape(sh)
                    removed_gifs += 1
                    continue
            except Exception:
                pass

        # Cover, transition, closing and judge-facing core pages should breathe.
        if si in {0, 3, 5, 6, 8, 9, 13, 14} and shape_is_large_outline_frame(sh):
            delete_shape(sh)
            removed_frames += 1

# ---------- 2) normalize visual effects across all remaining editable objects ----------
for slide in prs.slides:
    strip_effects(slide._element)
    soften_colors(slide._element)

    # Standardize large rounded cards: restrained outline, never luminous.
    for sh in slide.shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        if sh.width > SW * 0.16 and sh.height > SH * 0.10:
            try:
                if sh.line.color.type is not None and sh.line.color.rgb is not None:
                    old = str(sh.line.color.rgb)
                    new = soften_neon_hex(old)
                    sh.line.color.rgb = RGBColor.from_string(new)
                    if sh.line.width is None or sh.line.width > Pt(1.25):
                        sh.line.width = Pt(1.0)
            except Exception:
                pass

# ---------- 3) add a restrained editorial component system on key slides ----------
def add_kicker(slide, text: str):
    """A small gold label + hairline: learned from competition deck C8/C9, not another glowing box."""
    x = SW - Inches(2.15)
    y = Inches(0.54)
    line = slide.shapes.add_shape(1, x, y + Inches(0.16), Inches(0.34), Pt(1.1))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(214, 174, 93)
    line.line.fill.background()
    tb = slide.shapes.add_textbox(x + Inches(0.43), y, Inches(1.62), Inches(0.34))
    tf = tb.text_frame; tf.clear(); tf.margin_left = 0; tf.margin_right = 0
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.name = 'Microsoft YaHei'; r.font.size = Pt(9.5); r.font.bold = True
    r.font.color.rgb = RGBColor(218, 190, 124)


kickers = {
    3: '1234 方案总览',
    5: '模型①｜存量体检',
    6: '模型②｜增量推演',
    7: '双模耦合｜决策链',
    9: '三智｜角色工作流',
    13: '应用成效｜证据',
    14: '最终记忆点',
}
for idx, label in kickers.items():
    if idx < len(prs.slides):
        add_kicker(prs.slides[idx], label)

# ---------- 4) final pass: newly added components must also have no glow ----------
for slide in prs.slides:
    strip_effects(slide._element)
    soften_colors(slide._element)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)

NOTE.write_text(f'''# 智烤管家决赛PPT｜去荧光、组件化审美重构（10021320）

## 这次不是“再加一层光效”，而是做减法

本版以 `智烤管家_10021215_一等奖高阶光效价值深化版.pptx` 为内容基底，保留已经深化好的演讲逻辑、模型价值和证据截图，视觉上执行 **“母版原生、克制、舞台远读、证据优先”**。

### 已直接删除
- 全部动态扫光 GIF：共移除 **{removed_gifs}** 个，彻底取消“霓虹灯框”观感。
- 核心页的大面积空心装饰框：共移除 **{removed_frames}** 个，避免内容被边框包死。
- 全套 Glow / SoftEdge / Reflection；青蓝色阴影不再伪装成第二层发光。

### 颜色与线条
- 保留深蓝母版、白字、黄色记忆点。
- 电光青/荧光蓝降饱和为钢蓝/雾蓝；紫粉荧光统一压成低饱和蓝紫。
- 亮色粗描边压到约 1pt～1.25pt，形成“精密仪器感”，不是“电竞灯带感”。
- 白→青标题渐变保留，但不再叠 Glow。

## 真正按“创客25图片”学的组件类型

这次不再把“高级”理解成一堆发光框，而是按之前 401 张样式分类重新收口：

1. **C8 舞台极简组件**：封面、转场、结尾删大框，靠大标题、留白和一句结论形成记忆。
2. **C4 三步技术链**：模型页保留“步骤编号 + 方法 + 输出”，线条变细，流程本身成为视觉骨架。
3. **C2/C6 证据卡**：截图仍是主证据，但取消外发光；后台/小程序/BI看板像“证据窗”，不是装饰画。
4. **C9 模块化关系**：1234、双模耦合继续靠模块和连接关系表达，不再靠框的亮度抢注意力。
5. **C10 蓝青高级组件**：只保留蓝白渐变标题、轻玻璃感、金色题签；去掉霓虹边、扫光、角点爆亮。
6. **编辑式 Kicker**：关键页右上增加“小金线 + 小标题”，形成比赛级章节锚点，组件很轻，不抢正文。
7. **成效页大数字逻辑**：数字是主角，卡片只是承托；不再让每个 KPI 都像一个发光按钮。

## 评委视角的视觉原则
- 一页一个结论；
- 远距离先看标题与数字；
- 近距离再看流程与证据；
- 技术页“方法有骨架”，成效页“数据有主次”；
- 高级感来自留白、重复组件、线条精度和真实截图，而不是荧光强度。

输出文件：`{OUT.name}`
''', encoding='utf-8')

print(OUT)
print(NOTE)
