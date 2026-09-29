
from pathlib import Path
from docx import Document
from pptx import Presentation
import subprocess, sys, traceback

ROOT=Path('.')
OUT=Path('.analysis_extract_0930')
OUT.mkdir(exist_ok=True)

FILES = [
('official_report','读取资料/3之前项目官方叙述过的/附件2：宜昌市局+智烤管家-基于健康评估的烤房用投管一体化解决方案+数创成果报告.docx'),
('official_tech','读取资料/3之前项目官方叙述过的/附件3：宜昌市局+智烤管家-基于健康评估的烤房用投管一体化解决方案+技术资料说明书.docx'),
('draft_0922','读取资料/4一版初稿/智烤管家演讲讲稿_0922.docx'),
('draft_pp','读取资料/4一版初稿/智烤管家（PP主线）.doc'),
('draft_old','读取资料/4一版初稿/讲稿初稿.doc'),
('models_ppt','读取资料/2我的项目相关的零碎的/几个模型讲解-1.pptx'),
('ppt_0909','读取资料/2我的项目相关的零碎的/0909-1.pptx'),
('report_0909','读取资料/2我的项目相关的零碎的/数创报告-0909-1.docx'),
('v1_0929','决赛发布材料/0929五版发布稿/01_V1_双模型硬核版_最接近3D智慧对标风格.docx'),
('background','读取资料/2我的项目相关的零碎的/背景梳理.docx'),
('video_ref','读取资料/2我的项目相关的零碎的/视频画面呈现参考.docx'),
('tech_0907','读取资料/2我的项目相关的零碎的/技术报告-0907-2.docx'),
('health_life','读取资料/2我的项目相关的零碎的/健康寿命.docx'),
('site_algo','读取资料/2我的项目相关的零碎的/选址算法.docx'),
]

def docx_text(p):
    d=Document(p)
    out=[]
    for para in d.paragraphs:
        t=para.text.strip()
        if t: out.append(t)
    for ti,table in enumerate(d.tables,1):
        out.append(f'\n[TABLE {ti}]')
        for row in table.rows:
            vals=[c.text.replace('\n',' | ').strip() for c in row.cells]
            out.append('\t'.join(vals))
    return '\n'.join(out)

def pptx_text(p):
    prs=Presentation(p)
    out=[]
    for i,slide in enumerate(prs.slides,1):
        out.append(f'\n===== SLIDE {i} =====')
        texts=[]
        for shape in slide.shapes:
            if hasattr(shape,'text'):
                t=shape.text.strip()
                if t: texts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append('\t'.join(cell.text.strip() for cell in row.cells))
        out.extend(texts)
    return '\n'.join(out)

def doc_text(p):
    try:
        cp=subprocess.run(['antiword',str(p)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,check=True)
        for enc in ('utf-8','gb18030','latin1'):
            try: return cp.stdout.decode(enc)
            except: pass
        return cp.stdout.decode('utf-8','ignore')
    except Exception as e:
        return 'ANTIWORD ERROR: '+repr(e)+'\n'+subprocess.run(['strings','-el',str(p)],stdout=subprocess.PIPE).stdout.decode('utf-8','ignore')

for name,rel in FILES:
    p=ROOT/rel
    try:
        if p.suffix.lower()=='.docx':
            txt=docx_text(p)
        elif p.suffix.lower()=='.pptx':
            txt=pptx_text(p)
        elif p.suffix.lower()=='.doc':
            txt=doc_text(p)
        else:
            txt=p.read_text(encoding='utf-8',errors='ignore')
    except Exception:
        txt=traceback.format_exc()
    (OUT/f'{name}.txt').write_text(txt,encoding='utf-8')
    print(name, len(txt))
