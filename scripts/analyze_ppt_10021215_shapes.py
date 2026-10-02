from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SRC = Path('决赛发布材料/1002PPT决赛完善版/智烤管家_10021215_一等奖高阶光效价值深化版.pptx')
OUT = Path('scripts/_tmp_ppt_shape_audit_10021215.md')
prs = Presentation(SRC)


def emu_in(v):
    return round(v / 914400, 2)


def color_of(fill_or_line):
    try:
        c = fill_or_line.fore_color
        if c.type is not None and c.rgb is not None:
            return str(c.rgb)
        if c.type is not None and c.theme_color is not None:
            return f'theme:{c.theme_color}'
    except Exception:
        pass
    return '-'


def fill_desc(sh):
    try:
        ft = sh.fill.type
        if ft is None:
            return 'none'
        return f'{ft}:{color_of(sh.fill)}'
    except Exception:
        return '-'


def line_desc(sh):
    try:
        col = color_of(sh.line)
        w = round((sh.line.width or 0) / 12700, 2)
        return f'{col}/{w}pt'
    except Exception:
        return '-'

lines=[]
lines.append(f'# Shape audit: {SRC.name}')
lines.append(f'- slides: {len(prs.slides)}')
lines.append(f'- size: {emu_in(prs.slide_width)} x {emu_in(prs.slide_height)} in')

for si, slide in enumerate(prs.slides, start=1):
    texts=[]
    for sh in slide.shapes:
        if getattr(sh, 'has_text_frame', False) and sh.text.strip():
            texts.append(sh.text.strip().replace('\n','｜')[:120])
    lines.append(f'\n## S{si}  ' + (' / '.join(texts[:3]) if texts else ''))
    lines.append('|#|type|name|x|y|w|h|fill|line|text|')
    lines.append('|---:|---|---|---:|---:|---:|---:|---|---|---|')
    for i, sh in enumerate(slide.shapes, start=1):
        st = str(sh.shape_type)
        text=''
        if getattr(sh, 'has_text_frame', False):
            text=sh.text.strip().replace('\n','｜').replace('|','¦')[:160]
        nm=(sh.name or '').replace('|','¦')
        lines.append(f'|{i}|{st}|{nm}|{emu_in(sh.left)}|{emu_in(sh.top)}|{emu_in(sh.width)}|{emu_in(sh.height)}|{fill_desc(sh)}|{line_desc(sh)}|{text}|')

OUT.write_text('\n'.join(lines), encoding='utf-8')
print(OUT)
