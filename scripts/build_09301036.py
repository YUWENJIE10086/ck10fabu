from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

OUT=Path('决赛发布材料/09301036发布稿')
PATH=OUT/'01_智烤管家_湖北省创客大赛决赛双人发布稿_1234双模三智强化版_09301036.docx'
TITLE_FONT='方正小标宋简体'
BODY_FONT='仿宋_GB2312'
HEAD_FONT='黑体'

speech=[
('A','尊敬的各位领导、各位专家评委，大家好！我是宜昌市局“智同道合”创客小组的余文杰。'),
('B','大家好，我是项目组成员胡浩博。今天汇报的项目是《智烤管家——基于健康评估的烤房用投管一体化解决方案》。'),
('A','烤房，是烟叶提质增香的“匠心工坊”，也是烟农稳收增收的基石。但一线最头疼的，其实就三件事：一是“分不到”——旺季烤房靠人工协调，哪里有空、什么时候能用，烟农心里没数；二是“修不及”——设备一坏，报修、找人、找配件都靠人工；三是“投不准”——哪栋该修、哪个村该建，缺少量化依据。'),
('B','看上去是用房、抢修和投入三个问题。往深处看，真正卡在哪里？'),
('A','2025年，我们通过匿名问卷和线下走访继续追问，发现烟农愁调剂抢修，烟站愁底数不清、投入无据，合作社愁资金、责任和标准难落实。“三愁”背后，其实是一条链断了：资源没有统筹、投入没有量化、管护没有闭环。'),
('A','因此，我们提出“智烤管家”。以“一个数据库、两个模型、一个协同平台”为骨架，并把整套方案概括成“1234”：一库清底、双模决策、三智落地、四队长效，一个平台把它们串成闭环。'),
('B','第一步为什么是“一库清底”？'),
('A','因为底数不准，模型就没有可靠输入。我们以统一烤房编号为主线，把基础信息、部件状态、维修养护、使用和管护记录归到同一栋烤房名下。2026年，秭归2442座烤房完成全口径入库和健康评价，实现“一房一档、状态可查、过程可追”。'),
('B','家底清楚了，评委最想知道：2442座烤房，凭什么判断谁先修、还能用多久？'),
('A','第一套核心模型——烤房健康评估模型。我们抓住12个关键部件，把部件状态、指导寿命和现场评价标准化，再用AHP层次分析法，把专家对部件重要程度的判断转成可计算、可校验的权重；整栋健康分由部件健康80%、管护10%、人工评价10%共同形成。'),
('A','但“今天多少分”还不够。我们进一步用威布尔分布拟合部件寿命，结合结构件、功能件、维修记录、房龄和健康水平，预测剩余寿命，最终输出健康等级、剩余年限、风险部位和维修优先序。'),
('B','也就是说，它不是为了“打分”，而是回答三件事：现在健不健康、还能用多久、先修哪里。'),
('A','对。我们把静态台账变成了“健康—寿命—风险—维修”的动态体检。'),
('B','那第二个关键问题：该调的调了、该修的修了，一个村还是不够用，怎么证明它真的需要新建？'),
('A','这就是第二套核心模型——千亩村发展导引模型。它先把第一套模型的健康结果，与烤房容量、使用状态一起折算成“有效烤能”，再结合种烟规模、烘烤需求和发展条件判断真实缺口。健康在用房、闲置房、待修房，不再简单都算成“一栋”。'),
('A','进入新建论证后，用AHP确定主观权重、熵权法校正客观差异，再从地形、水文、电力、交通、烟田匹配、环境安全、集群效益和建设成本等维度评价候选点，用TOPSIS排序；同时用情景沙盘改变种植规模或新增栋数，实时重算烤能、负载和缺口。'),
('B','所以它连续回答四个问题：缺不缺、缺多少、建哪里、建后有没有改善。'),
('A','两个模型由此连成一条决策链：先看清存量，再推演增量，形成“先调剂、再修复，仍有结构性缺口才新建”的投入顺序。2026年，这套方法已用于建设方案论证，支撑争取县政府新建烤房29栋，其中肖家坪村7栋。'),
('A','【视频约1分钟】模型怎么算、平台怎么跑，下面请大家通过视频看一看。'),
('B','模型算清楚了，怎样真正落到烟农、烟站和合作社手里？'),
('A','第一，智用解忧。烟农通过微信查询附近烤房、预约调剂，设备故障时一键报修；平台结合烤房档案和风险部位快速派单、匹配人员和配件。磨坪烤房遭遇雷击后，就是通过线上报修和协同抢修，在3小时内恢复使用。'),
('A','第二，智投增效。存量看健康和风险，形成巡检、修复和资金优先序；增量看有效烤能和真实缺口，再做选址排序和沙盘推演。让投入从“凭经验申报”转向“先体检、再算缺口、最后看效果”。'),
('B','技术能把决策算清楚，但怎样避免平台建起来以后又回到“没人管”？'),
('A','第三，智管提质。我们把线上流程和线下机制一起改，组建烟叶调制、紧急抢修、日常管护、综合管理四支服务队伍，把申请、派单、处置、验收、评价、数据回写串成闭环，并配套管护制度和基金机制，让烤房有人管、有钱管、有标准地管。'),
('B','也就是：技术负责“算得准”，机制负责“跑得久”。'),
('A','最终看成效。一个烘烤季下来，2442座烤房全部入库并完成健康评价，筛出340多栋健康分低于60分的预警或危险烤房；完成48次烤房调剂；18次报修抢修全部控制在3小时以内，16条平台工单平均响应0.38小时；531份有效问卷综合满意度达到96.8%。'),
('B','两级合作社统筹管护1363栋烤房，占本年度在用烤房的81%；管护基金累计近190万元，四支服务队伍和制度标准同步落地。'),
('A','回到开场的三个问题：分不到，用“智用”盘活；修不及，用健康模型和协同抢修提速；投不准，用双模型和沙盘把投入算清。'),
('B','记住“1234”：一库清底、双模决策、三智落地、四队长效。'),
('A','从“有多少房”到“有多少有效烤能”，从“坏了再修”到“风险预警”，从“凭经验投”到“数据推演”。这就是“智烤管家”。我们的汇报完毕，谢谢大家！')
]

KEYS=['分不到','修不及','投不准','三愁','1234','一库清底、双模决策、三智落地、四队长效','2442座','烤房健康评估模型','12个关键部件','AHP层次分析法','80%','10%','威布尔分布','健康—寿命—风险—维修','千亩村发展导引模型','有效烤能','熵权法','TOPSIS','先调剂、再修复，仍有结构性缺口才新建','29栋','7栋','智用解忧','智投增效','智管提质','3小时内','340多栋','低于60分','48次','0.38小时','531份','96.8%','1363栋','81%','近190万元','技术负责“算得准”，机制负责“跑得久”']

def set_font(run,name,size,bold=None):
    run.font.name=name; run.font.size=Pt(size)
    if bold is not None: run.font.bold=bold
    rpr=run._element.get_or_add_rPr(); rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    for k in ('ascii','hAnsi','eastAsia','cs'): rfonts.set(qn('w:'+k),name)

def add_rich(p,text):
    hits=[k for k in KEYS if k in text]
    if not hits:
        r=p.add_run(text); set_font(r,BODY_FONT,16); return
    pat='('+'|'.join(map(re.escape,sorted(hits,key=len,reverse=True)))+')'
    for x in re.split(pat,text):
        if not x: continue
        r=p.add_run(x); set_font(r,BODY_FONT,16,x in KEYS)

OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(2.54); sec.bottom_margin=Cm(2.54)
sec.left_margin=Cm(3.175); sec.right_margin=Cm(3.175)
sec.header_distance=Cm(1.501); sec.footer_distance=Cm(1.75)
normal=doc.styles['Normal']; normal.font.name=BODY_FONT; normal.font.size=Pt(16)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),BODY_FONT)
normal.paragraph_format.line_spacing=Pt(27); normal.paragraph_format.space_after=Pt(0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14)
r=p.add_run('智烤管家——基于健康评估的烤房用投管一体化解决方案'); set_font(r,TITLE_FONT,22,True)
for role,text in speech:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.line_spacing=Pt(27); p.paragraph_format.space_after=Pt(0)
    r=p.add_run(role+'：'); set_font(r,HEAD_FONT,16,True)
    add_rich(p,text)
doc.save(PATH)
print(PATH)
