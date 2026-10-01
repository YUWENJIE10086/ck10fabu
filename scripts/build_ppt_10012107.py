from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

SRC=Path('ppt参考/3S共享烤房-0829.pptx')
OUTDIR=Path('决赛发布材料/10012107PPT母版原生组件版'); OUTDIR.mkdir(parents=True,exist_ok=True)
OUT=OUTDIR/'智烤管家_10012107_3S原母版逐页替换决赛版.pptx'
TMP=Path('.tmp_10012107'); TMP.mkdir(exist_ok=True)
media={'bg':'image17.png','mini':'image30.jpeg','web_a':'image33.png','web_b':'image34.png','web_c':'image35.png','repair_a':'image36.jpeg','repair_b':'image37.jpeg','bi1':'image39.png','bi2':'image40.png'}
with ZipFile(SRC) as z:
    for k,n in media.items():
        p=TMP/n; p.write_bytes(z.read('ppt/media/'+n)); media[k]=p
prs=Presentation(SRC)

def xt(sh,t):
    ns=sh._element.xpath('.//a:t')
    if ns:
        ns[0].text=t
        for x in ns[1:]: x.text=''

def mt(sh,t,sz=None):
    tf=sh.text_frame; p0=tf.paragraphs[0]; r0=p0.runs[0] if p0.runs else None
    name=r0.font.name if r0 else None; size=r0.font.size if r0 else None; bold=r0.font.bold if r0 else None; align=p0.alignment
    tf.clear()
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=line; r.font.name=name; r.font.size=Pt(sz) if sz else size; r.font.bold=bold

def S(sl,i,t): xt(sl.shapes[i],t)
def M(sl,i,t,sz=None): mt(sl.shapes[i],t,sz)
def rm(sh): sh._element.getparent().remove(sh._element)
def tb(sl,t,x,y,w,h,sz=24,c=(255,255,255),b=False,a=PP_ALIGN.LEFT):
    q=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); p=q.text_frame.paragraphs[0]; p.alignment=a; r=p.add_run(); r.text=t; r.font.name='Microsoft YaHei'; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=RGBColor(*c); return q
def frame(sl,x,y,w,h):
    q=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=RGBColor(2,33,64); q.fill.transparency=8; q.line.color.rgb=RGBColor(15,191,255); q.line.width=Pt(1.2); return q
def pic(sl,p,x,y,w,h): return sl.shapes.add_picture(str(p),Inches(x),Inches(y),Inches(w),Inches(h))

# 1 cover
s=prs.slides[0]
for i,t in [(4,'让每一座烤房，都有“体检”、有“决策”、有“管家”'),(5,'汇报单位：宜昌市烟草公司    汇报人：余文杰    2026.10'),(8,'基于健康评估的烤房用投管一体化解决方案')]: S(s,i,t)
for sh in [s.shapes[7],s.shapes[6]]: rm(sh)
tb(s,'智烤管家',2.35,1.45,5.3,.9,42,(255,255,255),True,PP_ALIGN.CENTER); tb(s,'ZHIKAO  MANAGER',3.25,2.23,3.5,.28,10,(48,211,255),False,PP_ALIGN.CENTER)
# 3 pains
s=prs.slides[2]
for i,t in [(3,'一针见血：烤房一线最头疼的3件事'),(16,'烤房不是“有没有”，关键是：好不好用、怎么投入、谁来持续管'),(10,'第一“忧”'),(11,'第二“忧”'),(12,'第三“忧”'),(7,'分不到'),(8,'修不及'),(9,'投不准')]: S(s,i,t)
M(s,13,'旺季供需不均\n人工协调、调剂慢'); M(s,14,'设备突发故障\n找人、找配件、进度不透明'); M(s,5,'哪栋该修、哪个村该建\n资金投入缺少量化依据')
# 4 root
s=prs.slides[3]
for i,t in [(0,'问题不在一个环节，而在“用·投·管”没有连起来'),(10,'资源无统筹'),(12,'投入无量化'),(14,'管护无闭环'),(18,'分不到'),(19,'投不准'),(20,'修不及')]: S(s,i,t)
M(s,11,'烟农调剂靠人工\n闲置与紧缺并存'); M(s,13,'底数不清、状态不明\n修复新建凭经验'); M(s,15,'责任、资金、标准分散\n过程反馈难沉淀')
# 5 1234
s=prs.slides[4]
for i,t in [(2,'智烤管家：记住“1234”，就记住整套解决方案'),(5,'1库'),(6,'清底'),(21,'Database'),(17,'2模'),(18,'决策'),(22,'Decision'),(11,'3智'),(12,'落地'),(23,'Application'),(7,'低代码采数 · 一房一档'),(19,'健康评估 · 发展导引'),(13,'智用 · 智投 · 智管'),(26,'1库清底'),(27,'2模决策'),(28,'3智落地'),(30,'数据'),(32,'算法'),(34,'应用')]: S(s,i,t)
M(s,8,'把2442座烤房\n“看清楚”'); M(s,20,'把“修哪里、建哪里”\n“算清楚”'); M(s,14,'把模型真正\n“跑起来”'); M(s,25,'一个数据库打底、两个模型决策、三个智慧应用落地，再用四支队伍把服务长期跑起来')
o=s.shapes[28]; e=deepcopy(o._element); s.shapes._spTree.insert_element_before(e,'p:extLst'); n=s.shapes[-1]; n.left=Inches(8.12); n.top=o.top; n.width=o.width; n.height=o.height; xt(n,'4队长效')
# 7 database
s=prs.slides[6]
for i,t in [(46,'1库  低代码采数 · 一房一档'),(37,'数据底座'),(5,'多源低代码采数'),(6,'统一入库'),(7,'持续回写'),(32,'建设年份'),(33,'部件状态'),(34,'维修历史'),(35,'使用记录'),(36,'管护评价'),(27,'统一烤房编号'),(28,'字段标准化 · 批量导入'),(29,'移动端核验'),(30,'异常校验'),(31,'2442座完整档案'),(1,'低门槛采集    可持续更新'),(48,'让模型的每一个判断，都有真实数据支撑')]: S(s,i,t)
for sh in [x for x in list(s.shapes) if x.shape_type==13 and x.left/Inches(1)>4.8]: rm(sh)
frame(s,5.25,1.13,2.2,3.65); pic(s,media['mini'],5.38,1.25,1.94,3.38); tb(s,'低代码采数',7.72,1.5,1.45,.34,17,(255,203,0),True); tb(s,'一次建档\n持续回写',7.65,2.05,1.65,.85,21,(255,255,255),True,PP_ALIGN.CENTER)
# 9 model1
s=prs.slides[8]
for i,t in [(22,'2模①  烤房健康评估模型'),(1,'单栋体检'),(5,'STEP 1  输入'),(6,'STEP 2  综合评分'),(7,'STEP 3  寿命与风险'),(17,'12个关键部件'),(18,'部件状态'),(19,'指导寿命'),(20,'现场评价'),(8,'健康分评估'),(11,'AHP层次分析法'),(9,'剩余寿命预测'),(12,'威布尔寿命分布'),(10,'风险部位排序'),(13,'健康等级'),(14,'剩余寿命'),(15,'风险部位'),(16,'维修优先序'),(21,'不是“打一个分”，而是回答：健不健康？还能用多久？先修哪里？'),(35,'把静态台账变成动态体检')]: S(s,i,t)
for sh in [x for x in list(s.shapes) if x.shape_type==13 and x.left/Inches(1)>4.8]: rm(sh)
for p,co in [(media['web_a'],(5.35,1,2,1.25)),(media['web_b'],(7.5,1,2,1.25)),(media['web_c'],(5.35,2.45,4.15,2.35))]: frame(s,*co); pic(s,p,co[0]+.08,co[1]+.08,co[2]-.16,co[3]-.16)
# 11 model1 results
s=prs.slides[10]
for i,t in [(9,'模型①落地｜从“坏了再修”到“风险前置”'),(1,'报修抢修全部控制'),(2,'≤3小时'),(4,'健康分低于60分'),(5,'340+栋'),(17,'平台工单平均响应'),(18,'0.38小时'),(20,'报修抢修'),(21,'18次')]: S(s,i,t)
M(s,11,'过去\n故障后再找人'); M(s,13,'现在\n模型预警 + 协同派单')
# 12 model2
s=prs.slides[11]
for i,t in [(17,'2模②  千亩村发展导引模型'),(19,'村级决策'),(5,'STEP 1  有效烤能'),(6,'STEP 2  组合评价'),(7,'STEP 3  建设建议'),(8,'健康结果'),(11,'容量 / 使用状态'),(9,'种烟规模'),(10,'发展条件'),(13,'AHP主观权重'),(14,'熵权客观校正'),(15,'有效烤能 + 真实缺口'),(12,'选址与新建建议'),(16,'助力决策    合理配置'),(28,'回答：真缺吗？缺多少？建哪里？')]: S(s,i,t)
frame(s,4.85,1,4.75,3.65); pic(s,media['bi1'],4.96,1.12,4.53,3.4)
# 13 sandbox
s=prs.slides[12]
for i,t in [(15,'2模②  情景沙盘｜先把未来演一遍'),(17,'沙盘推演'),(5,'STEP 1  设情景'),(6,'STEP 2  动态重算'),(7,'STEP 3  看改善'),(8,'当前规模'),(9,'+2% / +5% / +10%'),(10,'新增烤房数量'),(12,'重算烤能 / 负载 / 缺口'),(13,'比较建设前后'),(11,'输出建设建议'),(14,'回答：建不建、建多少、建后有没有改善'),(26,'投入之前，先看清“建后有没有用”')]: S(s,i,t)
frame(s,4.85,1,4.75,3.65); pic(s,media['bi2'],4.96,1.12,4.53,3.4)
# 14 video
s=prs.slides[13]; S(s,3,'模型算清楚以后，还要真正跑起来'); tb(s,'VIDEO · 约1分钟',3.55,2.05,2.9,.45,20,(255,201,0),True,PP_ALIGN.CENTER); tb(s,'从模型判断 → 平台协同 → 烟区一线',2.25,3.1,5.5,.55,24,(255,255,255),True,PP_ALIGN.CENTER)
# 10 smart use
s=prs.slides[9]
for i,t in [(22,'3智①  智用｜烟农一部手机完成调剂与报修'),(1,'智用'),(7,'查房调剂'),(26,'一键报修'),(11,'进度可见'),(16,'附近烤房'),(17,'可用状态'),(18,'在线预约'),(19,'调剂申请'),(27,'风险部位提示'),(20,'快速派单'),(21,'验收评价'),(15,'少打电话    少跑腿    少等待'),(25,'把服务真正送到烟农手机里')]: S(s,i,t)
for sh in [x for x in list(s.shapes) if x.shape_type==13 and x.left/Inches(1)>4.4]: rm(sh)
for p,co in [(media['repair_a'],(5.6,1.2,1.55,3.65)),(media['repair_b'],(7.55,1.2,1.55,3.65))]: frame(s,*co); pic(s,p,co[0]+.08,co[1]+.08,co[2]-.16,co[3]-.16)
# 15 three smart
s=prs.slides[14]
for sh in [x for x in list(s.shapes) if x.shape_type==13]: rm(sh)
pic0=s.shapes.add_picture(str(media['bg']),0,0,prs.slide_width,prs.slide_height); sp=s.shapes._spTree; el=pic0._element; sp.remove(el); sp.insert(2,el)
tb(s,'3智融合｜一端服务烟农 · 一屏管理全域',.55,.16,8.9,.55,25,(61,220,255),True,PP_ALIGN.CENTER)
for title,p,x in [('智用｜烟农端',media['web_a'],.55),('智投｜管理端',media['web_b'],3.55),('智管｜协同端',media['web_c'],6.55)]: frame(s,x,1,2.9,3.8); tb(s,title,x+.15,1.12,2.6,.35,16,(255,201,0),True); pic(s,p,x+.16,1.55,2.58,2.95)
for t,x in [('查房 / 报修',.9),('健康 / 缺口 / 投入',4),('派单 / 验收 / 回写',7)]: tb(s,t,x,5.02,2,.28,15,(255,255,255),True,PP_ALIGN.CENTER)
# 16 four teams
s=prs.slides[15]; S(s,0,'4队长效｜技术算得准，还要让机制跑得久'); M(s,2,'4支服务队\n烟叶调制｜紧急抢修\n日常管护｜综合管理\n\n申请 → 派单 → 处置\n→ 验收 → 评价 → 回写'); M(s,3,'长效机制\n制度标准｜基金机制\n责任到人｜评价闭环\n\n两级合作社统筹管护1363栋\n占本年度在用烤房81%'); S(s,5,'管护基金累计近190万元'); S(s,6,'把“平台有人用”变成“机制持续跑”')
for sh in [x for x in list(s.shapes) if x.shape_type==13]: rm(sh)
# 8 results
s=prs.slides[7]
for i,t in [(9,'落地成效｜结果经得起数字验证'),(1,'全量入库并完成健康评价'),(2,'2442座'),(4,'综合满意度'),(5,'96.8%'),(16,'合作社统筹管护'),(17,'81%'),(19,'管护基金累计'),(20,'≈190万')]: S(s,i,t)
M(s,11,'过去\n经验分散'); M(s,13,'现在\n数据决策'); M(s,14,'转\n\n变')
# 17 end
s=prs.slides[16]; S(s,0,'让每一座烤房，都有“体检”、有“决策”、有“管家”'); S(s,2,'一库清底 · 双模决策 · 三智落地 · 四队长效')
# order: drop blank slide2 + dense blueprint slide6
order=[0,2,3,4,6,8,10,11,12,13,9,14,15,7,16]; L=prs.slides._sldIdLst; ids=list(L)
for e in ids: L.remove(e)
for i in order: L.append(ids[i])
prs.save(OUT)
print(OUT)
