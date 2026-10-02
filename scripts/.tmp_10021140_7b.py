材料/1002PPT决赛完善版/智烤管家_10021140_高级科技动态一等奖精修版.stage1.pptx')
stamp='10021140'
out=src.parent/'智烤管家_10021140_高级科技动态一等奖精修版.pptx'
prs=Presentation(str(src))
LIGHT=RGBColor(210,238,248); WHITE=RGBColor(242,251,255); CYAN=RGBColor(122,235,255); YEL=RGBColor(255,207,55); GREEN=RGBColor(91,235,178)

def find(slide, needle):
    for sh in slide.shapes:
        if hasattr(sh,'text') and needle in sh.text:
            return sh

def paint(sh, color=LIGHT, size=None, bold=None):
    if not sh or not hasattr(sh,'text_frame'): return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.font.name='Microsoft YaHei'
            r.font.color.rgb=color
            if size: r.font.size=Pt(size)
            if bold is not None: r.font.bold=bold

# recolor replaced text elements
items=[
(2,'不是“烤房不够”',CYAN,10.5,False),(2,'旺季资源状态',LIGHT,10.5,False),(2,'故障突发后',LIGHT,10.5,False),(2,'维修、新建仍偏经验',LIGHT,10.5,False),(2,'表面是',WHITE,10.5,True),
(3,'2025匿名问卷',CYAN,10.2,False),(3,'用房调剂难',LIGHT,9.5,False),(3,'底数状态不清',LIGHT,9.5,False),(3,'资金不足',LIGHT,9.5,False),(3,'“三愁”不是三个孤立问题',WHITE,10.2,True),
(4,'一个协同平台，把“数据底座',CYAN,10.2,False),
(5,'低代码采数｜统一烤房编号',CYAN,10.2,False),(5,'数据采集覆盖18个关键部件',YEL,9.2,True),(5,'数据库不是一次性建档',WHITE,9.4,False),
(6,'三步把“人工看状态”',CYAN,10.5,False),(6,'输出：健康等级',YEL,10.5,True),
(8,'双模型不是两个独立打分器',WHITE,9.3,False),
(10,'三智不是三个菜单',CYAN,10.2,False),
(11,'管理人员：模型给出',CYAN,10.2,False),(11,'方案论证：县级29栋',YEL,10.0,True),
(12,'烟农 + 技术员：从',CYAN,10.0,False),(12,'实战：雷击停运',YEL,8.8,True),
(13,'合作社：线上把流程闭环',CYAN,10.0,False),
(14,'不只看“做了多少”',CYAN,10.2,False),(14,'找房靠问',WHITE,9.2,False),(14,'凭经验排查',WHITE,9.2,False),(14,'责任悬空',WHITE,9.2,False),(14,'成效的本质',YEL,9.8,True),
(15,'一库把家底变成数据资产',WHITE,10.0,False),
]
for slide_no, needle, color, size, bold in items:
    paint(find(prs.slides[slide_no-1],needle),color,size,bold)
# ensure slide 15 slogan is bright
paint(find(prs.slides[14],'好房充分用'),WHITE,15,True)
paint(find(prs.slides[14],'一库清底'),YEL,10.5,True)
paint(find(prs.slides[14],'谢谢大家'),WHITE,12,True)

prs.save(str(out))
print(out)

# final cleanup: remove decorative out-of-canvas cover connectors and stage1 temp
from pptx.enum.shapes import MSO_SHAPE_TYPE
pfix=Presentation(str(out)); s0=pfix.slides[0]
for sh in list(s0.shapes):
    if sh.shape_type==MSO_SHAPE_TYPE.LINE and (sh.left<0 or sh.top<0 or sh.left+sh.width>pfix.slide_width or sh.top+sh.height>pfix.slide_height):
        sh._element.getparent().remove(sh._element)
for sh in s0.shapes:
    if sh.left < 0 and abs(sh.left) < 100:
        delta=-sh.left; sh.left=0; sh.width=max(1,sh.width-delta)
pfix.save(str(out))
try: src.unlink()
except Exception: pass
print('FINAL',out)
