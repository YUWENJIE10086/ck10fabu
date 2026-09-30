from pathlib import Path
from docx import Document
OUT=Path(".tmp_video_0941"); OUT.mkdir(exist_ok=True)
for name,p in [
("mini","读取资料/2我的项目相关的零碎的/小程序视频演示脚本.docx"),
("ref","读取资料/2我的项目相关的零碎的/视频画面呈现参考.docx"),
("official","读取资料/3之前项目官方叙述过的/附件2：宜昌市局+智烤管家-基于健康评估的烤房用投管一体化解决方案+数创成果报告.docx")
]:
 d=Document(p); out=[]
 for para in d.paragraphs:
  if para.text.strip(): out.append(para.text.strip())
 for t in d.tables:
  for row in t.rows:
   vals=[c.text.replace("\n"," | ").strip() for c in row.cells]
   if any(vals): out.append("\t".join(vals))
 sec=d.sections[0]
 meta=[f"top={sec.top_margin.cm:.2f}",f"bottom={sec.bottom_margin.cm:.2f}",f"left={sec.left_margin.cm:.2f}",f"right={sec.right_margin.cm:.2f}",f"header={sec.header_distance.cm:.2f}",f"footer={sec.footer_distance.cm:.2f}"]
 (OUT/(name+".txt")).write_text("\n".join(meta)+"\n---\n"+"\n".join(out),encoding="utf-8")
 print(name,len(out),meta)
