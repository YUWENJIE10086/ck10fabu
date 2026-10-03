from pathlib import Path
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

BASE = Path('决赛发布材料/1002PPT决赛完善版/1003')
SRC = BASE / '智烤管家_10031649_A版精修终稿_文字渐变可编辑_小图标原生矢量_统一亮背景无光路.pptx'
OUT = Path('智烤管家_10031708_A版精修_通用字体兼容版_微软雅黑Arial_渐变可编辑.pptx')
EMU = 914400

prs = Presentation(SRC)

def set_typefaces(run, east_asia='Microsoft YaHei', latin='Arial'):
    run.font.name = latin
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = OxmlElement('a:ea')
        rPr.append(ea)
    ea.set('typeface', east_asia)
    latin_el = rPr.find(qn('a:latin'))
    if latin_el is None:
        latin_el = OxmlElement('a:latin')
        rPr.append(latin_el)
    latin_el.set('typeface', latin)

for slide in prs.slides:
    for sh in slide.shapes:
        if not getattr(sh, 'has_text_frame', False):
            continue
        y = sh.top / EMU
        h = sh.height / EMU
        whole_text = sh.text.strip()
        for p in sh.text_frame.paragraphs:
            for run in p.runs:
                if not run.text:
                    continue
                set_typefaces(run)
                if y < 1.30 and h >= 0.45:
                    run.font.bold = True
                elif len(whole_text.replace(' ', '')) <= 10 and h >= 0.45:
                    run.font.bold = True

prs.save(OUT)
print(OUT)
