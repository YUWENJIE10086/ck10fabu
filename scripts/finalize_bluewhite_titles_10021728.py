from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches

ROOT = Path('决赛发布材料/1002PPT决赛完善版')
SRC = ROOT / '智烤管家_10021714_创客25逐图学习蓝白母版重做版.pptx'
OUT = ROOT / '智烤管家_10021728_创客25逐图学习蓝白标题终版.pptx'

WHITE = RGBColor(0xFF,0xFF,0xFF)
ICE = RGBColor(0xC8,0xEE,0xFF)
SKY = RGBColor(0x79,0xDB,0xFF)
CYAN = RGBColor(0x54,0xF0,0xFF)

prs = Presentation(SRC)


def inch(v):
    return v / 914400.0


def set_shape_color(sh, color):
    if not getattr(sh, 'has_text_frame', False):
        return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            try:
                r.font.color.rgb = color
            except Exception:
                pass


def run_rgb(run):
    try:
        return run.font.color.rgb
    except Exception:
        return None


def run_size(run):
    try:
        return run.font.size.pt if run.font.size else None
    except Exception:
        return None


for si, slide in enumerate(prs.slides, 1):
    # 1) Every page title returns to one consistent blue-white family.
    for sh in slide.shapes:
        if not getattr(sh, 'has_text_frame', False) or not sh.text.strip():
            continue
        txt = sh.text.strip()
        y = inch(sh.top); w = inch(sh.width)

        if y <= 0.32 and w >= 3.0 and txt != '///':
            set_shape_color(sh, CYAN)
        elif 0.32 < y <= 0.86 and w >= 4.0:
            set_shape_color(sh, ICE)

    # 2) Explicit cover / ending title repair: no black fallback, no orange headline.
    if si == 1:
        for sh in slide.shapes:
            if not getattr(sh, 'has_text_frame', False):
                continue
            txt = sh.text.strip()
            if txt == '智烤':
                set_shape_color(sh, WHITE)
            elif txt == '管家':
                set_shape_color(sh, SKY)
            elif txt == 'ZHI KAO MANAGER':
                set_shape_color(sh, SKY)
            elif txt == '一库清底 · 双模决策 · 三智落地 · 四队长效':
                set_shape_color(sh, ICE)
    if si == 15:
        for sh in slide.shapes:
            if not getattr(sh, 'has_text_frame', False):
                continue
            txt = sh.text.strip()
            if txt == '智烤管家':
                set_shape_color(sh, SKY)
            elif txt == '从“有多少房”到“有多少有效烤能”':
                set_shape_color(sh, ICE)

    # 3) Fix every text run that fell back to black after Office/LibreOffice rendering.
    #    Small yellow tags are allowed; everything else stays white/blue.
    for sh in slide.shapes:
        if not getattr(sh, 'has_text_frame', False) or not sh.text.strip():
            continue
        y = inch(sh.top)
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                size = run_size(r)
                rgb = run_rgb(r)
                # Theme/gradient/no-RGB text is the main source of black fallback.
                if rgb is None:
                    r.font.color.rgb = WHITE if y >= 0.9 else ICE
                    continue
                # Actual black/dark text on dark background -> white.
                try:
                    if rgb[0] < 75 and rgb[1] < 75 and rgb[2] < 75:
                        r.font.color.rgb = WHITE
                        continue
                except Exception:
                    pass
                # Green was a leftover accent from older pages; normalize it to blue-white.
                if str(rgb).upper() in {'51E49B','45E8A4','38E6A1'}:
                    r.font.color.rgb = SKY
                # Orange/yellow is retained only for small Creator25-style micro tags.
                if str(rgb).upper() in {'FFC928','F1A63A','FFBE2E','FFD34D'}:
                    if (size or 0) > 12.2 or y < 0.9:
                        r.font.color.rgb = SKY

    # 4) S4 four-card labels: keep them one line, not 3+1-character wrapping.
    if si == 4:
        for sh in slide.shapes:
            if not getattr(sh, 'has_text_frame', False):
                continue
            if sh.text.strip() in {'一库清底','双模决策','三智落地','四队长效'}:
                sh.width = Inches(1.25)
                sh.text_frame.margin_left = Inches(0.01)
                sh.text_frame.margin_right = Inches(0.01)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(15.5)
                        r.font.color.rgb = WHITE

    # 5) S10 role cards and S14 result-card headings: blue-white, not yellow/green.
    if si in (10, 14):
        for sh in slide.shapes:
            if not getattr(sh, 'has_text_frame', False):
                continue
            txt = sh.text.strip()
            if txt in {'智投','智用','智管','烟农 / 技术员','烟站 / 管理人员','合作社'}:
                set_shape_color(sh, SKY)

prs.save(OUT)
print(OUT)
