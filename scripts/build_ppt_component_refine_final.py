from __future__ import annotations

from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import colorsys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SRC = Path('决赛发布材料/1002PPT决赛完善版/智烤管家_10021215_一等奖高阶光效价值深化版.pptx')
OUT_DIR = Path('决赛发布材料/1002PPT决赛完善版')
STAMP = datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%m%d%H%M')
OUT = OUT_DIR / f'智烤管家_{STAMP}_创客组件化克制高级决赛版.pptx'
NOTE = OUT_DIR / f'创客25组件化审美重构说明_{STAMP}.md'

# restrained palette: deep navy + mist blue + warm gold
NAVY = '0B1C2C'
NAVY_2 = '10283A'
MIST = '7EA9C6'
MIST_DARK = '5F829A'
MIST_SOFT = '91AABD'
GOLD = 'D6AE5D'
GOLD_SOFT = 'B99A5B'
WHITE = 'F2F6F8'
TEXT_MUTED = '9EB2C1'

if not SRC.exists():
    raise FileNotFoundError(SRC)

prs = Presentation(SRC)
SW, SH = prs.slide_width, prs.slide_height


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def remove_node(node):
    p = node.getparent()
    if p is not None:
        p.remove(node)


def delete_shape(sh):
    el = sh._element
    p = el.getparent()
    if p is not None:
        p.remove(el)


def hex_to_rgb(h: str):
    h = (h or '').strip().replace('#', '')
    if len(h) != 6:
        return None
    try:
        return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))
    except Exception:
        return None


def rgb_to_hex(vals):
    return ''.join(f'{max(0,min(255,int(round(v)))):02X}' for v in vals)


def soften_color(h: str) -> str:
    """Keep source identity but suppress fluorescent cyan/blue/purple."""
    raw = (h or '').upper()
    direct = {
        '19D7F5': MIST,
        '26D9FF': MIST,
        '02A8FF': MIST_DARK,
        '00FFFF': MIST,
        '00E5FF': MIST,
        '00D9FF': MIST,
        '18D8FF': MIST,
        '00BFFF': MIST_DARK,
    }
    if raw in direct:
        return direct[raw]
    vals = hex_to_rgb(raw)
    if not vals:
        return raw
    r,g,b = [v/255 for v in vals]
    hh,ll,ss = colorsys.rgb_to_hls(r,g,b)
    deg = hh*360
    if 165 <= deg <= 230 and ss > .68 and ll > .35:
        ss = min(ss, .38)
        ll = min(max(ll, .50), .68)
        rr,gg,bb = colorsys.hls_to_rgb(hh,ll,ss)
        return rgb_to_hex((rr*255,gg*255,bb*255))
    if 265 <= deg <= 335 and ss > .62 and ll > .38:
        hh = 220/360
        ss = .20
        ll = min(max(ll,.50),.64)
        rr,gg,bb = colorsys.hls_to_rgb(hh,ll,ss)
        return rgb_to_hex((rr*255,gg*255,bb*255))
    return raw


def strip_effects_and_soften(el):
    # remove the actual sources of the "电竞/霓虹" look
    for tag in ('glow', 'softEdge', 'reflection'):
        for node in list(el.xpath(f'.//a:{tag}')):
            remove_node(node)
    # cyan shadows read as a second luminous outline; neutral shadows are preserved
    for shdw in list(el.xpath('.//a:outerShdw')):
        cols = shdw.xpath('.//a:srgbClr')
        if any(c.get('val') and soften_color(c.get('val')) != c.get('val').upper() for c in cols):
            remove_node(shdw)
    for c in el.xpath('.//a:srgbClr'):
        old = c.get('val')
        if old:
            c.set('val', soften_color(old))
    for ln in el.xpath('.//a:ln'):
        try:
            w = int(ln.get('w') or '0')
        except Exception:
            w = 0
        cols = ln.xpath('.//a:srgbClr')
        is_accent = any(c.get('val') in {MIST, MIST_DARK, MIST_SOFT} for c in cols)
        if is_accent and w > 15875:  # 1.25 pt
            ln.set('w', '12700')


def set_fill(sh, color: str):
    try:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
    except Exception:
        pass


def set_line(sh, color: str, width_pt: float = 1.0):
    try:
        sh.line.color.rgb = rgb(color)
        sh.line.width = Pt(width_pt)
    except Exception:
        pass


def set_text_color(sh, color: str, bold=None):
    if not getattr(sh, 'has_text_frame', False):
        return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            try:
                r.font.color.rgb = rgb(color)
                if bold is not None:
                    r.font.bold = bold
            except Exception:
                pass


def add_rule(slide, x, y, w, h=0.025, color=GOLD_SOFT):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    return s


def shapes_near(slide, x=None, y=None, w=None, h=None, tol=.08):
    out=[]
    for sh in slide.shapes:
        xi, yi, wi, hi = [v/914400 for v in (sh.left, sh.top, sh.width, sh.height)]
        if x is not None and abs(xi-x)>tol: continue
        if y is not None and abs(yi-y)>tol: continue
        if w is not None and abs(wi-w)>tol: continue
        if h is not None and abs(hi-h)>tol: continue
        out.append(sh)
    return out

# 1) remove dynamic scan-light pictures everywhere (the deck contained 8 GIF overlays)
removed_gif = 0
for slide in prs.slides:
    for sh in list(slide.shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                if sh.image.ext.lower() == 'gif':
                    delete_shape(sh)
                    removed_gif += 1
            except Exception:
                pass

# 2) global de-neon pass, while retaining original motherboard gradients and content
for slide in prs.slides:
    strip_effects_and_soften(slide._element)
    # header separator becomes a quiet precision rule, not a laser line
    for sh in slide.shapes:
        yi, hi, wi = sh.top/914400, sh.height/914400, sh.width/914400
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and .62 <= yi <= .72 and hi <= .06 and wi > 7.5:
            set_fill(sh, MIST_DARK)
            try: sh.line.fill.background()
            except Exception: pass

# 3) S3 survey page: replace the triple neon halo by one disciplined data medallion
s3 = prs.slides[2]
ovals = [sh for sh in list(s3.shapes) if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
         and 1.15 < sh.top/914400 < 1.9 and 3.8 < sh.left/914400 < 4.7]
# keep the largest oval, remove the two inner glow layers
if ovals:
    ovals_sorted = sorted(ovals, key=lambda s: s.width*s.height, reverse=True)
    core = ovals_sorted[0]
    for sh in ovals_sorted[1:]:
        delete_shape(sh)
    core.left, core.top, core.width, core.height = Inches(4.03), Inches(1.28), Inches(1.94), Inches(.88)
    set_fill(core, NAVY_2)
    set_line(core, MIST, 1.0)
# 531 is the evidence, not the glow
for sh in s3.shapes:
    if getattr(sh, 'has_text_frame', False):
        t = sh.text.strip()
        if t == '531':
            set_text_color(sh, GOLD, True)
        elif t == '份有效问卷':
            set_text_color(sh, TEXT_MUTED, False)

# 4) S4 1234: keep the competition-grade 4-column component, refine its visual grammar
s4 = prs.slides[3]
for sh in s4.shapes:
    if getattr(sh, 'has_text_frame', False):
        t = sh.text.strip()
        if t in {'1','2','3','4'}:
            set_text_color(sh, GOLD, True)
        elif t in {'DATA','MODEL','SERVICE','MECHANISM'}:
            set_text_color(sh, MIST_SOFT, True)
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        xi, yi, wi, hi = [v/914400 for v in (sh.left, sh.top, sh.width, sh.height)]
        # four major cards
        if 1.5 < hi < 3.5 and 1.7 < wi < 2.1:
            set_line(sh, MIST_DARK, 1.0)
        # arrows and bottom status dots become restrained
        if .35 < wi < .55 and .35 < hi < .55:
            set_fill(sh, MIST_DARK); set_line(sh, MIST_DARK, .6)
        if .30 < wi < .45 and hi < .20 and yi > 4.2:
            set_fill(sh, GOLD_SOFT); set_line(sh, GOLD_SOFT, .4)
# tiny gold cap line gives the four columns one consistent component family
for x in (.90, 3.26, 5.62, 7.98):
    add_rule(s4, x, 1.67, .36, .025)

# 5) S6/S7 model pages: three-step evidence cards, with gold step index + quiet image frame
for idx in (5,6):
    slide = prs.slides[idx]
    for sh in slide.shapes:
        xi, yi, wi, hi = [v/914400 for v in (sh.left, sh.top, sh.width, sh.height)]
        if getattr(sh, 'has_text_frame', False):
            t = sh.text.strip()
            if t in {'01','02','03'}:
                set_text_color(sh, GOLD, True)
            if '技术高明之处' in t:
                set_text_color(sh, WHITE, True)
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            # major model cards
            if 2.65 < wi < 3.7 and 3.2 < hi < 3.7:
                set_line(sh, MIST_DARK, 1.0)
            # inner screenshot/evidence windows
            if 2.3 < wi < 3.4 and 1.45 < hi < 1.7:
                set_line(sh, MIST_SOFT, .8)
    # source images carry credibility, gold cap is enough for hierarchy
    starts = (.66, 3.84, 7.02) if idx == 5 else (.66, 3.56, 6.46)
    for x in starts:
        add_rule(slide, x, 1.47, .34, .025)

# 6) S8 coupling chain: dark badges + gold indices, steel-blue connectors
s8 = prs.slides[7]
for sh in s8.shapes:
    xi, yi, wi, hi = [v/914400 for v in (sh.left, sh.top, sh.width, sh.height)]
    if getattr(sh, 'has_text_frame', False) and sh.text.strip() in {'01','02','03','04'}:
        set_text_color(sh, GOLD, True)
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        if .65 < wi < .78 and .65 < hi < .78 and 1.45 < yi < 1.8:
            set_fill(sh, NAVY_2); set_line(sh, GOLD_SOFT, 1.0)
        elif .40 < wi < .55 and .40 < hi < .55 and 2.8 < yi < 3.2:
            set_fill(sh, MIST_DARK); set_line(sh, MIST_DARK, .6)
        elif 1.7 < wi < 2.0 and 1.5 < hi < 1.9:
            set_line(sh, MIST_DARK, 1.0)

# 7) S10 three roles: same card family, different role headline; left hairline instead of luminous frame
s10 = prs.slides[9]
for sh in s10.shapes:
    xi, yi, wi, hi = [v/914400 for v in (sh.left, sh.top, sh.width, sh.height)]
    if getattr(sh, 'has_text_frame', False):
        t = sh.text.strip()
        if t in {'智投','智用','智管'}:
            set_text_color(sh, GOLD, True)
        elif t in {'管理人员','烟农 + 技术员','合作社'}:
            set_text_color(sh, MIST_SOFT, True)
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 2.6 < wi < 2.9 and 3.1 < hi < 3.5:
        set_line(sh, MIST_DARK, 1.0)
for x in (.70, 3.66, 6.62):
    add_rule(s10, x, 1.76, .028, 2.62, GOLD_SOFT)

# 8) S14 KPI page: numbers are the memory anchors, frames recede
s14 = prs.slides[13]
KPI = {'48','18','96.8%','2442','340+','29','1363','81%','≈190万'}
for sh in s14.shapes:
    xi, yi, wi, hi = [v/914400 for v in (sh.left, sh.top, sh.width, sh.height)]
    if getattr(sh, 'has_text_frame', False):
        t = sh.text.strip()
        if t in KPI:
            set_text_color(sh, GOLD, True)
        elif t in {'烟农 / 技术员','烟站 / 管理人员','合作社'}:
            set_text_color(sh, WHITE, True)
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 2.5 < wi < 2.75 and 3.8 < hi < 4.2:
        set_line(sh, MIST_DARK, 1.0)
for x in (.68, 3.88, 7.08):
    add_rule(s14, x, 1.45, .42, .025)

# 9) S15 close: product mark + 4 restrained pills, no stage-neon frame
s15 = prs.slides[14]
for sh in s15.shapes:
    xi, yi, wi, hi = [v/914400 for v in (sh.left, sh.top, sh.width, sh.height)]
    if getattr(sh, 'has_text_frame', False):
        t = sh.text.strip()
        if t in {'1','2','3','4'}:
            set_text_color(sh, GOLD, True)
        elif t == 'ZHI KAO MANAGER':
            set_text_color(sh, MIST_SOFT, True)
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 1.2 < wi < 1.4 and .65 < hi < .9 and 3.3 < yi < 3.7:
        set_line(sh, MIST_DARK, .9)
add_rule(s15, 4.27, 1.40, 1.46, .025, GOLD_SOFT)

# 10) remove residual luminous effects introduced by source internals; leave native gradients otherwise intact
for slide in prs.slides:
    strip_effects_and_soften(slide._element)

OUT_DIR.mkdir(parents=True, exist_ok=True)
prs.save(OUT)

NOTE.write_text(f'''# 智烤管家决赛PPT｜创客25组件化克制高级重构（{STAMP}）

本版不是继续“加灯带”，而是以 `智烤管家_10021215_一等奖高阶光效价值深化版.pptx` 为内容基底，对 **组件语法** 做二次重构。

## 本次直接处理
- 删除动态扫光 GIF：**{removed_gif} 个**。
- 删除全套 Glow / SoftEdge / Reflection；青色外阴影一并去除。
- 电光青 / 荧光蓝压成雾蓝、钢蓝；保留母版深蓝与白→青标题渐变。
- 所有强调线改为 1pt 左右的“精密线”，不再做电竞灯带。

## 按“创客25图片”真正落到组件，而不是只换颜色
1. **C8 舞台极简**：封面、收尾、转场不再用大扫光框；用标题、留白、少量金线形成舞台记忆。
2. **C4 三步技术链**：模型①、模型②保持 01/02/03 三步卡，但把“发光框”改成“证据窗 + 步骤索引 + 技术输出”。
3. **C2/C6 证据卡**：真实截图继续作为证据主体，边缘只留细线；截图不再被霓虹装饰抢戏。
4. **C9 模块化关系**：1234 与双模耦合页保留模块链，序号改暖金、箭头改钢蓝，层级更清楚。
5. **C10 蓝青组件**：只留下深蓝玻璃感、雾蓝精密线、白青标题、暖金记忆点四种核心视觉语汇。
6. **成效大数字**：48、96.8%、2442、1363、≈190万等 KPI 统一成为暖金视觉锚点，卡片退后。
7. **调查证据**：531 页移除三层荧光光环，只保留一个克制的数据徽章。
8. **三智角色卡**：智投 / 智用 / 智管统一为三张角色工作流卡，用左侧细金线做识别，不再包一圈光。

## 保留不动
- 演讲稿主线和 15 页结构；
- 双模逻辑、1234闭环、三智/四队价值表达；
- 真实小程序、后台、模型和看板截图；
- 母版的深蓝底、蓝白渐变标题体系。

输出：`{OUT.name}`
''', encoding='utf-8')

print(f'OUTPUT={OUT.as_posix()}')
print(f'NOTE={NOTE.as_posix()}')
