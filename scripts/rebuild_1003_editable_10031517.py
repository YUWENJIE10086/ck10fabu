from pptx import Presentation
from pptx.util import Inches
from pathlib import Path
from PIL import Image
import numpy as np, cv2, io

STAMP='10031517'
ROOT=Path('.')
BASE=ROOT/'决赛发布材料/1002PPT决赛完善版/1003'
SRC_A=BASE/'智烤管家_142106_1比1蓝白山海光路抠图还原版_副本3.pptx'
SRC_B=BASE/'智烤管家_142106_1比1蓝白山海光路抠图还原版_副本2.pptx'
CLEAN_BG=BASE/'_assets'/f'background_clean_no_road_{STAMP}.jpg'
OUT=ROOT/f'智烤管家_{STAMP}_1003原图1比1文字图标分层可删除无背景光路版.pptx'
TMP=Path('/tmp')/f'zhikao1003_assets_{STAMP}'; TMP.mkdir(exist_ok=True)

ROIS={
1:[(2.0,.25,9.4,4.5)],2:[(.55,1.55,1.75,4.65),(7.55,1.65,1.70,1.35),(7.55,3.10,1.70,1.35),(7.55,4.55,1.70,1.35)],
3:[(2.55,1.45,8.35,4.55)],4:[(.45,1.60,3.05,4.20),(3.45,1.55,3.00,4.25),(6.45,1.55,3.00,4.25),(9.45,1.55,3.30,4.25)],
5:[(.30,3.55,8.15,2.55),(8.10,3.20,4.80,3.05)],6:[(.25,2.15,4.25,3.75),(4.30,2.10,4.45,3.75),(8.70,2.15,4.25,3.75)],
7:[(.25,2.00,6.35,4.10),(6.10,1.95,6.75,4.20)],8:[(.30,1.85,4.95,4.05),(4.90,2.25,3.55,3.75),(8.10,1.85,4.95,4.05)],
9:[(.30,2.15,4.10,3.95),(4.40,2.15,4.10,3.95),(8.50,2.15,4.45,3.95)],10:[(.15,2.10,4.25,4.00),(3.95,1.95,5.30,4.25),(9.00,2.10,4.20,4.00)],
11:[(.15,1.95,3.25,3.95),(3.20,1.95,3.35,3.95),(6.45,1.95,3.35,3.95),(9.65,1.95,3.55,3.95)],
12:[(.20,1.65,3.25,4.55),(3.20,1.65,3.20,4.55),(6.35,1.65,3.20,4.55),(9.45,1.65,3.55,4.55)],
13:[(.45,1.80,2.85,4.20),(3.30,1.80,2.85,4.20),(6.35,1.80,2.85,4.20),(9.30,1.80,3.00,4.20)],
14:[(.45,1.85,2.60,2.15),(3.35,1.85,2.55,2.15),(6.45,1.85,2.65,2.15),(9.55,1.85,2.75,2.15),(3.00,4.25,3.15,1.65),(7.10,4.25,3.15,1.65)],
15:[(.35,1.75,3.00,3.85),(3.40,1.75,3.00,3.85),(6.45,1.75,3.00,3.85),(9.50,1.75,3.20,3.85)]}

def alpha_from_crop(rgb,mode='art'):
    a=np.array(rgb).astype(np.int16); r,g,b=a[:,:,0],a[:,:,1],a[:,:,2]; v=np.max(a,axis=2); cyan=((g+b)//2-r)
    al=np.maximum((v-(52 if mode=='road' else 78))*(3.2 if mode=='road' else 3.0),(cyan-(12 if mode=='road' else 27))*(5.0 if mode=='road' else 4.5))
    al=np.clip(al,0,255).astype(np.uint8); al[v<58]=0
    return cv2.GaussianBlur(al,(3,3),0)

sources=[]
for p in [SRC_A,SRC_B]:
    r=Presentation(p); sources += [(r,s) for s in r.slides]
assert len(sources)==15
prs=Presentation(); prs.slide_width=Inches(13.3326); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
bgim=Image.open(CLEAN_BG).convert('RGB').resize((1600,900),Image.Resampling.LANCZOS); bgpath=TMP/'background.jpg'; bgim.save(bgpath,quality=95)
for gi,(srcprs,slide) in enumerate(sources,1):
    dst=prs.slides.add_slide(blank)
    p=dst.shapes.add_picture(str(bgpath),0,0,width=prs.slide_width,height=prs.slide_height); p.name='背景_无字无图标无底部光路'
    base=Image.open(io.BytesIO(slide.shapes[0].image.blob)).convert('RGB'); W,H=base.size; sx=W/(srcprs.slide_width/914400); sy=H/(srcprs.slide_height/914400)
    textmask=np.zeros((H,W),np.uint8); orig=[]
    for si,sh in enumerate(slide.shapes):
        if si and sh.shape_type==13:
            x=int(sh.left/914400*sx); y=int(sh.top/914400*sy); w=int(sh.width/914400*sx); h=int(sh.height/914400*sy)
            cv2.rectangle(textmask,(max(0,x-5),max(0,y-3)),(min(W-1,x+w+5),min(H-1,y+h+3)),255,-1); orig.append(sh)
    roadcut=int(5.45*sy)
    for ri,(x,y,w,h) in enumerate(ROIS[gi],1):
        x0=max(0,int(x*sx)); y0=max(0,int(y*sy)); x1=min(W,int((x+w)*sx)); y1=min(H,int((y+h)*sy)); crop=base.crop((x0,y0,x1,y1)); alpha=alpha_from_crop(crop)
        alpha[textmask[y0:y1,x0:x1]>0]=0
        if y1>roadcut: alpha[max(0,roadcut-y0):,:]=0
        if int((alpha>18).sum())<70: continue
        ap=TMP/f's{gi:02d}_visual_{ri:02d}.png'; Image.fromarray(np.dstack([np.array(crop),alpha])).save(ap)
        q=dst.shapes.add_picture(str(ap),Inches(x0/sx),Inches(y0/sy),width=Inches((x1-x0)/sx),height=Inches((y1-y0)/sy)); q.name=f'图标与图形组件_S{gi:02d}_{ri:02d}_可移动删除'
    ry0=int(5.25*sy); road=base.crop((0,ry0,W,H)); ra=np.array(road); aa=alpha_from_crop(road,'road'); aa[textmask[ry0:H,:]>0]=0
    hsv=cv2.cvtColor(ra,cv2.COLOR_RGB2HSV); strong=((hsv[:,:,2]>92)&(hsv[:,:,1]>40))|(hsv[:,:,2]>160); aa[~strong]=0
    if int((aa>15).sum())>100:
        rp=TMP/f's{gi:02d}_road.png'; Image.fromarray(np.dstack([ra,aa])).save(rp); q=dst.shapes.add_picture(str(rp),0,Inches(ry0/sy),width=prs.slide_width,height=Inches((H-ry0)/sy)); q.name=f'底部光路_S{gi:02d}_独立可删除'
    for ci,sh in enumerate(orig,1):
        ip=TMP/f's{gi:02d}_existing_{ci:03d}.{sh.image.ext or "png"}'; ip.write_bytes(sh.image.blob)
        q=dst.shapes.add_picture(str(ip),sh.left,sh.top,width=sh.width,height=sh.height); q.rotation=sh.rotation; q.name=f'文字与细节组件_S{gi:02d}_{ci:03d}_可移动删除'
prs.save(OUT)
print(OUT)
