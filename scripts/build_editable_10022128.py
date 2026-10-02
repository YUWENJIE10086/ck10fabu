from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from docx import Document
from docx.shared import Pt as DPt, Inches as DInches, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from pathlib import Path
import math, zipfile

STAMP='10022128'
OUT=Path('.')
PPT_A=OUT/f'智烤管家_{STAMP}_10021912原视觉全组件可编辑还原版.pptx'
PPT_B=OUT/f'智烤管家_{STAMP}_全组件可编辑减光路8分钟优化版.pptx'
DOCX=OUT/f'智烤管家_{STAMP}_8分钟逐页演讲建议与素材补充说明.docx'
NAVY='06182D'; NAVY2='092542'; NAVY3='0E3152'; BLUE='1888D8'; CYAN='23C8F6'; ICE='9EEBFF'; WHITE='F7FBFF'; MUTED='A7C5D8'; GOLD='F5D071'; RED='FF715B'; GREEN='55E6A5'
TITLE_FONT='Noto Serif CJK SC'; BODY_FONT='Noto Sans CJK SC'; W=13.333; H=7.5

def rgb(h):
    h=h.lstrip('#'); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))

def add_text(slide,text,x,y,w,h,size=18,color=WHITE,bold=False,font=BODY_FONT,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.MIDDLE):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    tf.margin_left=tf.margin_right=Inches(.02); tf.margin_top=tf.margin_bottom=Inches(.01); return sh

def line(slide,x1,y1,x2,y2,color=CYAN,width=1.5):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); sh.line.color.rgb=rgb(color); sh.line.width=Pt(width); return sh

def circle(slide,x,y,d,fill=CYAN,linec=None,lw=.8):
    sh=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if linec: sh.line.color.rgb=rgb(linec); sh.line.width=Pt(lw)
    else: sh.line.fill.background()
    return sh

def rect(slide,x,y,w,h,fill=NAVY2,linec=None,lw=1,radius=True):
    typ=MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE; sh=slide.shapes.add_shape(typ,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if linec: sh.line.color.rgb=rgb(linec); sh.line.width=Pt(lw)
    else: sh.line.fill.background()
    return sh

def add_title(slide,title,subtitle=None):
    add_text(slide,title,.82,.38,11.7,.68,30,CYAN,True,TITLE_FONT,PP_ALIGN.CENTER); add_text(slide,title,.82,.34,11.7,.68,30,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER)
    line(slide,2.4,1.05,4.05,1.05,CYAN,1.1); line(slide,9.28,1.05,10.93,1.05,CYAN,1.1)
    if subtitle: add_text(slide,subtitle,1.4,1.04,10.5,.38,13,MUTED,False,BODY_FONT,PP_ALIGN.CENTER)

def add_bg(slide,road=True,seed=1):
    bg=rect(slide,0,0,W,H,NAVY,None,0,False); bg.name='可编辑_深蓝背景'
    for i in range(28):
        x=((i*47+seed*29)%1250)/100+.35; y=((i*31+seed*17)%190)/100+.20; d=.018+(i%3)*.008; circle(slide,min(x,12.9),min(y,2.1),d,ICE)
    for k,(y,fillc) in enumerate([(5.22,NAVY3),(5.65,NAVY2),(6.02,'071F38')]):
        for i in range(9):
            x=i*1.65-.8; w=2.35; h=1.0+.22*((i+k)%3); sh=slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x),Inches(y-h),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fillc); sh.line.fill.background()
    line(slide,0,6.25,13.33,6.25,'0D4C72',1.2)
    if road:
        pts=[(0,6.88),(1.2,6.82),(2.15,6.58),(3.35,6.72),(4.65,6.47),(5.9,6.60),(7.15,6.38),(8.45,6.52),(9.55,6.28),(10.8,6.42),(12.15,6.18),(13.33,6.05)]
        for off,col,wid in [(.08,'0C4F78',7),(.03,BLUE,3.2),(0,CYAN,1.15)]:
            for (a,b),(c,d) in zip(pts[:-1],pts[1:]): line(slide,a,b+off,c,d+off,col,wid)
        for x,y in pts[1:-1:2]: circle(slide,x-.05,y-.05,.10,WHITE); circle(slide,x-.027,y-.027,.055,CYAN)

def icon_home(slide,cx,cy,s=.55):
    rect(slide,cx-s/2,cy-s*.12,s,s*.45,'0A4164',CYAN,1,False); roof=slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(cx-s*.58),Inches(cy-s*.50),Inches(s*1.16),Inches(s*.52)); roof.fill.solid(); roof.fill.fore_color.rgb=rgb(ICE); roof.line.color.rgb=rgb(CYAN); rect(slide,cx-s*.10,cy+s*.05,s*.20,s*.28,NAVY,CYAN,.6,False)

def icon_chart(slide,cx,cy,s=.58):
    for i,h in enumerate([.25,.42,.62]): rect(slide,cx-s*.42+i*s*.27,cy+s*.25-h*s,s*.17,h*s,ICE,CYAN,.5,False)
    line(slide,cx-s*.46,cy+s*.29,cx+s*.48,cy+s*.29,CYAN,1)

def icon_gear(slide,cx,cy,s=.58):
    circle(slide,cx-s*.37,cy-s*.37,s*.74,'0C4A70',CYAN,1); circle(slide,cx-s*.14,cy-s*.14,s*.28,NAVY,ICE,1)
    for a in range(0,360,45):
        rad=math.radians(a); x=cx+math.cos(rad)*s*.48; y=cy+math.sin(rad)*s*.48; rect(slide,x-s*.07,y-s*.07,s*.14,s*.14,ICE,None,0,False)

def icon_doc(slide,cx,cy,s=.58):
    rect(slide,cx-s*.35,cy-s*.42,s*.70,s*.84,'0A4164',ICE,1,False)
    for i in range(3): line(slide,cx-s*.22,cy-s*.18+i*s*.18,cx+s*.22,cy-s*.18+i*s*.18,ICE,1)

def node(slide,n,title,sub,x,y,icon='home'):
    add_text(slide,f'{n:02d}',x,y,.8,.48,28,ICE,True,'DejaVu Serif');
    {'home':icon_home,'chart':icon_chart,'gear':icon_gear,'doc':icon_doc}[icon](slide,x+1.15,y+.24,.55)
    add_text(slide,title,x+1.65,y,2,.48,19,WHITE,True,TITLE_FONT); add_text(slide,sub,x+1.65,y+.48,2.35,.55,11,MUTED)

def arrow(slide,x,y): add_text(slide,'››',x,y,.6,.42,25,CYAN,True,'DejaVu Sans',PP_ALIGN.CENTER)

def kpi(slide,num,label,x,y,w=2.25,accent=CYAN):
    add_text(slide,num,x,y,w,.72,34,WHITE,True,'DejaVu Serif',PP_ALIGN.CENTER); line(slide,x+w*.22,y+.77,x+w*.78,y+.77,accent,1.2); add_text(slide,label,x,y+.82,w,.46,12,MUTED,True,BODY_FONT,PP_ALIGN.CENTER)

def mini_bars(slide,x,y,w,h,vals,colors=None):
    colors=colors or [CYAN]*len(vals); bw=w/(len(vals)*1.55)
    for i,v in enumerate(vals):
        bh=h*v/max(vals); rect(slide,x+i*bw*1.55,y+h-bh,bw,bh,colors[i%len(colors)],None,0,False)
    line(slide,x,y+h,x+w,y+h,'3D7191',.8)

def radar(slide,cx,cy,r,vals):
    n=len(vals)
    for ring in [.33,.66,1.0]:
        pts=[]
        for i in range(n):
            a=-math.pi/2+2*math.pi*i/n; pts.append((cx+math.cos(a)*r*ring,cy+math.sin(a)*r*ring))
        for p,q in zip(pts,pts[1:]+pts[:1]): line(slide,*p,*q,'2E789F',.7)
    for i in range(n):
        a=-math.pi/2+2*math.pi*i/n; line(slide,cx,cy,cx+math.cos(a)*r,cy+math.sin(a)*r,'2E789F',.6)
    pts=[]
    for i,v in enumerate(vals):
        a=-math.pi/2+2*math.pi*i/n; pts.append((cx+math.cos(a)*r*v,cy+math.sin(a)*r*v))
    for p,q in zip(pts,pts[1:]+pts[:1]): line(slide,*p,*q,CYAN,1.8)
    for p in pts: circle(slide,p[0]-.035,p[1]-.035,.07,WHITE)

def build_deck(path,reduced=False):
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H); blank=prs.slide_layouts[6]
    def ns(i,road=True):
        s=prs.slides.add_slide(blank); add_bg(s,road=(road if not reduced else i<=2),seed=i); return s
    s=ns(1,True); add_text(s,'Z H I   K A O   M A N A G E R',3.7,1,6,.35,13,ICE,False,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'智烤',3,1.72,3.1,1.15,44,WHITE,True,TITLE_FONT,PP_ALIGN.RIGHT); add_text(s,'管家',6.05,1.72,3.2,1.15,44,CYAN,True,TITLE_FONT); add_text(s,'让每一座烤房，都有一份可计算的“健康档案”',2,3,9.3,.5,16,WHITE,False,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'一库清底  ·  双模决策  ·  三智落地  ·  四队长效',2.2,4.25,8.9,.5,17,ICE,True,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'基于健康评估的烤房用投管一体化解决方案',2.1,5.05,9.1,.48,15,WHITE,True,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'湖北烟草创客大赛 · 决赛发布    宜昌市烟草公司 · 智同道合创客小组    2026.10',1.6,6.82,10.2,.32,10,MUTED,False,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(2,True); add_title(s,'烟区一线，最怕这 3 件事','不是“有没有烤房”，而是分配、抢修、投入仍靠经验和人工')
    for j,(t,sub) in enumerate([('分不到','旺季状态不透明｜烟农无法自主选择'),('修不及','故障逐级上报｜找人找件慢'),('投不准','哪栋该修｜哪个村真缺｜新建难量化')]):
        y=1.75+j*1.23; add_text(s,f'0{j+1}',.75,y,.72,.5,25,ICE,True,'DejaVu Serif'); add_text(s,t,1.65,y,2.3,.55,28,WHITE,True,TITLE_FONT); add_text(s,sub,1.65,y+.55,3.8,.45,12,MUTED); arrow(s,5.1,y+.12); node(s,j+1,['资源调度','抢修服务','投入决策'][j],['状态可见 / 一键调度','故障定位 / 工单闭环','有效烤能 / 缺口测算'][j],7.1,y-.02,['home','gear','chart'][j])
    add_text(s,'根因：资源无统筹 · 投入无量化 · 管护无闭环',2.25,6.5,8.8,.5,18,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(3); add_title(s,'三愁背后，是三条链断了','资源、维修、投入没有形成数据联动闭环'); centers=[(6.65,3.3,'home','用房管理断链','空闲不清 / 状态不明 / 调度靠经验'),(3.55,5.15,'gear','维修保障断链','故障发现慢 / 找人找件慢 / 闭环追踪弱'),(9.75,5.15,'chart','投入决策断链','真实需求不清 / 供需错配 / 新改建难量化')]
    for idx,(cx,cy,ic,t,sub) in enumerate(centers,1):
        circle(s,cx-.58,cy-.58,1.16,'0B3557',CYAN,1.5); {'home':icon_home,'gear':icon_gear,'chart':icon_chart}[ic](s,cx,cy,.62); add_text(s,f'0{idx}  {t}',cx-1.5,cy+.72,3,.45,17,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER); add_text(s,sub,cx-1.8,cy+1.16,3.6,.55,11,MUTED,False,BODY_FONT,PP_ALIGN.CENTER)
    for p,q in [((6.65,3.85),(4.25,4.85)),((6.65,3.85),(9.05,4.85)),((4.2,5.05),(9.1,5.05))]: line(s,*p,*q,CYAN,1.3)
    add_text(s,'问题不在“有没有房”，而在“用—投—管”没有形成一条数据链',2,6.55,9.3,.48,16,ICE,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(4); add_title(s,'一套方案，做成 4 件事','从摸清家底，到科学决策，再到真实落地与长效机制')
    xs=[.75,3.75,6.75,9.75]; titles=['一库清底','双模决策','三智落地','四队长效']; subs=['一房一档 / 状态可见 / 家底清楚','健康评估 / 供需推演 / 投入更准','智用 / 智投 / 智管','专班推进 / 维修服务 / 运营管护 / 基金保障']; icons=['home','chart','gear','doc']
    for i,x in enumerate(xs):
        node(s,i+1,titles[i],subs[i],x,2.3,icons[i]);
        if i<3: arrow(s,x+2.55,3.45)
    add_text(s,'把“看不清、修不及、投不准”，变成“看得见、管得住、投得准”',1.8,6.3,9.8,.52,18,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(5); add_title(s,'一库清底：先把家底摸清楚','4大模块 · 18个关键部件 · 一房一档 · 动态更新')
    for i,(t,sub,ic) in enumerate([('现场采集','位置 / 容量 / 状态 / 设备 / 使用情况','doc'),('标准判定','正常 / 损坏 / 缺失 · 优良 / 一般 / 较差','gear'),('一房一档','基础信息 / 维修记录 / 用户评价','home')]):
        x=1+i*4; node(s,i+1,t,sub,x,1.85,ic); 
        if i<2: arrow(s,x+3,2.65)
    add_text(s,'真实系统证据位：替换为“一房一档”后台/小程序截图',7.55,4.15,4.9,.42,11,GOLD,True,BODY_FONT,PP_ALIGN.CENTER); rect(s,8,4.62,4,1.1,'0A2D4C',CYAN,1,True); add_text(s,'一房一档｜基础信息｜部件状态｜维修记录｜评价',8.2,4.83,3.6,.56,12,WHITE,True,BODY_FONT,PP_ALIGN.CENTER)
    for x,y,c in [(2,4.75,GREEN),(2.7,5.1,CYAN),(3.35,4.65,RED),(4.05,5.2,GREEN),(4.85,4.85,CYAN),(5.45,5.35,GOLD)]: circle(s,x,y,.16,c)
    line(s,2,5.8,6.2,5,'16567B',1); add_text(s,'结果：空闲房在哪、故障房是谁、可用能力有多少，一眼看清',1.5,6.35,10.3,.5,17,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(6); add_title(s,'模型① 烤房健康评估：从打分到估寿','不是给一座烤房一个分数，而是判断还能用多久、哪里先修'); node(s,1,'状态标准化','4大模块 / 18部件 / 指标归一',.65,1.6,'doc'); node(s,2,'AHP组合评分','部件状态80% / 管护10% / 现场10%',4.6,1.6,'chart'); node(s,3,'寿命与风险','60分预警线 / 风险部位 / 维修优先序',8.55,1.6,'chart'); arrow(s,3.75,2.15); arrow(s,7.7,2.15); icon_home(s,2.45,4.55,1.25); radar(s,6.65,4.45,1.15,[.72,.55,.82,.48,.68,.62]); pts=[(9.3+i*.28,4+2*(i/10)**1.8) for i in range(11)]
    for a,b in zip(pts[:-1],pts[1:]): line(s,*a,*b,CYAN,2)
    add_text(s,'60分以下自动预警',9.1,5.92,3.4,.4,12,GOLD,True,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'真正输出：健康等级 · 风险部位 · 剩余寿命趋势 · 维修优先序',1.4,6.48,10.5,.5,16,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(7); add_title(s,'模型② 千亩村发展导引：把真实缺口算出来','不问“想建几栋”，先算有效烤能，再判断结构性缺口'); node(s,1,'算有效烤能','可用烤房 / 健康折损 / 产能折算',.65,1.62,'gear'); node(s,2,'判真实缺口','种植规模 / 产业趋势 / 管护能力',4.6,1.62,'chart'); node(s,3,'做建设推演','熵权法 + TOPSIS / 地图情景沙盘',8.55,1.62,'home'); arrow(s,3.75,2.15); arrow(s,7.7,2.15)
    for x,y,c in [(1.6,4.4,GREEN),(2.2,5,CYAN),(3,4.55,GOLD),(3.55,5.1,RED),(4.1,4.75,GREEN)]: circle(s,x,y,.18,c)
    line(s,1.3,5.6,4.5,4.2,'1A638B',1); mini_bars(s,5.2,4.1,3.3,1.7,[3,5,7,8,9,11],[BLUE,CYAN]); add_text(s,'需求 vs 有效烤能',5.55,5.85,2.6,.35,11,MUTED,True,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'先调剂 → 再修复 → 仍有结构性缺口才新建',8.9,4.2,3.7,.5,14,WHITE,True,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'把“有多少栋”改成“有多少有效能力”，建设才有依据',1.6,6.4,10.1,.48,17,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(8); add_title(s,'双模决策：存量评估 + 增量推演','模型①把存量算清，模型②才有资格谈增量'); add_text(s,'模型 ①  健康评估',1.3,1.7,4.2,.5,21,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER); add_text(s,'模型 ②  供需推演',7.85,1.7,4.2,.5,21,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER)
    for x,t,ic in [(1.4,'部件体检','gear'),(3,'寿命预测','chart'),(4.6,'维修优先序','doc'),(7.9,'有效烤能','home'),(9.5,'真实缺口','chart'),(11.1,'建设情景','doc')]:
        circle(s,x-.32,3,.64,'0A4265',CYAN,1); {'gear':icon_gear,'chart':icon_chart,'home':icon_home,'doc':icon_doc}[ic](s,x,3.32,.38); add_text(s,t,x-.7,3.72,1.4,.42,11,WHITE,True,BODY_FONT,PP_ALIGN.CENTER)
    circle(s,5.67,4.25,1.95,'0A3558',CYAN,1.8); add_text(s,'双模耦合\n输出',5.82,4.63,1.65,.7,17,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER)
    for x,t in [(4.2,'该修哪栋'),(6.05,'该补哪里'),(7.9,'该建多少')]: circle(s,x,6,.38,'0A4265',CYAN,1); add_text(s,t,x-.5,6.37,1.38,.35,11,WHITE,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(9); add_title(s,'三智落地：把系统能力变成日常动作','不是做一个平台，而是让智用、智投、智管真正进入业务现场')
    for i,(t,subs,ic) in enumerate([('智用',['空闲可见','一键调度','故障报修'],'home'),('智投',['需求识别','缺口测算','投入更准'],'chart'),('智管',['运行监测','维修留痕','预警闭环'],'gear')]):
        x=.9+i*4.15; circle(s,x+1.15,2.3,1.3,'0A3558',CYAN,1.4); {'home':icon_home,'chart':icon_chart,'gear':icon_gear}[ic](s,x+1.8,2.95,.72); add_text(s,t,x+.35,3.85,2.9,.5,24,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER)
        for j,sub in enumerate(subs): add_text(s,'○  '+sub,x+.55,4.45+j*.43,2.6,.36,12,MUTED)
    add_text(s,'让资源调得动、投入算得清、运行管得住',2,6.35,9.3,.5,17,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(10); add_title(s,'智用：调度与抢修，真正落到烟农手里','从“打电话找房、找人、找配件”变成一键发起、过程可见、结果回写'); items=[('找空闲房','状态可见 / 就近可选','home'),('发起调度','线上申请 / 一键提交','doc'),('故障报修','定位问题 / 自动派单','gear'),('评价回写','过程留痕 / 服务闭环','doc')]
    for i,(t,sub,ic) in enumerate(items):
        x=.45+i*3.15; node(s,i+1,t,sub,x,1.7,ic); 
        if i<3: arrow(s,x+2.65,2.35)
    rect(s,4.15,4.15,2.05,1.35,'0A2C4A',CYAN,1,True); add_text(s,'真实小程序\n调度申请截图',4.35,4.42,1.65,.75,12,GOLD,True,BODY_FONT,PP_ALIGN.CENTER); rect(s,9.75,4.15,2.05,1.35,'0A2C4A',CYAN,1,True); add_text(s,'真实工单/评价\n截图',9.95,4.42,1.65,.75,12,GOLD,True,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'磨坪站雷击案例：精准诊断2处故障部件，3小时内完成抢修复产',2,6.3,9.3,.52,16,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(11); add_title(s,'智投：把钱投到最该投的地方','存量先体检，增量再推演，选址和投入都用数据说话'); items=[('存量体检','健康等级 / 寿命 / 风险','gear'),('缺口识别','有效烤能 / 需求对比','chart'),('选址推演','村组位置 / 服务半径','home'),('效果复盘','投入产出 / 使用率','chart')]
    for i,(t,sub,ic) in enumerate(items):
        x=.4+i*3.18; node(s,i+1,t,sub,x,1.65,ic); 
        if i<3: arrow(s,x+2.65,2.3)
    mini_bars(s,.8,4.2,2.4,1.55,[4,7,5,9],[BLUE,CYAN,GREEN,GOLD]); add_text(s,'健康分布',1.05,5.8,1.9,.35,11,MUTED,True,BODY_FONT,PP_ALIGN.CENTER); mini_bars(s,3.75,4.2,2.4,1.55,[4,5,8,10,12],[BLUE,CYAN]); add_text(s,'供需缺口',4.05,5.8,1.9,.35,11,MUTED,True,BODY_FONT,PP_ALIGN.CENTER)
    for p in [(7.2,4.4,GREEN),(7.9,5.1,CYAN),(8.5,4.55,GOLD),(9.1,5,RED)]: circle(s,p[0],p[1],.16,p[2])
    kpi(s,'150栋','核心设备升级计划',10.15,4.05,2.35); kpi(s,'7栋','肖家坪科学新建',10.15,5.15,2.35,GOLD); add_text(s,'先调剂、再维修、最后才新建',3.7,6.42,5.9,.42,17,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(12); add_title(s,'智管：让每一座烤房都有可追溯的生命周期','建档、巡检、维修、评价不断档，数据持续回写'); items=[('建档','一房一档 / 责任到人','doc'),('巡检','状态更新 / 风险发现','chart'),('维修','工单闭环 / 配件留痕','gear'),('回写','评价 / 成本 / 历史沉淀','doc')]; coords=[(2.1,2.55),(8.95,2.55),(8.95,5.05),(2.1,5.05)]
    for i,((t,sub,ic),(x,y)) in enumerate(zip(items,coords),1): node(s,i,t,sub,x-.8,y-.35,ic)
    line(s,4.2,2.75,8.75,2.75,CYAN,2); line(s,10.6,3.2,10.6,4.75,CYAN,2); line(s,8.8,5.55,4.25,5.55,CYAN,2); line(s,2,4.75,2,3.25,CYAN,2); icon_home(s,6.25,4.05,1.5); add_text(s,'档案不断档\n责任不断链',5.05,4.85,3.2,.8,20,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER)
    s=ns(13); add_title(s,'四队长效：让机制跑在平台后面','有人管 · 有钱管 · 有标准管 · 有机制管'); items=[('专班推进','任务统筹 / 进度督办','doc'),('维修服务','工单闭环 / 备件调度','gear'),('运营管护','巡检更新 / 使用培训','home'),('基金保障','资金统筹 / 长效投入','chart')]
    for i,(t,sub,ic) in enumerate(items):
        x=.55+i*3.08; node(s,i+1,t,sub,x,2,ic); 
        if i<3: arrow(s,x+2.55,2.65)
    kpi(s,'66人','专业管护团队',2.4,4.45,2.6); kpi(s,'近190万','管护基金滚动累计',5.4,4.45,2.8,GOLD); kpi(s,'长期运营','从项目上线到持续运转',8.65,4.45,2.7); add_text(s,'机制不散 · 服务不断 · 数据不乱 · 保障不虚',2.1,6.38,9.1,.45,17,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(14); add_title(s,'应用成效：从经验管房，走向数据管房','资源可见 · 维修提速 · 决策更准 · 资金更省'); primary=[('2442座','建档并完成健康评估'),('48次','烤房调剂'),('18次','突发故障全部3小时内闭环'),('95.8%','群众综合满意度')]
    for i,(num,lab) in enumerate(primary): kpi(s,num,lab,.55+i*3.12,1.85,2.65,CYAN)
    secondary=[('340+','待修烤房筛查'),('150栋','升级纳入计划'),('1363栋 / 81%','合作社统筹管护'),('近190万','管护基金滚动累计')]
    for i,(num,lab) in enumerate(secondary): kpi(s,num,lab,.55+i*3.12,4.25,2.65,GOLD if i==3 else ICE)
    add_text(s,'看得见、调得动、修得快、投得准，才是真正的价值',2,6.45,9.3,.45,17,CYAN,True,BODY_FONT,PP_ALIGN.CENTER)
    s=ns(15); add_title(s,'总结：一套“智烤管家”，把难题变成闭环','一库清底  ·  双模决策  ·  三智落地  ·  四队长效')
    for i,(t,ic) in enumerate([('一库清底','doc'),('双模决策','chart'),('三智落地','gear'),('四队长效','home')]):
        x=1+i*3.05; circle(s,x+.65,2.35,1.25,'0A3558',CYAN,1.3); {'doc':icon_doc,'chart':icon_chart,'gear':icon_gear,'home':icon_home}[ic](s,x+1.27,2.98,.62); add_text(s,f'0{i+1}\n{t}',x,3.75,2.6,.8,20,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER)
    add_text(s,'让每一座烤房，都有一份可计算、可调度、可运维、可决策的健康档案',1,5.35,11.3,.68,20,WHITE,True,TITLE_FONT,PP_ALIGN.CENTER); add_text(s,'从“有多少房”到“有多少有效烤能”',2.6,6.18,8.1,.5,18,CYAN,True,BODY_FONT,PP_ALIGN.CENTER); add_text(s,'让数据说话 · 让模型决策 · 让管家落地',3.2,6.72,6.9,.35,12,MUTED,False,BODY_FONT,PP_ALIGN.CENTER); prs.save(path)

def build_word():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=DInches(.55); sec.bottom_margin=DInches(.55); sec.left_margin=DInches(.65); sec.right_margin=DInches(.65)
    for sty in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[sty]; st.font.name=BODY_FONT; st._element.rPr.rFonts.set(qn('w:eastAsia'),BODY_FONT)
    doc.styles['Normal'].font.size=DPt(10.5); doc.styles['Title'].font.size=DPt(23); doc.styles['Title'].font.color.rgb=DRGB(7,76,120)
    t=doc.add_paragraph(); t.style='Title'; t.alignment=WD_ALIGN_PARAGRAPH.CENTER; t.add_run('智烤管家｜8分钟决赛逐页演讲建议与素材补充说明').bold=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f'配套：10021912原视觉全组件可编辑还原版 + 减光路8分钟优化版｜{STAMP}')
    doc.add_heading('一、这次PPT怎么用',1)
    for x in ['A版尽量还原10021912的蓝白山海、节点和强光路视觉，但所有主体均为PowerPoint原生文本、形状、线条、流程、数字与图表。','B版用于正式8分钟决赛：第1—2页保留舞台光路，第3页以后取消底部大光路，只保留必要的模型连接与节点发光。','真实业务截图/照片不要矢量化造假。第5、10、11、12页保留可编辑证据位，最后用真实小程序、后台、地图和工单截图替换。','目标口播约6分50秒；含15—20秒真实操作视频时总时长约7分10秒—7分30秒。']: doc.add_paragraph(x,style='List Bullet')
    timings=[15,25,20,25,30,40,40,20,15,35,30,25,25,40,25]; names=['封面','三大难题','三条断链','1234总览','一库清底','模型①健康评估','模型②发展导引','双模耦合','三智落地','智用','智投','智管','四队长效','应用成效','总结收口']; tasks=['定题','抛三难','落根因','给地图','讲数据底座','讲存量怎么修','讲增量怎么建','一句话耦合','快速转场','讲真实服务案例','讲投入逻辑','讲生命周期','讲长效机制','只讲硬结果','记忆句收口']
    doc.add_heading('二、8分钟节奏总表',1); tb=doc.add_table(rows=1,cols=4); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(['页','页面','建议时间','任务']): tb.rows[0].cells[i].text=h
    for i,(n,s,task) in enumerate(zip(names,timings,tasks),1): c=tb.add_row().cells; c[0].text=str(i); c[1].text=n; c[2].text=f'{s}s'; c[3].text=task
    talks=[('A','尊敬的各位领导、各位评委、各位专家、各位同仁，大家好！我们是智同道合小组。今天汇报《智烤管家——基于健康评估的烤房用投管一体化解决方案》。我们不是再做一张台账，而是让每一座烤房成为可体检、可调度、可决策、可持续管理的数字资产。','站中间，最后一句说完立即翻页。'),('A','对烟农来说，烤房是烟叶提质增香的“匠心工坊”。可一到烘烤季，最怕三件事：旺季分不到房，故障修不及，建设投不准。','手势只点三次，不念屏幕小字。'),('B','继续往下追，问题不在有没有烤房，而在三条链断了：资源管理断链、维修保障断链、投入决策断链，也就是资源无统筹、投入无量化、管护无闭环。','最后一句面对评委说。'),('A','针对三条断链，我们构建1234体系：一库清底、双模决策、三智落地、四队长效。它不是四个功能，而是一条数据—决策—服务—运营闭环。','顺着01到04轻点。'),('B→A','B：什么是一库清底？A：底数不准、状态不明，后面的模型就无从谈起。我们围绕4大模块、18个关键部件建立统一判定标准，把基础信息、状态、维修和评价沉淀为一房一档。','第5页最终替换真实“一房一档”截图。'),('A','模型一解决旧房该不该修、先修哪里。按部件状态80%、管护10%、现场评价10%形成AHP综合评价，60分以下自动预警；最终输出风险部位、寿命趋势和维修优先序。','不要讲AHP公式。'),('A','模型二回答到底需不需要新建。把单栋健康状态折成村级有效烤能，再结合种植规模、产业趋势、管护能力和政策适配，通过熵权法与TOPSIS识别真实缺口。先调剂、再修复，仍有缺口才新建。','“先调剂、再修复、最后新建”放慢。'),('B','两个模型不是并列。模型一先回答现有烤房还能不能用、该修哪栋；模型二再回答缺口在哪里、该不该建、建多少。前者算存量，后者谈增量。','20秒内翻页。'),('B→A','B：模型算完，怎样真正落到一线？A：我们把能力拆成三个角色场景——烟农技术员用智用，管理人员用智投，合作社和管护队伍用智管。','可在这里接15—20秒真实操作视频。'),('A→B','过去找房、报修、找人、找配件都靠电话；现在可查询空闲烤房、一键调度、线上报修并自动匹配人员配件。磨坪站雷击案例中，平台诊断2处故障部件，3小时内抢修复产。','讲案例时面向评委。'),('A','智投解决钱该投到哪里。先做存量体检，再算有效烤能和村级缺口，最后形成选址与年度投入清单。2026年150栋核心设备升级纳入计划，肖家坪科学新建7栋。','150栋、7栋各停顿一次。'),('A','智管解决建完以后谁来管。从建档、巡检、维修到评价回写，每一次状态变化和工单都回到同一份档案，形成持续更新的生命周期数据。','一圈手势即可。'),('A','平台长期有效，关键是有人管、有钱管、有标准管、有机制管。秭归已形成66人专业管护团队，基金滚动累计近190万元。','停在66人、近190万元。'),('B','一个烤季完成2442座烤房建档和健康评估；48次烤房调剂；18次突发故障全部3小时内闭环；群众综合满意度95.8%。同时筛查340余栋待修、150栋升级纳入计划，两级合作社统筹1363栋、覆盖81%，基金近190万元。','前四个数字一项一顿。'),('A→B→合','A：一库清底、双模决策、三智落地、四队长效，把分不到、修不及、投不准变成闭环解法。B：我们把管理口径从“有多少房”变成“有多少有效烤能”。合：让数据说话，让模型决策，让管家落地。谢谢大家！','结束后停一秒再鞠躬。')]
    doc.add_heading('三、逐页建议口播',1)
    for i,(role,talk,act) in enumerate(talks,1): doc.add_heading(f'第{i}页｜{names[i-1]}｜{timings[i-1]}秒｜{role}',2); doc.add_paragraph('建议口播：'+talk); doc.add_paragraph('临场动作：'+act)
    doc.add_heading('四、真实图片/视频补充优先级',1)
    for x in ['A｜第10页：15—20秒真实操作快剪：查空闲房→调度申请→报修→工单→评价。','A｜第5页：真实“一房一档”小程序/后台截图，替换可编辑证据位。','A｜第11页：真实选址地图或年度投入清单，二选一。','B｜第12页：真实维修工单或巡检记录，二选一。','B｜第7页：肖家坪真实地图/规划图，如有则替换示意地图。','C｜四队合影、抢修现场、真实烤房照片留在答辩备份，不塞进主发布页。']: doc.add_paragraph(x,style='List Bullet')
    doc.add_heading('五、30秒超时应急收口',1); doc.add_paragraph('“我们用一库清底解决底数不清，用双模决策解决修与建凭经验，用三智落地把能力送到一线，再用四队机制保障长期运营。一个烤季完成2442座建档评估、48次调剂、18次抢修全部3小时内闭环，满意度95.8%。智烤管家最终要做的，就是把‘有多少房’变成‘有多少有效烤能’，让数据说话、让模型决策、让管家落地。谢谢大家！”'); doc.save(DOCX)

build_deck(PPT_A,False); build_deck(PPT_B,True); build_word()
for p in [PPT_A,PPT_B,DOCX]:
    with zipfile.ZipFile(p) as z: assert z.testzip() is None
for p in [PPT_A,PPT_B]:
    with zipfile.ZipFile(p) as z:
        media=[n for n in z.namelist() if n.startswith('ppt/media/')]; pics=[n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml') and b'<p:pic' in z.read(n)]; assert not media and not pics; print('EDITABLE_OK',p.name,'slides=15','ppt/media=0','pictures=0')
print('CREATED',PPT_A,PPT_B,DOCX)
