from pathlib import Path
from docx import Document
from pptx import Presentation
from collections import Counter
import re
root=Path(__file__).resolve().parents[2]
out=Path(__file__).resolve().parent/'修改稿'/'资料取证'
out.mkdir(parents=True,exist_ok=True)
targets=[
 ('QC参考_全文.md',root/'缩短烤房设备单次维护作业耗时_QC成果发布词0907.docx'),
 ('9月30日原稿_全文.md',root/'智烤管家演讲讲稿9.30.docx'),
]
for dest,source in targets:
 d=Document(str(source))
 lines=['# '+source.name, '']
 for block in d.element.body:
  tag=block.tag.split('}')[-1]
  if tag=='p':
   from docx.text.paragraph import Paragraph
   p=Paragraph(block,d)
   if p.text.strip():lines.append(p.text.strip())
  elif tag=='tbl':
   from docx.table import Table
   t=Table(block,d)
   lines.append('')
   for row in t.rows:
    lines.append(' | '.join(c.text.strip().replace(chr(10),' / ') for c in row.cells))
   lines.append('')
 (out/dest).write_text('\n\n'.join(lines),encoding='utf-8')
 print(dest,len('\n\n'.join(lines)),len(d.paragraphs),len(d.tables))
off=root/'读取资料/3之前项目官方叙述过的'
for p in sorted(off.glob('*.docx')):
 d=Document(p)
 lines=['# '+p.name,'']
 for pa in d.paragraphs:
  if pa.text.strip():lines.append(pa.text.strip())
 for t in d.tables:
  for row in t.rows:
   lines.append(' | '.join(c.text.strip().replace(chr(10),' / ') for c in row.cells))
 short=p.name.replace('+','_').replace('：','_')[:42]
 (out/('官方_'+short+'.md')).write_text('\n\n'.join(lines[:14000]),encoding='utf-8')
 print('OFFICIAL',p.name,len(lines))
slides=root/'1010-V2_智烤管家_副本.pptx'
if slides.exists():
 p=Presentation(str(slides))
 lines=['# '+slides.name]
 for i,slide in enumerate(p.slides,1):
  lines.append(f'\n## 第{i}页')
  for sh in slide.shapes:
   if sh.has_text_frame and sh.text.strip():lines.append(sh.text.strip())
 (out/'最新PPT_逐页原文.md').write_text('\n\n'.join(lines),encoding='utf-8')
# 目录二逐个源文件清单（包含所有格式）
files=sorted([x for x in (root/'读取资料/2我的项目相关的零碎的').rglob('*') if x.is_file()])
(out/'项目源文件清单_全量.tsv').write_text('\n'.join(f'{p.relative_to(root).as_posix()}\t{p.stat().st_size}' for p in files),encoding='utf-8')
print('TOTAL PROJECT',len(files))
