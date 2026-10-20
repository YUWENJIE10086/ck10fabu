# -*- coding: utf-8 -*-
"""V6: QC-referenced final Word deliverables in 1019/v6/修改稿.
The pre-existing V5 50k/10k studies are explicitly retained as inherited research,
not presented as 50k newly written words. V6 adds QC-specific chapters and audit.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re,zipfile,json
ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).resolve().parent
MD=OUT/'10201004_正式稿全文_QC融合_含PPT页码.md'
SRC=MD.read_text(encoding='utf-8')
QC=(OUT/'资料取证/QC参考_全文.md').read_text(encoding='utf-8')
OFF=(OUT/'资料取证/官方_附件2_宜昌市局_智烤管家-基于健康评估的烤房用投管一体化解决方案_数创成果报告.d.md').read_text(encoding='utf-8')
V5=ROOT/'1019/v5/修改稿'
STAMP='10201004'
def han(s):return len(re.findall('[\u4e00-\u9fff]',s))
def add(d,s,style=None):return d.add_paragraph(s,style=style)
def head(d,s):return add(d,s,'Heading 2')
def table(d,h,rows):
 t=d.add_table(rows=1,cols=len(h));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for cell,s in zip(t.rows[0].cells,h):cell.text=str(s)
 for row in rows:
  cells=t.add_row().cells
  for cell,s in zip(cells,row):cell.text=str(s)
 return t
def basic(title,subtitle):
 d=Document();sec=d.sections[0];sec.left_margin=Cm(2.25);sec.right_margin=Cm(2.25);sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.0)
 st=d.styles['Normal'];st.font.name='仿宋_GB2312';st.font.size=Pt(11);st._element.rPr.rFonts.set(qn('w:eastAsia'),'仿宋_GB2312')
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=p.add_run(title);r.font.name='黑体';r.font.size=Pt(16);r.bold=True;r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
 add(d,subtitle)
 return d
segments=[];segment=None
for ln in SRC.splitlines():
 if ln.startswith('## 【'):
  m=re.match(r'## 【(P[^】]+)】(.+)',ln)
  assert m,ln
  segment=[m.group(1),m.group(2),[]];segments.append(segment)
 elif ln.startswith(('A：','B：')):
  assert segment is not None
  segment[2].append((ln[0],ln[2:]))
A=sum(han(t) for pg,topic,ls in segments for who,t in ls if who=='A')
B=sum(han(t) for pg,topic,ls in segments for who,t in ls if who=='B')
assert (A,B)==(1094,465),(A,B)
assert len(segments)==16,len(segments)
def script(stage):
 d=Document(ROOT/'智烤管家演讲讲稿9.30.docx')
 for item in list(d._element.body):
  if item.tag!=qn('w:sectPr'):d._element.body.remove(item)
 if stage:
  d.sections[0].top_margin=Cm(1.65);d.sections[0].bottom_margin=Cm(1.65)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=p.add_run('智烤管家——基于健康评估的烤房用投管一体化解决方案');r.bold=True;r.font.name='黑体';r.font.size=Pt(14)
 r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
 for pg,topic,lines in segments:
  if stage:
   p=d.add_paragraph()
   p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(0)
   r=p.add_run(pg+'  '+topic);r.font.name='黑体';r.font.size=Pt(9);r.font.color.rgb=RGBColor(87,102,117)
   if pg=='P2':
    p=add(d,'（视频剪至30秒，AB静默；原始片长约41.5秒。）')
    for x in p.runs:x.font.size=Pt(10)
  for who,txt in lines:
   p=d.add_paragraph();p.paragraph_format.space_after=Pt(2 if stage else 4);p.paragraph_format.line_spacing=1.05 if stage else 1.11
   r=p.add_run(who+'：');r.bold=True;r.font.name='黑体';r.font.size=Pt(14);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
   r=p.add_run(txt);r.font.name='仿宋_GB2312';r.font.size=Pt(14);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'仿宋_GB2312')
 return d
script(False).save(OUT/f'{STAMP}_04_智烤管家V6_QC融合_8分钟AB正式演讲净稿.docx')
script(True).save(OUT/f'{STAMP}_05_智烤管家V6_25页PPT双人演讲舞台稿.docx')

qcinsights=[
('QC证据链的精髓','QC原稿先有40分钟基准、26分钟目标，再通过14.9/12.2分钟症结调查、扫码与专用工具措施、60次作业17.9分钟结果建立因果闭环。V6学习的是证据组织方式，不把QC的数值当成本项目数据。'),
('B角身份调整','QC原稿B以扎根现场的李师傅身份追问维护难题，符合15分钟QC场景。V6的B为项目组成员、评审视角提问者，依次追问“为何要管家”“算法为何可信”“实际效果与长期运行如何证明”，不假扮评委。'),
('保留专业并缩至8分钟','QC中现状调查、目标设定、要因验证、实施、结果和巩固均有多轮介绍；本项目只用三愁三缺交代问题，双模解释方法，三智用真实数据验证，四队说明标准化持续。'),
('明确不能套用QC效益','40降至17.9分钟、63.6万元年效益、扫码调取12秒、842台燃烧机、60次维护均属于金魔方QC课题；智烤没有对应同口径前后对照，不能改写成智烤的实际节时和收益。'),
('技术兼容两种评委','信息评委看输入、权重、边界、输出；烟叶评委看有效烤能、旺季调剂、响应速度、队伍责任。每个技术段落须紧接业务价值，并说明人工审核与现场核验边界。'),
('语言节奏','只在三愁三缺、1234亮相和结尾用对仗；技术段保持规范短句，不用“太棒了”“神奇”等捧哏或夸大词。')
]
# 01: twelve txt award learning retained plus V6 focused QC chapter
aw=Document(V5/'10192342_01_V5_十二篇获奖现场解说词学习与创新借鉴.docx')
head(aw,'V6追加：QC发布词（2026）完整研读与12篇获奖讲词的共同方法')
for h,b in qcinsights:head(aw,h);add(aw,b)
aw.save(OUT/f'{STAMP}_01_获奖12篇及QC成果发布词深度学习.docx')
# 02: 50k previous source-indexed volume retained, 8 original diagrams
depth=Document(V5/'10192342_02_V5_项目五万字级全文件研究与双模创新增订.docx')
head(depth,'V6新增：从QC的效果验证链到智烤的真实落地链')
for h,b in qcinsights:head(depth,h);add(depth,b)
head(depth,'技术实质：两模不只是并排展示，而是共同决定“先修后建”的投入次序')
for t in [
 '模型1先整合基础档案、关键部件、管护维修和评价资料，对12项算法参数进行健康和寿命计算；按部件健康80%、管护10%、评价10%形成整栋健康分，同时给出风险部位排序、寿命提醒和维修优先建议。模型未能单凭分数代替现场维修诊断。',
 '模型2以千亩村为评价单元，将模型1的健康状态与名义烤房数量、实际使用条件结合折算有效烤能Q，再与理论需求D形成缺口判断；候选点使用AHP和熵权法组合赋权，用TOPSIS排序，通过地图沙盘模拟新增数量和空间布局。',
 '创新的逻辑连接在于：健康评估改变有效烤能估计，后者改变村级缺口与新建建议；与传统“只凭名义栋数和烟农呼声”相比，可同时考虑维修价值和增量配置。',
 'QC借鉴的是按“症结—措施—效果—固化”说明每个环节，不把其他QC课题的40分钟、17.9分钟、63.6万元效益直接转换为智烤成绩。',
 'V6定稿优先采用附件2正式报送数据：2442全口径建档、健康分低于60的340多栋预警研判、48次调剂、18次报修、其中16条平台工单0.38小时平均响应、531份有效问卷96.8%、1363/1672约81%管护覆盖、基金总额近190万元。']:
 add(depth,t)
head(depth,'资料覆盖与附录说明')
add(depth,'本研究卷继承V5逐文件研究与原始资料摘录、图片及资料路径，保留50,000字以上的研究与引述总篇幅；V6新增的是QC发布方法与双模落地的核验章节，不能将继承原文重复称为“V6新写5万字”。')
inventory=(OUT/'资料取证/项目源文件清单_全量.tsv').read_text(encoding='utf8').splitlines()
add(depth,f'该资料夹共登记{len(inventory)}个文件（含系统隐藏文件；实质资料比总数少），逐项路径见同目录“资料取证/项目源文件清单_全量.tsv”，不把不同版本、重复副本混为独立成效证据。')
depth.save(OUT/f'{STAMP}_02_五万字项目研究卷_含QC与双模创新增订.docx')
# 03: official 10k source-backed volume kept, 3 original visuals
off=Document(V5/'10192342_03_V5_官方叙述一万字级研究校订版.docx')
head(off,'V6重点：QC成果发布规范与数创成果报送的不同要求')
for h,b in qcinsights:head(off,h);add(off,b)
head(off,'正式报送的实施结果与计划目标必须分清')
for t in [
 '附件2报告正式载明：2442座全口径烤房信息入库；16条平台工单平均响应时间0.38小时；531份有效问卷综合满意度96.8%。这些是已报告的运行结果。',
 '资料中的预期“满意率达到95%以上”、建设推荐或维修计划，与已完成工程验收不是同义。既不能把340多栋预警全部写成已修复，也不能将早期版本中的150栋升级、肖家坪7栋建设冒称正式已完成。',
 '秭归两级合作社统筹管护1363栋（本年度在用1672栋），覆盖率约81%；基金总额近190万元，不能说成余额。',
 'PPT第6页的18项清查和第8页12个模型参数是不同评价层级；PPT第8页寿命估算流程与第9页Weibull文字须在现场备答中统一参数路径。',
 '角色定位：A是发布人，B代表项目组模拟评审视角抛出专业问题；B不以评委身份发言，也不让A/B两人轮流背大段同质化介绍。']:
 add(off,t)
off.save(OUT/f'{STAMP}_03_官方叙述一万字研究_成果报告与QC规范校订.docx')

compare=basic('V6修改对照｜9.30原稿、V1—V5与QC成果发布词','重点是保留项目最初优势、恢复真正技术和证明、剔除无验收依据的完成性表述。')
table(compare,['原有优点/争议','V6采用方案','PPT页码'],[
('9.30“匠心工坊”与“三缺”','保留；用烟农三愁引出行业三缺','P3–P4'),
('“1234”总方案','四组数字各对应清楚业务作用','P5'),
('四大部位18项清查及建档','保留真实采集范围与2442入库','P6'),
('健康模型加权与风险处置','AHP、80/10/10、寿命及风险部位TOP3','P7–P9'),
('发展导引不足以一句“地图沙盘”带过','数需评推、Q/D、AHP+熵权、TOPSIS、情景推演','P10–P11'),
('QC因果论证','B从评审视角追问、A以模型输入输出与运行实绩作答','全篇'),
('智用报修与调剂','48调剂、18报修、16条0.38小时严格分样本','P14–P15'),
('智投和智管','340多栋风险研判、真实派单验收回写','P16–P19'),
('四队与组织保障','职责、制度、1363/1672和基金总额近190万','P20–P21'),
('项目初稿95.8%满意度','使用正式报送531份96.8%','P22'),
('150栋/7栋已完成、66人等历史表述','转为备查线索，未有验收证据不报完成','答辩备查'),
('跨县推广','秭归一季试点，跨区需校准参数和机制','P23–P24')])
compare.save(OUT/f'{STAMP}_06_原稿及V1V5与QC融合修改对照.docx')
rehearse=basic('V6逐页彩排卡｜A/B 70/30，含30秒视频的8分钟发布','按220汉字/分钟的保守口播估算，仍需现场逐秒计时。')
table(rehearse,['统计量','计时或判定'],[
('A口述',f'{A}汉字'),('B口述',f'{B}汉字'),('A比例',f'{A/(A+B)*100:.2f}%'),
('220字/分',f'{(A+B)/220*60:.1f}秒'),('视频/切页', '30秒＋20秒'),('目标总时长',f'{(A+B)/220*60+50:.1f}秒（应控制≤480秒）')])
for pg,title,ls in segments:
 head(rehearse,f'{pg}｜{title}')
 for who,txt in ls:add(rehearse,who+'：'+txt)
 if pg=='P2':add(rehearse,'【导演提示】开场视频必须剪至30秒，原素材约41.5秒。')
 if pg in ('P7—P9','P10—P11'):add(rehearse,'【导演提示】AHP/TOPSIS读法提前统一，数需评推四字稍作停顿。')
 if pg in ('P14—P15','P16—P17','P18—P19'):add(rehearse,'【导演提示】内嵌演示视频不额外完整加播，可同步短片或静帧。')
head(rehearse,'评委追问5条')
for t in [
 '健康分分母：18项清查≠12项模型参数，赛前准备字段映射表。',
 '寿命估算：P8短板加权公式与P9 Weibull描述需统一实际实施方法，不报未经验证的预测精度。',
 '0.38小时：16条工单平均响应，不是18次故障的平均维修完成时间。',
 '基金：近190万元指基金总额，不等于当期结余。',
 '有效烤能：需交代健康状态变化怎样进入新建缺口计算和沙盘推演。']:
 add(rehearse,t)
rehearse.save(OUT/f'{STAMP}_07_逐页演讲彩排与评委追问卡.docx')
qcdoc=basic('QC成果发布词专项学习与V6成果化改写','所研究的QC课题是《缩短烤房设备单次维护作业耗时》，约15分钟；V6项目上台只有8分钟。')
for h,b in qcinsights:head(qcdoc,h);add(qcdoc,b)
table(qcdoc,['QC现场展开','V6取舍','原因'],[
('调查基准、工序拆解、要因分析','三愁三缺概括问题','8分钟内先交代重点'),
('三大对策、参数与实施过程','一库和双模简洁展示真实算法','保留技术可信度'),
('60次实测17.9分钟','用本项目真实运行指标替代','不能移植QC实验样本'),
('四份标准化文件和巩固期','四队、两项制度、基金与回写','体现持续运营'),
('B作为李师傅提供现场经验','B提出评审疑问并从业务角度解释','身份不同，避免模仿他人')])
head(qcdoc,'V6评委追问结构')
for t in ['第一问：为什么需要数字管家？——对应烟农三愁和管理三缺。','第二问：模型凭什么判断？——对应数据源、AHP健康分和Q/D/TOPSIS。','第三问：应用效果怎样、谁来持续做？——对应48次调剂、18次报修、531份问卷、四队与基金。']:add(qcdoc,t)
qcdoc.save(OUT/f'{STAMP}_08_QC发布词专项学习与V6成果化改写.docx')
# Validate generated docs and zip
docs=sorted(OUT.glob(STAMP+'_0*.docx'))
assert len(docs)==8,[x.name for x in docs]
depthcnt=sum(han(p.text) for p in depth.paragraphs)
officialcnt=sum(han(p.text) for p in off.paragraphs)
assert depthcnt>=50000 and len(depth.inline_shapes)>=8,(depthcnt,len(depth.inline_shapes))
assert officialcnt>=10000 and len(off.inline_shapes)>=3,(officialcnt,len(off.inline_shapes))
checks=[f'# V6 complete Word audit; A={A};B={B}',f'depth={depthcnt} chars, {len(depth.inline_shapes)} diagrams',f'official={officialcnt} chars,{len(off.inline_shapes)} diagrams']
for f in docs:checks.append(f'{f.name}\t{f.stat().st_size} bytes')
(OUT/f'{STAMP}_Word生成核验.md').write_text('\n'.join(checks),encoding='utf8')
zipfile_path=OUT/f'{STAMP}_智烤管家V6_全部8份Word.zip'
with zipfile.ZipFile(zipfile_path,'w',zipfile.ZIP_DEFLATED) as z:
 for f in docs:z.write(f,arcname=f.name)
 z.write(OUT/f'{STAMP}_Word生成核验.md',arcname='Word生成核验.md')
assert zipfile_path.stat().st_size>500000
print('\n'.join(checks));print('ZIP',zipfile_path.stat().st_size)
