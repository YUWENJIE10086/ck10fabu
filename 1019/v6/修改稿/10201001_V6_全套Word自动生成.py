# -*- coding: utf-8 -*-
"""V6 Word builder. Inherit full sourced V5 studies, append V6/QC analysis, rebuild professional A/B script."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).resolve().parent
SOURCE=OUT/"10201001_智烤管家V6_决赛双人演讲稿_源稿.md"
V5=ROOT/"1019/v5/修改稿"
STAMP="10201001"
HAN=lambda x:len(re.findall(r"[\u4e00-\u9fff]",str(x)))
groups=[]
for line in SOURCE.read_text(encoding="utf-8").splitlines():
 if line.startswith("## 【P"):
  m=re.match(r"## 【(P[^】]+)】(.+)",line);groups.append([m.group(1),m.group(2),[]])
 elif line.startswith(("A：","B：")):groups[-1][2].append((line[:1],line[2:]))
A=sum(HAN(t) for _,_,ls in groups for r,t in ls if r=="A")
B=sum(HAN(t) for _,_,ls in groups for r,t in ls if r=="B")
assert A+B<=1550 and .69<A/(A+B)<.71,(A,B)
assert len(groups)==14,len(groups)
def writ(d,t,style=None):return d.add_paragraph(t,style=style)
def section(d,t):writ(d,t,"Heading 2")
def table(d,heads,rows):
 tab=d.add_table(rows=1,cols=len(heads));tab.style="Table Grid"
 for c,s in zip(tab.rows[0].cells,heads):c.text=str(s)
 for row in rows:
  cells=tab.add_row().cells
  for c,s in zip(cells,row):c.text=str(s)
def font(r,name,size,bold=False):
 r.font.name=name;r.font.size=Pt(size);r.bold=bold
 r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),name)
def doc_speech(stage):
 d=Document(ROOT/"1004/智烤管家演讲讲稿9.30.docx")
 for n in list(d._element.body):
  if n.tag!=qn("w:sectPr"):d._element.body.remove(n)
 sec=d.sections[0]
 if stage:sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 font(p.add_run("智烤管家——基于健康评估的烤房用投管一体化解决方案"),"黑体",14,True)
 for num,title,ls in groups:
  if stage:
   p=d.add_paragraph();p.paragraph_format.keep_with_next=True
   p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(1)
   r=p.add_run(num+"｜"+title);font(r,"黑体",9,True);r.font.color.rgb=RGBColor(75,100,130)
   if num=="P2":
    p=d.add_paragraph();font(p.add_run("（开场视频30秒；原素材约41.5秒，现场AB静默）"),"楷体",9)
  for role,text in ls:
   p=d.add_paragraph()
   p.paragraph_format.space_after=Pt(2 if stage else 4)
   p.paragraph_format.line_spacing=1.06 if stage else 1.12
   font(p.add_run(role+"："),"黑体",14,True)
   font(p.add_run(text),"仿宋_GB2312",14)
 path=OUT/(f"{STAMP}_"+("05_V6_25页PPT双人舞台口播稿.docx" if stage else "04_V6_8分钟正式AB演讲净稿.docx"))
 d.save(path)
 return path
speech=doc_speech(False);stage=doc_speech(True)
# Research 01: all twelve award scripts retained + QC full structure analysis.
d=Document(V5/"10192342_01_V5_十二篇获奖现场解说词学习与创新借鉴.docx")
section(d,"V6增补专题：十五分钟QC成果发布词的证据组织及双人发布方式")
for title,content in [
("题材边界","QC《缩短烤房设备单次维护作业耗时》是另一项目，其维护40分钟、目标26分钟、实测17.9分钟、63.6万元效益等只用于研究QC论证，不得列为智烤管家成绩。"),
("QC的七环逻辑","现场问题—现状调查—关键症结—目标及措施—措施实施—现场效果—标准化巩固。V6凝练为三愁三缺→1234→模型依据→三智成效→四队长效。"),
("B的身份不是捧哏","QC稿B为一线师傅，承担实操经验及目标难度、成本、长效追问。V6改为评委视角引导人，专门问风险是否可追溯、为何选址、效果有何证据、谁执行与支付。"),
("技法：量化证据有因果链","QC稿先找出维护耗时主因，再说明工具如何针对主因，之后检验结果。V6的两模不能只报算法名称，必须解释存量健康状态如何改变有效烤能，进而改变新建缺口。"),
("技法：实效后说固化","QC以标准化和巩固检验证明持续性。V6以四队、两项管护制度、合作社基金和档案回写证明运行不止一个展示页面。"),
("语言控制","保留原稿“匠心工坊”“资源无统筹、投入无量化、管护无闭环”；避免空泛夸赞、重复插科打诨。押韵重点放在开头、1234总结与尾声。")
]:section(d,title);writ(d,content)
table(d,["QC发布稿可学之处","V6演讲落点"],[
("关键问题先于方案","三愁→三缺，而不是开场堆功能"),
("B有专业身份并发问","B是评委关注点引导人，问题是模型、实效、可持续"),
("比对基线/目标/结果","智烤只说同口径已核指标，响应和维修完结分开"),
("成果变标准方法","四队+制度+基金+回写构成运行机制")])
d.save(OUT/f"{STAMP}_01_V6_12篇优秀讲稿与QC发布方法学习.docx")
# Research 02: complete prior sourced V5 53k research + genuine V6 QC comparative section
d=Document(V5/"10192342_02_V5_项目五万字级全文件研究与双模创新增订.docx")
section(d,"V6专题研究：由QC证据链延伸出的双模型可解释性")
for t in [
"QC项目采用维护工序拆解与现场验证说明改进有效，其指标样本和项目目标与本项目不同，不作结果借用。可以迁移的是“识别瓶颈—定位机制—措施验证—长期巩固”的研究逻辑。",
"智烤管家首先必须准确记录烤房身份与部件状态：四大部位18项属于现场清查；健康模型图使用12个计算部件参数。两个层级需要专门的字段映射表，不应直接称数据矛盾已解决。",
"健康模型以基础档案、部件数据、检查评价、管护维修四类数据为输入，AHP确定部件权重，80%部件状态、10%管护、10%人工评价形成整栋健康分；还估算剩余寿命，生成风险烤房、风险部位前三及维修优先级。",
"PPT第8页的状态修正与寿命合成公式、第9页提及的Weibull拟合仍需研发团队统一现行方法；在确认前，正式发布不声称当前系统严格采用某一具体分布模型。",
"千亩村发展导引按“数需评推”组织。数据入模后，使用健康状态折算有效烤能Q，与理论烤能需求D比较；采用AHP与熵权法组合赋权、TOPSIS点位排序；通过地图沙盘模拟新增栋数与位置并比较供需改善。",
"两模之间不是简单并列，而是健康评估结果改变有效烤能，继而影响新建需求、推荐数量与点位排序。这使“先维护存量、再科学投入增量”具备可解释的管理依据。",
"三智将模型输出转为具体业务；四队与两项制度、基金提供责任主体。派单→维修→验收→评价→数据回写→模型复算，构成可验证闭环。",
"成效数据须区别统计分母：2442全量建档、340余栋低分预警、48次调剂、18次报修、16条平台工单平均响应0.38h、531份有效问卷满意度96.8%、1363栋统筹管护/年度在用1672栋约81%、基金总额近190万元。"
]:writ(d,t)
d.save(OUT/f"{STAMP}_02_V6_项目五万字级全文件研究与QC方法增订.docx")
# Official 03: prior sourced 11k record plus V6 courtroom wording
d=Document(V5/"10192342_03_V5_官方叙述一万字级研究校订版.docx")
section(d,"V6增订：从QC发布法到八分钟决赛表达的正式叙述标准")
for t in [
"角色：A为主发布人，承担技术、业务、结论；B为评委视角引导人，依问题递进提问，禁止机械捧哏、反复口语感叹。台词按实际汉字计数约70%和30%。",
"技术：模型1必须讲AHP部件权重、80/10/10综合健康分、寿命估计、风险部位TOP3与维修清单；模型2必须讲Q有效烤能、D理论需求、AHP＋熵权、TOPSIS排序与地图沙盘。",
"应用：不把“预警研判对象”说成“已经全部修复”，不把“工单响应时间”误当“维修完工时间”，不把“管护基金总额”说成“年度可支配余额”。",
"QC项目独立：原QC稿相关40分钟、17.9分钟、经济效益、专利材料不得认定为智烤管家自有数据或项目实绩。",
"节奏：严格将30秒开场视频计入8分钟，保留至少20秒换页与停顿；未经实测不得宣布准时通过。PPT的其他嵌入演示视频如完整加播，必须等量删减台词。"
]:writ(d,t)
d.save(OUT/f"{STAMP}_03_V6_官方叙述一万字级研究及QC规范.docx")
# 06: 25-page crosswalk
d=Document();d.add_heading("V6｜25页PPT逐页口播与8分钟彩排卡",0)
titles=["封面","视频","三愁","三缺","1234方案","一房一档","双模总览","健康模型计算","健康模型三步","村级导引技术图","数需评推","三类用户","三智协同","智用调剂报修","服务48次/18次","维修清单和资金","新建沙盘","智管闭环","AI助手","管护覆盖与基金","调抢养综","运行成效","创新价值","推广条件","结束语"]
def page_numbers(pg):
 m=re.match(r"P(\d+)(?:[—-]P?(\d+))?",pg)
 return list(range(int(m.group(1)),int(m.group(2) or m.group(1))+1))
mapping={i:(pg,title,ls) for pg,title,ls in groups for i in page_numbers(pg)}
assert sorted(mapping)==list(range(1,26))
for i,t in enumerate(titles,1):
 section(d,f"第{i}页｜{t}")
 pg,name,ls=mapping[i]
 if i==2:writ(d,"（视频30秒，AB静默；需从原约41.5秒剪辑。）")
 elif i==page_numbers(pg)[0]:
  for who,s in ls:writ(d,who+"："+s)
 else:writ(d,"本页与前页共用一组解说，保留图示支撑，不增加重复口述。")
 if i in (8,9):writ(d,"技术追问：P8合成寿命与P9 Weibull拟合的现行关系待技术组最终确认。")
 if i in (14,16,17,19):writ(d,"该页视频演示不能叠加完整播放，否则会突破8分钟。")
section(d,"AB比例及计时")
table(d,["统计","结果"],[
("A角",A),("B角",B),("A占比",f"{100*A/(A+B):.2f}%"),("B占比",f"{100*B/(A+B):.2f}%"),
("口播220汉字每分",f"{(A+B)*60/220:.1f}秒"),("视频","30秒"),("换页停顿","约20秒"),
("总预算",f"{(A+B)*60/220+50:.1f}秒（非实测）")])
d.save(OUT/f"{STAMP}_06_V6_25页PPT逐页对应与8分钟彩排.docx")
# 07: comparative and questions
d=Document();d.add_heading("V6｜9.30原稿、历版与QC方法融合对照及备答",0)
table(d,["来源","保留优势","修改用途"],[
("9.30原稿","匠心工坊、三缺、管理实效与基金机制","保留语言力量和业务信息，纠正未经证实完成性描述"),
("V1","现场冲突与双人节奏","三愁→三缺建立立题"),
("V2","体检与选址的可理解类比","简明生活化解释技术内核"),
("V3","双模型完整步骤","模型输入→计算→结果和管理用途"),
("V4","激情凝练、四队闭环","正式对仗、长效机制"),
("V5","双模型相互耦合","健康改变有效烤能影响新建投入"),
("15分钟QC发布词","B带身份、现场证据、目标验证与标准化","B代表评委追问，A作专业论证，不套用QC数据")
])
section(d,"评委技术问答")
for q,a in [
("为什么一库不只是电子台账？","一房一档关联关键部件、维修、评价与空间信息，为风险计算和情景演练提供可追溯输入。"),
("AHP权重是否可靠？","专家赋权须记录权重来源及一致性检验，健康分还取决于部件现场状态、管护、评价与寿命信息。"),
("如何理解剩余寿命？","属于依据年限与状态的估算，不等于某日精确故障预报；P8与P9的两种计算表述需由研发澄清。"),
("18清查项和12参数如何对应？","前者是现场清查范围，后者是当前模型图计算参数；应提供字段映射，不能自行推断完全相等。"),
("为什么选址要用三个算法？","AHP表示专业权重、熵权补充数据差异、TOPSIS排序候选点；排序要接受现场踏勘与审批约束。"),
("两模创新为何成立？","健康改变存量有效烤能，存量有效烤能决定新建需求；维修与投资因此可以在同一套指标链上协同。"),
("为什么能够长期运行？","四支队伍承担抢修、调制、巡检和验收，制度与基金保证流程，维修结果回写健康档案，支持长期迭代。"),
("什么是可以复核的试点成效？","2442建档、340+低分预警、48调剂、18报修、16条工单平均响应0.38h、531问卷96.8%、1363/1672管护约81%、基金总额近190万。")
]:section(d,"问："+q);writ(d,"答："+a)
writ(d,"QC项目自身的维护40→17.9分钟、节约金额等数据属于独立QC研究，不移作《智烤管家》成果。")
d.save(OUT/f"{STAMP}_07_V6_初稿历版QC对照及评委问答.docx")
summary=f"口述A={A} B={B} 总={A+B} A比例={100*A/(A+B):.2f}%。\nV6全套7份Word；研究02继承V5完整5万字级资料并增订，研究03继承V5完整1万字级官方研究并增订。"
(OUT/f"{STAMP}_交付与验收说明.md").write_text(summary,encoding="utf8")
print(summary)
