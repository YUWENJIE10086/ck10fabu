# -*- coding: utf-8 -*-
"""Generate editable .docx files for the 1019 presentation deliverables.
Generated Word files are based on the corresponding reviewed Markdown sources.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
repo=root.parent
orig=repo/"智烤管家演讲讲稿9.30.docx"
for src in sorted(root.glob("1019*.md")):
    if "决赛演讲稿" not in src.name and "获奖解说词" not in src.name:
        continue
    script="决赛演讲稿" in src.name
    doc=Document(orig) if script and orig.exists() else Document()
    if script:
        b=doc.element.body
        for el in list(b):
            if el.tag!=qn("w:sectPr"):
                b.remove(el)
    sec=doc.sections[0]
    sec.top_margin=Cm(2.4);sec.bottom_margin=Cm(2.3)
    sec.left_margin=Cm(2.65);sec.right_margin=Cm(2.45)
    normal=doc.styles["Normal"]
    normal.font.name="宋体";normal.font.size=Pt(14 if script else 10.5)
    normal.font.color.rgb=RGBColor(20,29,41)
    normal.paragraph_format.space_after=Pt(5 if script else 7)
    normal.paragraph_format.line_spacing=1.22 if script else 1.18
    for style in ["Heading 1","Heading 2","Title"]:
        st=doc.styles[style]
        st.font.name="黑体"
        st.font.color.rgb=RGBColor(24,71,125)
    doc.styles["Heading 1"].font.size=Pt(14)
    doc.styles["Heading 2"].font.size=Pt(11)
    doc.styles["Title"].font.size=Pt(17)
    lines=src.read_text(encoding="utf-8").splitlines()
    for line in lines:
        line=line.strip()
        if not line:
            continue
        if line.startswith("# "):
            p=doc.add_paragraph(line[2:],style="Title")
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_paragraph(line[3:],style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:],style="Heading 2")
        elif line.startswith("> "):
            p=doc.add_paragraph(line[2:])
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.size=Pt(9)
        elif script and (line.startswith("A：") or line.startswith("B：")):
            p=doc.add_paragraph()
            p.paragraph_format.keep_together=False
            marker=line[:2]
            r=p.add_run(marker)
            r.bold=True;r.font.name="黑体";r.font.size=Pt(14)
            p.add_run(line[2:])
        elif script and line.startswith("（合）"):
            p=doc.add_paragraph()
            r=p.add_run(line[:3]);r.bold=True;r.font.name="黑体"
            p.add_run(line[3:])
        elif script and line.startswith("【"):
            p=doc.add_paragraph(line)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size=Pt(11);r.italic=True;r.font.color.rgb=RGBColor(120,120,120)
        else:
            doc.add_paragraph(line)
    dest=src.with_suffix(".docx")
    doc.save(dest)
    print(f"Generated {dest} ({dest.stat().st_size} bytes)")
