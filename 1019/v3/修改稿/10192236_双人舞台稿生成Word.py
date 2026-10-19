# -*- coding: utf-8 -*-
"""1019/v3/修改稿：基于9.30原始Word格式，生成高感染力双人舞台稿、净稿、修改说明。
逐字以同目录Markdown为准；AB比例与视频计时由源文本实时核算。
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT=Path(__file__).resolve().parents[3]
DIR=Path(__file__).resolve().parent
ORIGINAL=ROOT/"1004"/"智烤管家演讲讲稿9.30.docx"
SOURCE=DIR/"10192230_智烤管家V3_决赛激情朗诵版_含PPT页码.md"
STAMP="10192236"
assert ORIGINAL.exists(), ORIGINAL
assert SOURCE.exists(), SOURCE
raw=SOURCE.read_text(encoding="utf-8")
lines=raw.splitlines()

def hans(s):
    return len(re.findall(r"[\u4e00-\u9fff]",str(s)))

A,B=0,0
blocks=[]
current=None
for line in lines:
    line=line.strip()
    if line.startswith("## 【"):
        current={"heading":line[3:],"lines":[]}
        blocks.append(current)
    elif line.startswith(("A：","B：")):
        assert current is not None
        current["lines"].append(line)
        if line.startswith("A："):A+=hans(line[2:])
        else:B+=hans(line[2:])
    elif line.startswith("（") and current:
        current["lines"].append(line)
assert len(blocks)==16,(len(blocks),[x["heading"] for x in blocks])
assert A+B<=1600,(A,B)
assert 0.69<A/(A+B)<0.71,(A,B,A/(A+B))
PPT_PAGES=[]
for blk in blocks:
    m=re.search(r"【P(\d+)(?:[—–-]P?(\d+))?】",blk["heading"])
    assert m,blk["heading"]
    lo,hi=int(m.group(1)),int(m.group(2) or m.group(1))
    PPT_PAGES.extend(range(lo,hi+1))
assert sorted(PPT_PAGES)==list(range(1,26)),PPT_PAGES

def clear_body(d):
    for el in list(d._element.body):
        if el.tag!=qn("w:sectPr"):
            d._element.body.remove(el)
def make_doc():
    d=Document(str(ORIGINAL))
    clear_body(d)
    return d
def add_title(d):
    p=d.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(8)
    r=p.add_run("智烤管家——基于健康评估的烤房用投管一体化解决方案")
    r.font.name="黑体";r.font.size=Pt(14);r.bold=True
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),"黑体")
def runtext(p,text):
    # only selected key phrases bold, to make live recital easier to scan
    keywords=[
       "想用，调不到；想修，等不及；想问，问不明！",
       "资源无统筹、投入无量化、管护无闭环！",
       "一库见底、两模决策、三智落地、四队长效",
       "数、需、评、推",
       "服务不掉线，责任不断档！",
       "五百三十一份有效问卷",
       "百分之九十六点八",
       "数据不睡觉，服务不断线"
    ]
    pattern="("+ "|".join(re.escape(x) for x in sorted(keywords,key=len,reverse=True)) +")"
    parts=re.split(pattern,text)
    for part in parts:
        if not part:continue
        r=p.add_run(part)
        r.font.name="仿宋_GB2312"
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"仿宋_GB2312")
        r.font.size=Pt(14)
        if part in keywords:r.bold=True
def speaker(d,line,compact=False):
    role,content=line.split("：",1)
    p=d.add_paragraph()
    p.paragraph_format.space_after=Pt(2 if compact else 4)
    p.paragraph_format.line_spacing=1.06 if compact else 1.13
    r=p.add_run(role+"：")
    r.bold=True;r.font.name="黑体";r.font.size=Pt(14)
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"黑体")
    runtext(p,content)
def subtitle(d,line):
    p=d.add_paragraph()
    p.paragraph_format.space_before=Pt(3)
    p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.keep_with_next=True
    r=p.add_run(line)
    r.font.name="黑体";r.font.size=Pt(9)
    r.font.color.rgb=RGBColor(93,103,119)
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"黑体")
def aside(d,line):
    p=d.add_paragraph()
    p.paragraph_format.space_after=Pt(4)
    r=p.add_run(line)
    r.italic=True;r.font.size=Pt(10)
    r.font.color.rgb=RGBColor(105,105,105)
def base_staging(d,notes=True):
    add_title(d)
    for bl in blocks:
        if notes:
            subtitle(d,bl["heading"])
            cue=""
            head=bl["heading"]
            if "【P1】" in head:cue="【舞台】微笑亮相，B说“管家”时轻快，A说项目名时停顿半拍"
            elif "【P3】" in head:cue="【舞台】A放慢，让评委先看到烟农；B三句痛点层层递进"
            elif "【P4】" in head:cue="【舞台】三缺排比加重音，后面三组“未必”再推进一层"
            elif "【P5】" in head:cue="【舞台】四个数字不抢语速，A高声提出，B节奏接力"
            elif "【P7" in head:cue="【舞台】模型一讲清“输入—计算—输出”，B负责让评委会心一笑"
            elif "【P10" in head:cue="【舞台】“数、需、评、推”四字重读；AHP/TOPSIS清晰而不过度拉长"
            elif "【P14" in head:cue="【舞台】B模拟烟农急切，而非夸张表演；0.38小时必须说平均响应"
            elif "【P20" in head:cue="【舞台】B提问后稍停顿，A用“靠四队”回答，建立力量感"
            elif "【P22" in head:cue="【舞台】报出531份、96.8%时适当放慢，让数字落地"
            elif "【P25" in head:cue="【舞台】最后一组排比抬头看观众，不催促谢幕"
            if cue:subtitle(d,cue)
        for line in bl["lines"]:
            if line.startswith(("A：","B：")):speaker(d,line,notes)
            elif notes:aside(d,line)

# with precise PPT pages and performance cues
stage=make_doc()
base_staging(stage,True)
stage_path=DIR/f"{STAMP}_01_智烤管家V3_决赛强化双人舞台稿_含PPT与语气提示.docx"
stage.save(str(stage_path))

# printable clean original-format version, no slide markers
clean=make_doc()
base_staging(clean,False)
clean_path=DIR/f"{STAMP}_02_智烤管家V3_决赛冲刺_正式演讲净稿.docx"
clean.save(str(clean_path))

# compact, constructive comparison and rehearsal sheet
review=Document()
sec=review.sections[0]
sec.left_margin=Cm(2.4);sec.right_margin=Cm(2.4)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.0)
normal=review.styles["Normal"]
normal.font.name="仿宋_GB2312";normal.font.size=Pt(11)
normal._element.rPr.rFonts.set(qn("w:eastAsia"),"仿宋_GB2312")
def ptext(t,bold=False):
    p=review.add_paragraph()
    p.paragraph_format.space_after=Pt(5)
    r=p.add_run(t)
    r.bold=bold
    return p
p=review.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("智烤管家·V3舞台强化修改说明与8分钟校验")
r.bold=True;r.font.size=Pt(16);r.font.name="黑体"
ptext("本轮只修改讲稿与舞台说法，不改原25页PPT。主要吸收V1生动问答、V2“家庭医生”式生活化比喻，保留V3健康模型与发展模型的算法、已核实数据及四队实操机制。")
ptext("一、V1—V3优势怎样重新融合",True)
table=review.add_table(rows=1,cols=3)
table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,tx in zip(table.rows[0].cells,["过去的优势","本次新版写法","现场价值"]):c.text=tx
rows=[
("V1“请一位管家”亮相","给山里的烤房，请一位懂技术、会算账、能跑腿的管家","开头有画面、有角色"),
("V1“打分、找病灶、开清单”","健康模型：数据入模→权重及寿命测算→风险与维修清单","技术不再像论文宣读"),
("V2“家庭医生”通俗比喻","好比请来一位体检医生，先打分、找病灶、再开清单","信息与烟叶两类评委都听得懂"),
("9.30原稿“匠心工坊”","一座烤房，守着烟农一年的收成","保留温度与真实场景"),
("原稿“三缺”归因","有房未必调得动、有钱未必投得准、有人未必管得久","排比增强辨识度"),
("V3技术核心","健康分80/10/10；有效烤能、AHP/熵权/TOPSIS及数需评推","专业度与关键方法不丢"),
("V3运营落地","48次调剂、18次报修、16单平均0.38小时；1363栋、81%及近190万基金","数据、流程、人员资金完整闭环"),
("V1结尾“数据不睡觉”","数据不睡觉，服务不断线；聚数成方，一脉致远","收尾有节奏和记忆点")
]
for row in rows:
    cs=table.add_row().cells
    for c,tx in zip(cs,row):c.text=tx
ptext("二、A/B台词及总时长",True)
ptext(f"A角{A}个口述汉字（{A/(A+B)*100:.2f}%），B角{B}个（{B/(A+B)*100:.2f}%）；合计{A+B}字，接近7:3。PPT标题和括号内提示均不朗读。")
for speed in (220,230,240):
    spoken=(A+B)/speed*60
    total=spoken+30+25
    ptext(f"语速{speed}汉字/分钟：口述约{spoken/60:.2f}分钟 + 视频30秒 + 换页/停顿25秒 ≈ {int(total//60)}分{int(round(total%60)):02d}秒。")
ptext("开场视频原素材约41.5秒。上述时长必须以裁剪后30秒为前提；第14、16、17、19页视频若再独立完整播放，极可能超时。建议两人至少完成3次实测彩排。")
ptext("三、三项专业边界",True)
ptext("① 第6页四大部位十八个清查项，与第8页十二项健康模型参数不是同一个口径，不能混作一项。② 第8页图中寿命合成测算与第9页威布尔表述，应由技术团队统一参数和现行算法；为避免误导，台上只说“结合使用年限估算剩余寿命”。③ 0.38小时是16条平台工单的平均响应，不是全部18条报修的平均维修完结时长。")
ptext("四、落实口径与现场价值",True)
ptext("正式成果优先采用：秭归2442座全口径建档、健康分低于60的340多座预警对象、48次调剂、18次报修、16条工单平均响应0.38小时、531份问卷综合满意度96.8%、两级合作社统筹管护1363栋/年度在用1672栋约81%、基金总额近190万元。")
ptext("原稿中的“150栋升级已完成”“肖家坪7栋已建成”、雷击案例具体维修点、影像12478张和6704条评价等，如后续要放回主稿，建议先补齐同年度验收或工作台账佐证。")
review_path=DIR/f"{STAMP}_03_V1V2V3优势融合_逐项修改说明与计时核对.docx"
review.save(str(review_path))

readme=[
"# 1019/v3/修改稿｜舞台感染力强化版",
"",
f"- 源文件：{SOURCE.name}",
f"- A角色：{A}口述汉字；B角色：{B}口述汉字；合计：{A+B}；A占比：{A/(A+B)*100:.2f}%；B占比：{B/(A+B)*100:.2f}%。",
"- 30秒开场视频是**裁剪目标**；原PPT视频约41.5秒。",
"- 25页PPT均有对应口播/静默页面（部分连续页合并讲解）。",
"- 原9.30稿的页面和正文风格继承于原始docx（标题黑体14磅、A/B角色黑体14磅、正文仿宋_GB2312 14磅）。",
"- Word输出：舞台稿、净稿、修改说明与计时核对。",
"- 已保留双模技术及真实结果，不虚构未证实的150栋/7栋验收成果。"
]
(DIR/f"{STAMP}_核验和交付说明.md").write_text("\n".join(readme),encoding="utf-8")
print("\n".join(readme))
print("OUTPUTS:",*[x.name for x in [stage_path,clean_path,review_path]])

# verified source-to-Word build trigger

# rerun with table-heavy document QA correction
